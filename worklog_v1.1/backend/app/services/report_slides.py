"""보고자료 PPT의 참고 슬라이드: 업무일지 본문의 표·간트, 프로젝트 마일스톤 일정을 PPT 도형으로 그린다.

- 그림(캡처)이 아니라 PowerPoint 표·막대 도형이라 받은 사람이 PPT에서 바로 고칠 수 있다.
- 글꼴은 주간업무 PPT와 같은 규칙(ea = LG스마트체 Regular, latin = Arial Narrow).
- 한 장에 다 들어가지 않으면 같은 제목에 "(계속)"을 붙여 다음 장으로 나눈다(표는 머리글 행을 반복).
이 모듈은 DB를 모른다. 블록(TableBlock/GanttBlock)을 받아 슬라이드를 만든다.
"""
from __future__ import annotations

import math
import os
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from weekly_report.fonts import EA_REGULAR, LATIN_FONT
from weekly_report.textmetrics import weighted_length
from weekly_report.worklog import tiptap_lines

GRAY_TEXT = "595959"
GRID = "D9D9D9"
BORDER = "808080"
HEADER_FILL = "E7E6E6"
BAR = "4472C4"
BAR_DONE = "2F5597"
MARKER = "C00000"
STATUS_COLOR = {"완료": "A6A6A6", "진행": "4472C4", "지연": "E06666", "예정": "9DC3E6", "보류": "D9D9D9", "취소": "D9D9D9"}

MARGIN_X = Inches(0.4)
TITLE_Y = Inches(0.3)
BODY_Y = Inches(1.05)
BOTTOM_PAD = Inches(0.55)


# ── 블록 ─────────────────────────────────────────────────────────────────────

@dataclass
class CellSpec:
    text: str
    colspan: int = 1
    rowspan: int = 1
    header: bool = False


@dataclass
class TableBlock:
    kind_label: str
    subtitle: str
    rows: list[list[CellSpec]]


@dataclass
class GanttRow:
    label: str
    start: date | None
    end: date | None
    color: str = BAR
    progress: int | None = None
    baseline: tuple[date, date] | None = None
    actual_end: date | None = None
    note: str = ""


@dataclass
class GanttBlock:
    kind_label: str
    subtitle: str
    rows: list[GanttRow]
    display: tuple[date, date] | None = None
    marker: date | None = None  # 보고 기준일 (빨간 세로선)
    legend: str = ""


Block = TableBlock | GanttBlock


def _iso(value: Any) -> date | None:
    try:
        return date.fromisoformat(value) if isinstance(value, str) else None
    except ValueError:
        return None


def blocks_from_document(document: dict | None, subtitle: str) -> list[Block]:
    """Tiptap 문서에서 표·간트 node를 순서대로 꺼낸다 (본문 텍스트는 AI 정리본이 맡으므로 넣지 않는다)."""
    found: list[Block] = []
    doc = (document or {}).get("doc") if isinstance(document, dict) else None

    def walk(node: Any) -> None:
        if not isinstance(node, dict):
            return
        kind = node.get("type")
        if kind == "table":
            rows = []
            for tr in node.get("content") or []:
                cells = []
                for cell in (tr or {}).get("content") or []:
                    attrs = cell.get("attrs") or {}
                    text = "\n".join(tiptap_lines({"type": "doc", "content": cell.get("content") or []}))
                    cells.append(CellSpec(text, max(1, int(attrs.get("colspan") or 1)), max(1, int(attrs.get("rowspan") or 1)),
                                          cell.get("type") == "tableHeader"))
                if cells:
                    rows.append(cells)
            if rows:
                found.append(TableBlock("업무일지 표", subtitle, rows))
            return
        if kind == "gantt":
            attrs = node.get("attrs") or {}
            rows = []
            for item in attrs.get("items") or []:
                s, e = _iso(item.get("startDate")), _iso(item.get("endDate"))
                if not isinstance(item, dict) or s is None or e is None:
                    continue
                pct = item.get("progressPercent")
                rows.append(GanttRow(str(item.get("label") or "").strip() or "(이름 없음)", s, e,
                                     progress=int(pct) if isinstance(pct, (int, float)) else None))
            rng = attrs.get("displayRange") or None
            display = (_iso(rng.get("start")), _iso(rng.get("end"))) if isinstance(rng, dict) else None
            if rows:
                found.append(GanttBlock("업무일지 간트", subtitle, rows,
                                        display if display and all(display) and display[0] <= display[1] else None))
            return
        for child in node.get("content") or []:
            walk(child)

    walk(doc)
    return found


# ── 글자·도형 도우미 ───────────────────────────────────────────────────────

def _fonts(run, size: float, bold: bool = False, color: str | None = None) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    for tag, face in (("a:latin", LATIN_FONT), ("a:ea", EA_REGULAR)):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", face)


def _write(tf, lines: list[str], size: float, *, bold: bool = False, color: str | None = None, align=None) -> None:
    tf.word_wrap = True
    for i, line in enumerate(lines or [""]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None:
            para.alignment = align
        run = para.add_run()
        run.text = line
        _fonts(run, size, bold, color)


def _textbox(slide, x, y, w, h, lines: list[str], size: float, *, bold=False, color=None, align=None, anchor=MSO_ANCHOR.TOP, name=None):
    box = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(max(int(w), 1)), Emu(max(int(h), 1)))
    if name:
        box.name = name
    tf = box.text_frame
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = anchor
    _write(tf, lines, size, bold=bold, color=color, align=align)
    return box


def _rect(slide, x, y, w, h, fill: str, *, line: str | None = None, shape=MSO_SHAPE.RECTANGLE):
    r = slide.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(max(int(w), 1)), Emu(max(int(h), 1)))
    style = r._element.find(qn("p:style"))  # 테마 도형 스타일(그림자·글자색) 대신 아래에서 직접 지정
    if style is not None:
        r._element.remove(style)
    if fill:
        r.fill.solid()
        r.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        r.fill.background()
    if line:
        r.line.color.rgb = RGBColor.from_string(line)
        r.line.width = Pt(0.75)
    else:
        r.line.fill.background()
    return r


def _vline(slide, x, y1, y2, color: str, width: float = 0.5, dash: bool = False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(x)), Emu(int(y1)), Emu(int(x)), Emu(int(y2)))
    style = c._element.find(qn("p:style"))  # 테마 선 스타일(그림자·두께 효과)을 쓰지 않는다
    if style is not None:
        c._element.remove(style)
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(width)
    if dash:
        ln = c.line._get_or_add_ln()
        prst = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.append(prst)
    return c


def _cell_border(cell, color: str = BORDER, width_pt: float = 0.75) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnB", "a:lnT", "a:lnR", "a:lnL"):  # 앞쪽에 순서대로(lnL, lnR, lnT, lnB) 넣는다 — 채우기보다 앞이어야 함
        old = tcpr.find(qn(tag))
        if old is not None:
            tcpr.remove(old)
        ln = tcpr.makeelement(qn(tag), {"w": str(int(Pt(width_pt))), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr)
        ln.append(fill)
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "solid"}))
        tcpr.insert(0, ln)


def blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name in ("빈 화면", "Blank"):
            return layout
    return min(prs.slide_layouts, key=lambda l: len(l.placeholders))


def _new_slide(prs, title: str, subtitle: str):
    slide = prs.slides.add_slide(blank_layout(prs))
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    width = prs.slide_width - 2 * MARGIN_X
    _textbox(slide, MARGIN_X, TITLE_Y, width, Inches(0.42), [title], 16, bold=True, name="appendix_title")
    _textbox(slide, MARGIN_X, TITLE_Y + Inches(0.45), width, Inches(0.25), [subtitle], 9, color=GRAY_TEXT, name="appendix_subtitle")
    return slide


# ── 표 ──────────────────────────────────────────────────────────────────────

def _grid(rows: list[list[CellSpec]]) -> tuple[int, int, list[tuple[int, int, int, int, CellSpec]]]:
    occupied: set[tuple[int, int]] = set()
    placed = []
    n_cols = 0
    for r, row in enumerate(rows):
        c = 0
        for cell in row:
            while (r, c) in occupied:
                c += 1
            rs = min(cell.rowspan, len(rows) - r)
            for rr in range(r, r + rs):
                for cc in range(c, c + cell.colspan):
                    occupied.add((rr, cc))
            placed.append((r, c, rs, cell.colspan, cell))
            c += cell.colspan
            n_cols = max(n_cols, c)
    return len(rows), max(n_cols, 1), placed


def _lines_for(text: str, width_emu: int, size: float) -> int:
    usable_pt = max(width_emu / 12700 - 7, 10)  # 칸 안쪽 여백
    per_line = max(usable_pt / size, 1.0)  # 한글 1자 = 글자 크기 폭(보수적)
    return sum(max(1, math.ceil(weighted_length(line) / per_line)) for line in (text.split("\n") if text else [""]))


def _table_pages(block: TableBlock, avail_w: int, avail_h: int) -> tuple[list[int], int, list[list[int]], float, list, list[int]]:
    n_rows, n_cols, placed = _grid(block.rows)
    size = 9.0 if n_cols <= 8 else 8.0 if n_cols <= 12 else 7.0
    weights = [4.0] * n_cols
    for r, c, rs, cs, cell in placed:
        if cs == 1:
            longest = max((weighted_length(line) for line in cell.text.split("\n")), default=0.0)
            weights[c] = max(weights[c], min(longest, 30.0) + 2.0)
    total = sum(weights)
    widths = [int(avail_w * w / total) for w in weights]
    line_h = size * 1.25 * 12700
    heights = [int(Inches(0.26))] * n_rows
    for r, c, rs, cs, cell in placed:
        if rs == 1:
            need = int(_lines_for(cell.text, sum(widths[c:c + cs]), size) * line_h + Inches(0.08))
            heights[r] = max(heights[r], need)
    header = 0
    while header < n_rows and all(cell.header for r, c, rs, cs, cell in placed if r == header):
        header += 1
    header = min(header, n_rows - 1) if n_rows > 1 else 0
    pages: list[list[int]] = []
    current: list[int] = []
    used = sum(heights[:header])
    for r in range(header, n_rows):
        if current and used + heights[r] > avail_h:
            pages.append(current)
            current, used = [], sum(heights[:header])
        current.append(r)
        used += heights[r]
    if current or not pages:
        pages.append(current)
    return widths, header, pages, size, placed, heights


def _draw_table(slide, block: TableBlock, x, y, widths, heights_all, header, body_rows, size, placed) -> None:
    page_rows = list(range(header)) + body_rows
    index = {r: i for i, r in enumerate(page_rows)}
    n_cols = len(widths)
    shape = slide.shapes.add_table(len(page_rows), n_cols, Emu(int(x)), Emu(int(y)), Emu(sum(widths)),
                                   Emu(sum(heights_all[r] for r in page_rows)))
    shape.name = "appendix_table"
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, w in enumerate(widths):
        tbl.columns[j].width = Emu(w)
    for i, r in enumerate(page_rows):
        tbl.rows[i].height = Emu(heights_all[r])
    for i in range(len(page_rows)):
        for j in range(n_cols):
            cell = tbl.cell(i, j)
            _cell_border(cell)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
    for r, c, rs, cs, cell in placed:
        if r not in index:
            continue
        top = index[r]
        last = max(index[rr] for rr in range(r, r + rs) if rr in index)
        target = tbl.cell(top, c)
        if last > top or cs > 1:
            target.merge(tbl.cell(last, min(c + cs - 1, n_cols - 1)))
        target.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = target.text_frame
        for extra in tf.paragraphs[1:]:  # 병합하면 가려진 칸의 빈 문단이 옮겨 온다 → 지운다
            extra._p.getparent().remove(extra._p)
        for run in tf.paragraphs[0].runs:
            run._r.getparent().remove(run._r)
        _write(tf, cell.text.split("\n") if cell.text else [""], size, bold=cell.header)
        if cell.header:
            target.fill.solid()
            target.fill.fore_color.rgb = RGBColor.from_string(HEADER_FILL)


def table_slides(prs, block: TableBlock, title: str) -> list:
    avail_w = prs.slide_width - 2 * MARGIN_X
    avail_h = prs.slide_height - BODY_Y - BOTTOM_PAD
    widths, header, pages, size, placed, heights = _table_pages(block, avail_w, avail_h)
    slides = []
    for number, body_rows in enumerate(pages):
        slide = _new_slide(prs, title + (" (계속)" if number else ""), block.subtitle)
        _draw_table(slide, block, MARGIN_X, BODY_Y, widths, heights, header, body_rows, size, placed)
        slides.append(slide)
    return slides


# ── 간트 ─────────────────────────────────────────────────────────────────────

def _md(d: date) -> str:
    return f"{d.month}/{d.day}"


def _ticks(lo: date, hi: date) -> list[tuple[date, str]]:
    days = (hi - lo).days + 1
    if days <= 35:
        step = max(1, math.ceil(days / 18))
        return [(lo + timedelta(days=i), _md(lo + timedelta(days=i))) for i in range(0, days, step)]
    if days <= 200:
        first = lo + timedelta(days=(7 - lo.weekday()) % 7)  # 월요일
        out = []
        d = first
        while d <= hi:
            out.append((d, _md(d)))
            d += timedelta(days=7 if days <= 100 else 14)
        return out
    out = []
    d = date(lo.year, lo.month, 1)
    if d < lo:
        d = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    while d <= hi:
        out.append((d, f"'{d.year % 100:02d}.{d.month}"))
        d = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return out


def gantt_slides(prs, block: GanttBlock, title: str) -> list:
    dated = [r for r in block.rows if r.start and r.end]
    lo_candidates = [r.start for r in dated] + [r.baseline[0] for r in block.rows if r.baseline]
    hi_candidates = [r.end for r in dated] + [r.baseline[1] for r in block.rows if r.baseline] + \
                    [r.actual_end for r in block.rows if r.actual_end]
    if block.display:
        lo, hi = block.display
    elif lo_candidates and hi_candidates:
        lo, hi = min(lo_candidates), max(hi_candidates)
    else:
        lo = hi = block.marker or date.today()
    if block.marker and not block.display:
        lo, hi = min(lo, block.marker), max(hi, block.marker)
    if hi < lo:
        lo, hi = hi, lo
    days = (hi - lo).days + 1

    label_w = Inches(2.4)
    x0 = MARGIN_X + label_w
    chart_w = prs.slide_width - MARGIN_X - x0
    header_h = Inches(0.3)
    row_h = Inches(0.34)
    legend_h = Inches(0.25) if block.legend else 0
    avail = prs.slide_height - BODY_Y - BOTTOM_PAD - header_h - legend_h
    per_page = max(1, int(avail // row_h))
    chunks = [block.rows[i:i + per_page] for i in range(0, len(block.rows), per_page)] or [[]]
    day_w = chart_w / days

    def x_of(d: date) -> float:
        return x0 + (min(max(d, lo), hi + timedelta(days=1)) - lo).days * day_w

    slides = []
    for number, rows in enumerate(chunks):
        slide = _new_slide(prs, title + (" (계속)" if number else ""), block.subtitle)
        top = BODY_Y
        bottom = top + header_h + row_h * len(rows)
        _rect(slide, MARGIN_X, top, label_w + chart_w, header_h, HEADER_FILL)
        _textbox(slide, MARGIN_X, top, label_w, header_h, ["항목"], 9, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        for d, text in _ticks(lo, hi):
            x = x_of(d)
            _vline(slide, x, top + header_h, bottom, GRID)
            _textbox(slide, x, top, Inches(0.6), header_h, [text], 7, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)
        for i, row in enumerate(rows):
            y = top + header_h + row_h * i
            if i % 2:
                _rect(slide, MARGIN_X, y, label_w + chart_w, row_h, "F7F7F7")
            span = f"{_md(row.start)}~{_md(row.end)}" if row.start and row.end else "일정 미정"
            _textbox(slide, MARGIN_X, y, label_w - Inches(0.05), row_h * 0.62, [row.label], 9, anchor=MSO_ANCHOR.BOTTOM)
            _textbox(slide, MARGIN_X, y + row_h * 0.58, label_w - Inches(0.05), row_h * 0.42,
                     [span + (f" · {row.note}" if row.note else "")], 7, color=GRAY_TEXT)
            if row.start and row.end and row.end >= lo and row.start <= hi:
                bx, ex = x_of(row.start), x_of(row.end + timedelta(days=1))
                bar_h = row_h * 0.42
                by = y + (row_h - bar_h) / 2 - (Inches(0.03) if row.baseline else 0)
                _rect(slide, bx, by, ex - bx, bar_h, row.color)
                if row.progress:
                    _rect(slide, bx, by, (ex - bx) * max(0, min(row.progress, 100)) / 100, bar_h, BAR_DONE)
                if row.baseline:
                    sx, sy = x_of(row.baseline[0]), x_of(row.baseline[1] + timedelta(days=1))
                    _rect(slide, sx, by + bar_h + Inches(0.03), sy - sx, Inches(0.05), "7F7F7F")
            if row.actual_end and lo <= row.actual_end <= hi:
                ax = x_of(row.actual_end) + day_w / 2
                size = Inches(0.12)
                _rect(slide, ax - size / 2, y + row_h / 2 - size / 2, size, size, "404040", shape=MSO_SHAPE.DIAMOND)
        if block.marker and lo <= block.marker <= hi:
            mx = x_of(block.marker) + day_w / 2
            _vline(slide, mx, top + header_h, bottom, MARKER, 1.25, dash=True)
        border = _rect(slide, MARGIN_X, top, label_w + chart_w, bottom - top, "", line=BORDER)
        border.name = "appendix_gantt_frame"
        if block.legend:
            _textbox(slide, MARGIN_X, bottom + Inches(0.06), label_w + chart_w, legend_h, [block.legend], 7, color=GRAY_TEXT)
        slides.append(slide)
    return slides


# ── 덱에 끼워 넣기 ─────────────────────────────────────────────────────────

@dataclass
class AppendixGroup:
    after_slide: int  # 원래 덱에서 이 번호(1부터) 슬라이드 뒤에 넣는다
    project_name: str
    blocks: list[Block] = field(default_factory=list)


def build_slides(prs, group: AppendixGroup) -> list:
    slides = []
    for block in group.blocks:
        title = f"[참고] {group.project_name} – {block.kind_label}"
        slides += table_slides(prs, block, title) if isinstance(block, TableBlock) else gantt_slides(prs, block, title)
    return slides


def _set_text_keep_format(shape, text: str) -> None:
    para = shape.text_frame.paragraphs[0]
    runs = para.runs
    if not runs:
        _fonts(para.add_run(), 9)
        runs = para.runs
    runs[0].text = text
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)


def insert_appendix(path: Path, groups: list[AppendixGroup]) -> int:
    """원래 슬라이드 순서를 지키며 그룹별 참고 슬라이드를 끼워 넣고, page_no 칸이 있으면 번호를 다시 매긴다."""
    groups = [g for g in groups if g.blocks]
    if not groups:
        return 0
    prs = Presentation(str(path))
    sld_list = prs.slides._sldIdLst
    original = list(sld_list)
    page_proto = next((sh for s in prs.slides for sh in s.shapes if sh.name == "page_no"), None)
    show_master = prs.slides[0]._element.get("showMasterSp") if len(prs.slides) else None
    added: dict[int, list] = {}
    count = 0
    for g in groups:
        before = len(sld_list)
        slides = build_slides(prs, g)
        new_ids = list(sld_list)[before:]
        if show_master is not None:  # 원래 장이 마스터 도형(쪽 번호 등)을 숨기면 참고 장도 같게
            for s in slides:
                s._element.set("showMasterSp", show_master)
        if page_proto is not None:
            for s in slides:
                box = _textbox(s, page_proto.left, page_proto.top, page_proto.width, page_proto.height, [""], 9, bold=True,
                               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="page_no")
                box.text_frame.paragraphs[0].runs[0].text = ""
        added.setdefault(g.after_slide, []).extend(new_ids)
        count += len(slides)
    for el in list(sld_list):
        sld_list.remove(el)
    for i, el in enumerate(original, 1):
        sld_list.append(el)
        for extra in added.get(i, []):
            sld_list.append(extra)
    for extra in added.get(0, []):  # 0 = 맨 앞 (쓰지 않지만 안전하게)
        sld_list.insert(0, extra)
    if page_proto is not None:
        total = len(prs.slides)
        for number, slide in enumerate(prs.slides, 1):
            for sh in slide.shapes:
                if sh.name == "page_no" and sh.has_text_frame:
                    _set_text_keep_format(sh, f"{number} / {total}")
    tmp = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    prs.save(str(tmp))
    tmp.replace(path)
    return count
