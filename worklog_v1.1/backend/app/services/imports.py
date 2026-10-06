"""조직/사용자 Excel·CSV 가져오기. preview → commit, 오류/모호 행이 하나라도 있으면 전량 적용 중단.

- 검증은 transaction 밖(순수 함수)에서 끝내고, 적용만 짧은 단일 transaction 으로 처리한다.
- 매핑 우선순위: 내부 userId → externalKey → 사번. 이름/팀만으로 자동 merge 하지 않는다(모호 → 차단).
- 사번 셀이 숫자 형식이면 이미 앞자리 0이 사라졌을 수 있으므로 오류로 알린다("0" 같은 문자열 사번은 유효).
- 수식 셀('=...')은 값으로 추론하지 않고 오류로 처리한다.
"""
from __future__ import annotations

import csv
import hashlib
import io
import secrets
from datetime import timedelta
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..errors import ApiError, conflict, not_found, validation
from ..models import ImportPreview, Organization, User, utcnow
from .common import audit, bump, mark_export, stamp_new
from .masters import master_revision, norm_text

ORG_COLUMNS = ["코드", "구분", "조직명", "상위 코드"]
USER_COLUMNS = ["이름", "팀 코드", "팀 경로", "사번", "사용자 코드", "사용자 ID", "사용 여부"]
# 양식은 한글 머리글을 쓰고, 내부 키와 예전 영문 머리글도 그대로 받는다. 비교는 공백/대소문자를 무시한다.
_ALIASES_RAW = {
    "코드": "externalKey", "구분": "kind", "조직명": "name", "상위코드": "parentExternalKey",
    "이름": "name", "성명": "name", "팀코드": "teamExternalKey", "팀경로": "teamPath", "팀": "teamPath",
    "사번": "employeeNumber", "사용자코드": "externalKey", "사용자id": "userId", "사용여부": "active",
    **{k.lower(): k for k in ("externalKey", "kind", "name", "parentExternalKey", "teamExternalKey", "teamPath",
                              "employeeNumber", "userId", "active")},
}
KIND_ALIASES = {"division": "division", "team": "team", "담당조직": "division", "본부": "division", "팀": "team"}
GUIDE = {
    "organizations": [
        ("코드", "조직을 구분하는 고유 코드(선택). 다시 가져올 때 같은 조직으로 인식하는 기준입니다. 텍스트로 입력하세요."),
        ("구분", "'담당 조직' 또는 '팀'"),
        ("조직명", "필수"),
        ("상위 코드", "팀이 속한 담당 조직의 코드(선택). 담당 조직은 비워 두세요."),
    ],
    "users": [
        ("이름", "필수"),
        ("팀 코드", "소속 팀의 코드. '팀 경로'와 둘 중 하나는 필요합니다."),
        ("팀 경로", "코드를 모를 때: '담당 조직/팀' 형식(예: 운영담당/업무개선팀)"),
        ("사번", "선택. 비워도 등록됩니다. 앞자리 0이 있으면 반드시 텍스트로 입력하세요(숫자 서식이면 오류로 안내합니다)."),
        ("사용자 코드", "선택. 사번이 없을 때 같은 사람을 구분하는 코드"),
        ("사용자 ID", "선택. 프로그램에서 내보낸 파일을 수정해 다시 가져올 때만 사용"),
        ("사용 여부", "선택. 사용 / 비활성 (비우면 변경하지 않음)"),
    ],
}


def _norm_header(h: Any) -> str:
    key = "".join(str(h).split()).lower()
    return _ALIASES_RAW.get(key, str(h).strip())
TRUE = {"1", "true", "y", "yes", "활성", "사용", "예"}
FALSE = {"0", "false", "n", "no", "비활성", "미사용", "아니오"}
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
PREVIEW_TTL = timedelta(hours=1)


TEMPLATE_ROWS = 500  # 드롭박스/서식을 미리 적용해 두는 행 수(2~501행). 더 필요하면 복사해서 늘리면 된다.


def _workbook_template() -> bytes:
    """조직 + 사용자를 한 파일에서 입력하는 양식.

    - 구분 / 사용 여부 / 상위 코드 / 사용자 시트의 팀 코드는 드롭박스(팀 코드는 조직 시트에 적은 코드 목록).
    - 사용자 시트의 '팀 이름(자동)' 열은 수식으로 채워지며 가져오기에서는 읽지 않는다(알 수 없는 열 + 빈 행 판단에서 제외).
    - 코드/사번 열은 텍스트 서식(앞자리 0 보존).
    """
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    last = TEMPLATE_ROWS + 1
    wb = Workbook()
    ws_o = wb.active
    ws_o.title = "조직"
    ws_u = wb.create_sheet("사용자")
    guide = wb.create_sheet("작성 안내")
    head_fill = PatternFill("solid", fgColor="DCE6FA")
    auto_fill = PatternFill("solid", fgColor="EEEEEE")

    ws_o.append(ORG_COLUMNS)
    ws_u.append(["이름", "팀 코드", "팀 이름(자동)", "사번", "사용자 코드", "사용 여부"])
    for ws, widths in ((ws_o, (16, 14, 26, 16)), (ws_u, (16, 16, 26, 14, 16, 12))):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True), head_fill, Alignment(horizontal="center")
        ws.freeze_panes = "A2"
    for col in (1, 4):  # 조직: 코드, 상위 코드
        for r in range(2, last + 1):
            ws_o.cell(row=r, column=col).number_format = "@"
    for col in (2, 4, 5):  # 사용자: 팀 코드, 사번, 사용자 코드
        for r in range(2, last + 1):
            ws_u.cell(row=r, column=col).number_format = "@"

    def dv(formula: str, rng: str, ws, title: str, msg: str, strict: bool = True) -> None:
        d = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                           errorStyle="stop" if strict else "warning", errorTitle=title, error=msg, showInputMessage=False)
        ws.add_data_validation(d)
        d.add(rng)

    dv('"담당 조직,팀"', f"B2:B{last}", ws_o, "구분 선택", "'담당 조직' 또는 '팀' 중에서 선택하세요.")
    dv(f"$A$2:$A${last}", f"D2:D{last}", ws_o, "상위 코드", "이 시트 '코드' 열에 적은 담당 조직 코드 중에서 선택하세요. (이미 등록된 코드라면 계속 진행)", strict=False)
    dv(f"'조직'!$A$2:$A${last}", f"B2:B{last}", ws_u, "목록에 없는 팀 코드",
       "조직 시트에 적은 코드가 아닙니다. 이미 등록되어 있는 팀의 코드라면 '예'로 계속 진행하세요.", strict=False)
    dv('"사용,비활성"', f"F2:F{last}", ws_u, "사용 여부", "'사용' 또는 '비활성' 중에서 선택하세요.")
    ws_u["C1"].comment = Comment("팀 코드를 고르면 조직 시트의 조직명이 자동으로 표시됩니다. 직접 입력하지 마세요(가져오기에서 읽지 않습니다).", "Worklog")
    for r in range(2, last + 1):
        c = ws_u.cell(row=r, column=3, value=f'=IF(B{r}="","",IFERROR(VLOOKUP(B{r},\'조직\'!$A$2:$C${last},3,FALSE),"※ 조직 시트에 없는 코드"))')
        c.fill, c.font = auto_fill, Font(color="666666")

    guide.append(["순서", "설명"])
    for row in (
        ("1. 조직 시트", "코드 · 구분(드롭박스) · 조직명 · 상위 코드(드롭박스)를 적습니다. 팀의 상위 코드에는 소속 '담당 조직'의 코드를 고릅니다. 코드는 사용자 시트가 팀을 고르는 기준이 되므로 팀에는 코드를 적는 것을 권장합니다."),
        ("2. 사용자 시트", "이름 · 팀 코드(드롭박스: 조직 시트에 적은 코드) · 사번(선택) · 사용자 코드(선택) · 사용 여부(드롭박스)를 적습니다. 팀 코드를 고르면 '팀 이름(자동)'이 표시되니 맞는 팀인지 확인하세요."),
        ("사번", "비워도 등록됩니다. 앞자리 0이 있어도 텍스트 서식이라 그대로 유지됩니다. 서식을 숫자로 바꾸면 오류로 안내합니다."),
        ("3. 가져오기", "화면의 '조직·사용자 > Excel/CSV 가져오기'에서 '조직 + 사용자 (한 파일)'로 올립니다. 미리보기에서 오류나 모호한 행이 하나라도 있으면 조직·사용자 모두 적용되지 않습니다."),
        ("기존 데이터", "코드가 같으면 같은 조직으로 보고 이름 등을 수정합니다. 파일에 없는 조직·사용자는 삭제하거나 비활성화하지 않습니다."),
        ("주의", "시트 이름('조직', '사용자')과 머리글(1행)을 바꾸지 마세요. '팀 이름(자동)' 열은 수식이라 지우거나 덮어써도 가져오기에는 영향이 없습니다."),
    ):
        guide.append(list(row))
    guide.column_dimensions["A"].width = 16
    guide.column_dimensions["B"].width = 110
    for c in guide[1]:
        c.font, c.fill = Font(bold=True), head_fill
    for r in guide.iter_rows(min_row=2):
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def template(entity: str, fmt: str) -> tuple[bytes, str, str]:
    xlsx_mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if entity == "workbook":
        if fmt != "xlsx":
            raise validation("조직+사용자 한 파일 양식은 xlsx 만 제공합니다.")
        return _workbook_template(), xlsx_mime, "worklog-import-template.xlsx"
    cols = ORG_COLUMNS if entity == "organizations" else USER_COLUMNS
    if fmt == "csv":
        buf = io.StringIO()
        csv.writer(buf).writerow(cols)
        return ("﻿" + buf.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", f"{entity}-template.csv"
    wb = Workbook()
    ws = wb.active
    ws.title = "입력"
    ws.append(cols)
    text_cols = {"코드", "상위 코드", "팀 코드", "사번", "사용자 코드", "사용자 ID"}
    for idx, name in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = 18
        if name in text_cols:
            for r in range(2, 502):
                ws.cell(row=r, column=idx).number_format = "@"  # 텍스트 형식: 앞자리 0 보존
    guide = wb.create_sheet("작성 안내")  # 첫 번째 시트만 읽으므로 안내 시트는 가져오기에 영향 없음
    guide.append(["열 이름", "설명"])
    for row in GUIDE[entity]:
        guide.append(list(row))
    guide.column_dimensions["A"].width = 16
    guide.column_dimensions["B"].width = 90
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{entity}-template.xlsx"


# ── 파일 읽기 ───────────────────────────────────────────────────────────────

KNOWN_KEYS = {"externalKey", "kind", "name", "parentExternalKey", "teamExternalKey", "teamPath", "employeeNumber", "userId", "active"}
SHEET_ALIASES = {"조직": "organizations", "organizations": "organizations", "organization": "organizations", "org": "organizations",
                 "사용자": "users", "users": "users", "user": "users"}


def _sheet_kind(title: str) -> str | None:
    return SHEET_ALIASES.get("".join(str(title).split()).lower())


def _load_xlsx(data: bytes):
    try:
        return load_workbook(io.BytesIO(data), read_only=False, data_only=False)
    except Exception:  # noqa: BLE001 - zip/xml 손상 전반
        raise validation("xlsx 파일을 읽을 수 없습니다.") from None


def _rows_from_grid(grid: list[list[Any]]) -> list[dict[str, Any]]:
    if not grid:
        raise validation("파일이 비어 있습니다.")
    header = [_norm_header(h) if h is not None else "" for h in grid[0]]
    known = [j for j, h in enumerate(header) if h in KNOWN_KEYS]
    rows = []
    for i, raw in enumerate(grid[1:], start=2):
        # 비어 있는 행 판단은 '인식하는 열'만 본다: 양식의 자동 입력 열(팀 이름 수식 등)이 채워져 있어도 빈 행으로 취급해야 한다.
        if all((raw[j] if j < len(raw) else None) is None or str(raw[j]).strip() == "" for j in known):
            continue
        rows.append({"_row": i, **{header[j]: (raw[j] if j < len(raw) else None) for j in range(len(header)) if header[j]}})
    if len(rows) > MAX_ROWS:
        raise validation(f"한 번에 {MAX_ROWS}행까지 가져올 수 있습니다.")
    return rows


def _read_table(filename: str, data: bytes, encoding: str, sheet: str | None = None) -> list[dict[str, Any]]:
    """CSV 또는 xlsx 한 시트를 읽는다. xlsx 에 '조직'/'사용자' 시트가 있으면 sheet 로 해당 시트를 고르고, 없으면 첫 시트."""
    if len(data) > MAX_IMPORT_BYTES:
        raise ApiError(413, "FILE_TOO_LARGE", "가져오기 파일이 너무 큽니다.")
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig" if encoding in {"utf-8", "utf8", "utf-8-sig"} else encoding)
        except (UnicodeDecodeError, LookupError):
            raise validation("CSV를 읽을 수 없습니다. UTF-8로 저장하거나 인코딩을 선택해 주세요.") from None
        grid = [row for row in csv.reader(io.StringIO(text))]
    elif name.endswith(".xlsx"):
        wb = _load_xlsx(data)
        ws = next((w for w in wb.worksheets if sheet and _sheet_kind(w.title) == sheet), wb.worksheets[0])
        grid = [[c.value for c in row] for row in ws.iter_rows()]
    else:
        raise ApiError(415, "UNSUPPORTED_MEDIA", "xlsx 또는 csv 파일만 가져올 수 있습니다.")
    return _rows_from_grid(grid)


def _read_workbook(filename: str, data: bytes) -> tuple[list[dict], list[dict]]:
    """한 파일(xlsx)의 '조직' 시트와 '사용자' 시트를 함께 읽는다. 둘 중 없는 시트는 빈 목록."""
    if len(data) > MAX_IMPORT_BYTES:
        raise ApiError(413, "FILE_TOO_LARGE", "가져오기 파일이 너무 큽니다.")
    if not filename.lower().endswith(".xlsx"):
        raise ApiError(415, "UNSUPPORTED_MEDIA", "조직+사용자 한 파일 가져오기는 xlsx 만 지원합니다. CSV 는 조직/사용자를 따로 가져오세요.")
    wb = _load_xlsx(data)
    found: dict[str, list[dict]] = {}
    for ws in wb.worksheets:
        kind = _sheet_kind(ws.title)
        if kind and kind not in found:
            grid = [[c.value for c in row] for row in ws.iter_rows()]
            found[kind] = _rows_from_grid(grid) if grid else []
    if not found:
        raise validation("'조직' 또는 '사용자' 시트를 찾을 수 없습니다. 내려받은 양식의 시트 이름을 바꾸지 마세요.")
    return found.get("organizations", []), found.get("users", [])


def _text_cell(row: dict, col: str, errors: list[str], *, label: str) -> str | None:
    v = row.get(col)
    if v is None:
        return None
    if isinstance(v, bool) or isinstance(v, (int, float)):
        errors.append(f"{label}이(가) 숫자 형식입니다. 앞자리 0이 이미 사라졌을 수 있으니 텍스트 형식으로 다시 입력해 주세요.")
        return None
    s = str(v).strip()
    if s.startswith("="):
        errors.append(f"{label}에 수식이 있습니다. 값을 직접 입력해 주세요.")
        return None
    return s or None  # "0" 은 유지


def _bool_cell(v: Any, errors: list[str]) -> bool | None:
    if v is None or str(v).strip() == "":
        return None
    t = str(v).strip().lower()
    if t in TRUE:
        return True
    if t in FALSE:
        return False
    errors.append("활성 여부 값을 알 수 없습니다. (true/false, 1/0, Y/N)")
    return None


# ── 미리보기 ────────────────────────────────────────────────────────────────

def _preview_orgs(s: Session, rows: list[dict]) -> list[dict]:
    existing = list(s.execute(select(Organization).where(Organization.deleted_at.is_(None))).scalars())
    by_key = {o.external_key: o for o in existing if o.external_key}
    out: list[dict] = []
    keys_in_file: dict[str, int] = {}
    for r in rows:
        errs: list[str] = []
        key = _text_cell(r, "externalKey", errs, label="코드")
        kind = KIND_ALIASES.get("".join(str(r.get("kind") or "").split()).lower())
        name = (str(r.get("name") or "").strip()) or None
        parent_key = _text_cell(r, "parentExternalKey", errs, label="상위 코드")
        if kind is None:
            errs.append("구분(kind)은 division 또는 team 이어야 합니다.")
        if not name:
            errs.append("조직명이 필요합니다.")
        if key:
            keys_in_file[key] = keys_in_file.get(key, 0) + 1
        out.append({"row": r["_row"], "values": {"externalKey": key, "kind": kind, "name": name, "parentExternalKey": parent_key},
                    "errors": errs, "action": None, "matchedId": None})
    file_by_key = {o["values"]["externalKey"]: o for o in out if o["values"]["externalKey"]}
    for o in out:
        v = o["values"]
        if v["externalKey"] and keys_in_file[v["externalKey"]] > 1:
            o["errors"].append("파일 안에서 코드가 중복됩니다.")
        parent_kind = None
        if v["parentExternalKey"]:
            pk = v["parentExternalKey"]
            if pk == v["externalKey"]:
                o["errors"].append("자기 자신을 상위 조직으로 지정할 수 없습니다.")
            elif pk in file_by_key:
                parent_kind = file_by_key[pk]["values"]["kind"]
            elif pk in by_key:
                parent_kind = by_key[pk].kind
            else:
                o["errors"].append(f"상위 조직 코드({pk})를 찾을 수 없습니다.")
            if parent_kind and parent_kind != "division":
                o["errors"].append("상위 조직은 담당 조직(division)이어야 합니다.")
            if v["kind"] == "division":
                o["errors"].append("division은 상위 조직을 가질 수 없습니다.")
        match = by_key.get(v["externalKey"]) if v["externalKey"] else None
        if match is not None:
            o["matchedId"] = match.id
            changed = (match.name != v["name"] or match.kind != v["kind"] or
                       (_parent_key_of(match, existing) or None) != (v["parentExternalKey"] or None))
            if match.kind != v["kind"]:
                o["errors"].append("기존 조직의 구분(kind)은 바꿀 수 없습니다.")
            o["action"] = "update" if changed else "unchanged"
        elif not v["externalKey"]:
            same = [e for e in existing if e.kind == v["kind"] and e.name == v["name"]]
            o["action"] = "ambiguous" if same else "new"
            if same:
                o["errors"].append("같은 이름의 조직이 이미 있습니다. 같은 조직이면 코드를 지정하고, 다른 조직이면 구분되는 이름을 사용해 주세요.")
        else:
            o["action"] = "new"
        if o["errors"] and o["action"] != "ambiguous":
            o["action"] = "error"
    return out


def _parent_key_of(org: Organization, all_orgs: list[Organization]) -> str | None:
    if not org.parent_id:
        return None
    parent = next((p for p in all_orgs if p.id == org.parent_id), None)
    return parent.external_key if parent else None


def _team_candidates(orgs: list[Organization], file_orgs: list[dict]) -> tuple[list[dict], set[str]]:
    """사용자의 팀 후보 목록. DB 의 팀 + 같은 파일 조직 시트의 팀(신규/수정)을 합친다.

    후보: {id(DB id 또는 None=이번에 새로 만들 팀), row(조직 시트 행 번호 또는 None), key, name, parent(상위 조직명), bad(그 조직 행에 오류)}
    같은 파일에서 수정되는 기존 팀은 파일의 값(이름/상위)을 기준으로 삼는다. division_keys 는 '팀이 아닌 코드' 안내용.
    """
    name_by_key = {o.external_key: o.name for o in orgs if o.external_key}
    file_by_id = {fo["matchedId"]: fo for fo in file_orgs if fo.get("matchedId")}
    for fo in file_orgs:
        k, n = fo["values"].get("externalKey"), fo["values"].get("name")
        if k and n:
            name_by_key[k] = n
    parent_name_of = {o.id: next((p.name for p in orgs if p.id == o.parent_id), None) for o in orgs}
    cands: list[dict] = []
    division_keys = {o.external_key for o in orgs if o.kind == "division" and o.external_key}
    for t in (o for o in orgs if o.kind == "team"):
        fo = file_by_id.get(t.id)
        if fo:
            v = fo["values"]
            cands.append({"id": t.id, "row": fo["row"], "key": t.external_key or v.get("externalKey"), "name": v.get("name") or t.name,
                          "parent": name_by_key.get(v.get("parentExternalKey") or "") or parent_name_of.get(t.id), "bad": bool(fo["errors"])})
        else:
            cands.append({"id": t.id, "row": None, "key": t.external_key, "name": t.name, "parent": parent_name_of.get(t.id), "bad": False})
    for fo in file_orgs:
        v = fo["values"]
        if fo.get("matchedId"):
            continue
        if v.get("kind") == "team":
            cands.append({"id": None, "row": fo["row"], "key": v.get("externalKey"), "name": v.get("name"),
                          "parent": name_by_key.get(v.get("parentExternalKey") or ""), "bad": bool(fo["errors"])})
        elif v.get("kind") == "division" and v.get("externalKey"):
            division_keys.add(v["externalKey"])
    return cands, division_keys


def _preview_users(s: Session, rows: list[dict], file_orgs: list[dict] | None = None) -> list[dict]:
    orgs = list(s.execute(select(Organization).where(Organization.deleted_at.is_(None))).scalars())
    cand_teams, division_keys = _team_candidates(orgs, file_orgs or [])
    users = list(s.execute(select(User).where(User.deleted_at.is_(None))).scalars())
    by_id = {u.id: u for u in users}
    by_ext = {u.external_key: u for u in users if u.external_key}
    by_emp = {u.employee_number: u for u in users if u.employee_number}
    out: list[dict] = []
    seen: dict[str, dict[str, int]] = {"userId": {}, "externalKey": {}, "employeeNumber": {}}
    for r in rows:
        errs: list[str] = []
        name = (str(r.get("name") or "").strip()) or None
        if not name:
            errs.append("이름이 필요합니다.")
        emp = _text_cell(r, "employeeNumber", errs, label="사번")
        ext = _text_cell(r, "externalKey", errs, label="사용자 코드")
        uid = _text_cell(r, "userId", errs, label="userId")
        team_key = _text_cell(r, "teamExternalKey", errs, label="팀 코드")
        team_path = _text_cell(r, "teamPath", errs, label="팀 경로")
        active = _bool_cell(r.get("active"), errs)
        team = None
        if team_key:
            hits = [t for t in cand_teams if t["key"] == team_key]
            if len(hits) == 1:
                team = hits[0]
            elif team_key in division_keys:
                errs.append(f"코드({team_key})는 팀이 아니라 담당 조직입니다. 소속 팀의 코드를 선택해 주세요.")
            else:
                errs.append(f"팀 코드({team_key})를 찾을 수 없습니다. 조직 시트에 먼저 적거나 이미 등록된 팀 코드를 입력해 주세요.")
        elif team_path:
            parts = [p.strip() for p in team_path.replace(">", "/").split("/") if p.strip()]
            hits = [t for t in cand_teams if parts and t["name"] == parts[-1]]
            if len(parts) > 1:
                hits = [t for t in hits if t["parent"] == parts[-2]]
            if len(hits) == 1:
                team = hits[0]
            else:
                errs.append("팀 경로로 팀을 하나로 특정할 수 없습니다. 팀 코드를 사용해 주세요." if hits else f"팀({team_path})을 찾을 수 없습니다.")
        else:
            errs.append("팀 코드 또는 팀 경로가 필요합니다.")
        if team and team["bad"]:
            errs.append(f"이 팀은 조직 시트 {team['row']}행에 오류가 있어 사용할 수 없습니다. 조직 시트를 먼저 고쳐 주세요.")
        for k, v in (("userId", uid), ("externalKey", ext), ("employeeNumber", emp)):
            if v:
                seen[k][v] = seen[k].get(v, 0) + 1
        out.append({"row": r["_row"], "values": {"name": name, "teamId": team["id"] if team else None, "teamOrgRow": team["row"] if team else None,
                                                  "teamName": team["name"] if team else None,
                                                  "employeeNumber": emp, "externalKey": ext, "userId": uid, "active": active},
                    "errors": errs, "action": None, "matchedId": None})
    for o in out:
        v = o["values"]
        for k in ("userId", "externalKey", "employeeNumber"):
            if v[k] and seen[k][v[k]] > 1:
                o["errors"].append(f"파일 안에서 {k} 값이 중복됩니다.")
        cands: dict[str, User] = {}
        if v["userId"]:
            u = by_id.get(v["userId"])
            if u is None:
                o["errors"].append("userId에 해당하는 기존 사용자가 없습니다.")
            else:
                cands["userId"] = u
        if v["externalKey"] and v["externalKey"] in by_ext:
            cands["externalKey"] = by_ext[v["externalKey"]]
        if v["employeeNumber"] and v["employeeNumber"] in by_emp:
            cands["employeeNumber"] = by_emp[v["employeeNumber"]]
        distinct = {u.id for u in cands.values()}
        if len(distinct) > 1:
            o["errors"].append("userId/코드/사번이 서로 다른 기존 사용자를 가리킵니다.")
        match = next(iter(cands.values())) if cands else None
        if match is not None and len(distinct) == 1:
            o["matchedId"] = match.id
            # 매칭 사용자가 가진 다른 키와 충돌 확인(예: 사번을 바꾸려는데 다른 사람이 사용 중)
            if v["externalKey"] and v["externalKey"] in by_ext and by_ext[v["externalKey"]].id != match.id:
                o["errors"].append("사용자 코드가 다른 사용자에게 이미 사용 중입니다.")
            if v["employeeNumber"] and v["employeeNumber"] in by_emp and by_emp[v["employeeNumber"]].id != match.id:
                o["errors"].append("사번이 다른 사용자에게 이미 사용 중입니다.")
            changed = (match.name != v["name"] or match.team_id != v["teamId"] or
                       (v["employeeNumber"] is not None and match.employee_number != v["employeeNumber"]) or
                       (v["externalKey"] is not None and match.external_key != v["externalKey"]) or
                       (v["active"] is not None and match.active != v["active"]))
            o["action"] = "update" if changed else "unchanged"
        elif not cands and not v["userId"]:
            has_key = bool(v["externalKey"] or v["employeeNumber"])
            same = [u for u in users if v["teamId"] and u.name == v["name"] and u.team_id == v["teamId"]]  # 새로 만들 팀(teamId 없음)에는 기존 사용자가 없다
            if not has_key and same:
                o["action"] = "ambiguous"
                o["errors"].append("같은 팀에 같은 이름의 사용자가 이미 있습니다. 사번 또는 사용자 코드로 같은 사람인지 구분해 주세요.")
            else:
                o["action"] = "new"
        if o["errors"] and o["action"] != "ambiguous":
            o["action"] = "error"
        elif o["action"] is None:
            o["action"] = "error"
    return out


ACTIONS = ("new", "update", "unchanged", "ambiguous", "error")
SECTION_LABEL = {"organizations": "조직", "users": "사용자"}


def _summary(rows: list[dict]) -> dict:
    return {**{k: sum(1 for r in rows if r["action"] == k) for k in ACTIONS}, "total": len(rows)}


def create_preview(s: Session, entity: str, filename: str, data: bytes, encoding: str = "utf-8") -> dict:
    """entity: organizations | users | workbook(한 파일의 '조직'+'사용자' 시트를 함께 검증, 전량 적용/전량 중단)"""
    if entity not in {"organizations", "users", "workbook"}:
        raise validation("entity는 workbook, organizations, users 중 하나여야 합니다.")
    if entity == "workbook":
        org_rows, user_rows = _read_workbook(filename, data)
        org_res = _preview_orgs(s, org_rows) if org_rows else []
        user_res = _preview_users(s, user_rows, org_res)  # 같은 파일에서 새로 만드는 팀도 사용자의 팀 후보가 된다
        sections = {"organizations": org_res, "users": user_res}
        stored: Any = sections
    else:
        rows = _read_table(filename, data, encoding.lower(), sheet=entity)
        res = _preview_orgs(s, rows) if entity == "organizations" else _preview_users(s, rows)
        sections = {entity: res}
        stored = res
    all_rows = [r for rs in sections.values() for r in rs]
    summary = _summary(all_rows)
    has_errors = summary["error"] > 0 or summary["ambiguous"] > 0
    token = secrets.token_urlsafe(24)
    s.add(ImportPreview(token=token, entity=entity, file_hash=hashlib.sha256(data).hexdigest(), rows=stored,
                        has_errors=has_errors, base_revision=master_revision(s), expires_at=utcnow() + PREVIEW_TTL))
    return {"previewToken": token, "entity": entity, "summary": summary,
            "rows": all_rows if entity != "workbook" else [],
            "sections": [{"entity": k, "label": SECTION_LABEL[k], "summary": _summary(v), "rows": v} for k, v in sections.items() if v or entity != "workbook"],
            "canCommit": not has_errors and bool(all_rows)}


# ── 적용 ────────────────────────────────────────────────────────────────────

def commit_preview(s: Session, token: str, actor_id: str | None) -> dict:
    p = s.get(ImportPreview, token)
    if p is None:
        raise not_found("가져오기 미리보기")
    if p.committed_at is not None:
        raise conflict("PREVIEW_CONSUMED", "이미 적용된 미리보기입니다.")
    if p.expires_at < utcnow():
        raise conflict("PREVIEW_EXPIRED", "미리보기가 만료되었습니다. 파일을 다시 올려 주세요.")
    if p.has_errors:
        raise ApiError(422, "IMPORT_HAS_ERRORS", "오류 또는 모호한 행이 있어 적용할 수 없습니다. 파일을 수정해 다시 올려 주세요.")
    if p.base_revision != master_revision(s):
        raise conflict("PREVIEW_STALE", "미리보기 이후 조직/사용자 정보가 변경되었습니다. 파일을 다시 올려 미리보기를 새로 만들어 주세요.")
    rows = p.rows
    detail: dict[str, Any] = {}
    if p.entity == "workbook":
        # 하나의 트랜잭션: 조직을 먼저 적용하고(새 팀 ID 확정), 그 ID 로 사용자를 연결한다. 어느 단계든 실패하면 전부 롤백된다.
        oc = {"created": 0, "updated": 0, "unchanged": 0}
        uc = {"created": 0, "updated": 0, "unchanged": 0}
        row_to_id = _apply_orgs(s, rows.get("organizations", []), actor_id, oc)
        _apply_users(s, rows.get("users", []), actor_id, uc, row_to_id)
        counts = {k: oc[k] + uc[k] for k in oc}
        detail = {"organizations": oc, "users": uc}
    elif p.entity == "organizations":
        counts = {"created": 0, "updated": 0, "unchanged": 0}
        _apply_orgs(s, rows, actor_id, counts)
    else:
        counts = {"created": 0, "updated": 0, "unchanged": 0}
        _apply_users(s, rows, actor_id, counts)
    p.committed_at = utcnow()
    audit(s, f"import.{p.entity}", "import", token[:36], actor_id, {**counts, **detail})
    mark_export(s, "master")
    return {"entity": p.entity, **counts, **detail}


def _apply_orgs(s: Session, rows: list[dict], actor_id: str | None, counts: dict) -> dict[int, str]:
    """조직을 적용하고 {조직 시트 행 번호: 조직 ID} 를 돌려준다(같은 파일의 사용자가 새 팀을 참조할 때 사용)."""
    key_to_id = {o.external_key: o.id for o in s.execute(select(Organization).where(Organization.external_key.is_not(None))).scalars()}
    row_to_id: dict[int, str] = {}
    ordered = sorted(rows, key=lambda r: 0 if r["values"]["kind"] == "division" else 1)  # division 먼저
    for r in ordered:
        v = r["values"]
        if r["action"] == "unchanged":
            counts["unchanged"] += 1
            row_to_id[r["row"]] = r["matchedId"]
            continue
        parent_id = key_to_id.get(v["parentExternalKey"]) if v["parentExternalKey"] else None
        if r["action"] == "new":
            org = Organization(name=v["name"], kind=v["kind"], parent_id=parent_id, external_key=v["externalKey"])
            stamp_new(org, actor_id)
            s.add(org)
            s.flush()
            if org.external_key:
                key_to_id[org.external_key] = org.id
            row_to_id[r["row"]] = org.id
            counts["created"] += 1
        else:
            org = s.get(Organization, r["matchedId"])
            org.name, org.parent_id = v["name"], parent_id
            bump(org, actor_id)
            row_to_id[r["row"]] = org.id
            counts["updated"] += 1
    s.flush()
    return row_to_id


def _apply_users(s: Session, rows: list[dict], actor_id: str | None, counts: dict, row_to_id: dict[int, str] | None = None) -> None:
    for r in rows:
        v = r["values"]
        if r["action"] == "unchanged":
            counts["unchanged"] += 1
            continue
        # 같은 파일에서 새로 만든 팀은 조직 시트 행 번호(teamOrgRow)로 방금 확정된 ID 를 찾는다
        team_id = v["teamId"] or (row_to_id or {}).get(v.get("teamOrgRow"))
        if not team_id:
            raise ApiError(409, "TEAM_NOT_RESOLVED", "사용자의 팀을 확정하지 못했습니다. 미리보기를 다시 만들어 주세요.")
        if r["action"] == "new":
            u = User(name=v["name"], team_id=team_id, employee_number=v["employeeNumber"], external_key=v["externalKey"],
                     active=True if v["active"] is None else v["active"])
            stamp_new(u, actor_id)
            s.add(u)
            counts["created"] += 1
        else:
            u = s.get(User, r["matchedId"])
            u.name, u.team_id = v["name"], team_id
            if v["employeeNumber"] is not None:
                u.employee_number = v["employeeNumber"]
            if v["externalKey"] is not None:
                u.external_key = v["externalKey"]
            if v["active"] is not None:
                u.active = v["active"]
            bump(u, actor_id)
            counts["updated"] += 1
        s.flush()
