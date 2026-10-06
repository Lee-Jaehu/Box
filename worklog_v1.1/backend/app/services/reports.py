"""보고자료(PPT) 생성. weekly_report 패키지(backend/weekly_report, D:\\claude\\weekly-report 에서 가져옴)를 서비스에 연결한다.

흐름 (작업 1건)
  1) DB → WorkLog export 형식(JSON)을 작업공간(data/reports/workspace/data/worklog)에 새로 쓴다 (JSON 사본 폴더를 읽지 않음: 사본은 늦을 수 있다)
  2) 정리(AI): 주간 = ISO 주차별 weekly_rollup + cumulative_update, 기간 = period_rollup + cumulative_update
  3) PPT: 주간업무 양식(과제당 1~2장) 또는 경영진 1장 요약 양식(과제당 1장), 월간 종합(현황표)
  4) 참고 슬라이드(선택): 업무일지 표·간트, 프로젝트 마일스톤 일정 → PPT 도형
AI 연결
  - 서버 설정 ai_api_url·ai_api_key가 있으면 서버가 직접 호출(live). 모든 사용자가 같은 설정을 쓴다.
  - 없으면 붙여넣기(paste): 작업이 'AI 응답 필요'로 멈추고, 화면에서 프롬프트를 복사해 다른 AI에 보낸 뒤 받은 JSON을 붙여 넣으면 이어서 진행한다.
  - 받은 응답은 프롬프트 해시와 함께 저장한다. 업무일지가 바뀌어 프롬프트가 달라지면 예전 응답을 쓰지 않는다(다시 묻는다).
작업 상태는 DB가 아니라 data/reports/jobs/*.json 에 둔다(생성물이라 스키마 변경 없이 추가). 작업은 서버 안 작업 스레드 하나가 순서대로 처리한다.
"""
from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import shutil
import threading
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from sqlalchemy import select

from ..config import Settings
from ..documents import title_preview
from ..errors import ApiError
from ..models import DailyLog, Milestone, Project
from . import exports as export_svc
from .report_slides import STATUS_COLOR, AppendixGroup, Block, GanttBlock, GanttRow, TableBlock, blocks_from_document, insert_appendix

from weekly_report import prompt_vars as pv
from weekly_report import sources as wr_sources
from weekly_report.ai import (RESPONSE_SHAPES, AIError, ChatCompletionsAdapter, ExaoneClient, MockResponseMissing, ResponseFormatError,
                               UrlLibTransport, read_payload)
from weekly_report.core import ValidationError, load_json, week_range
from weekly_report.ppt.budget import BudgetError
from weekly_report.codes import PeopleTable
from weekly_report.ppt.compose import updated_at_label
from weekly_report.ppt.render import open_deck_template
from weekly_report.pptgen import find_template, generate_ppt, prepare_ppt
from weekly_report.report.generate import exec_fill, generate_monthly
from weekly_report.report.render import find_report_template, inspect_report, render_report
from weekly_report.report_vars import month_weeks
from weekly_report.summary import run_summary
from weekly_report.team import (FONT_SIZES, ITEM_MAR_IN, Block as SummaryBlock, SummaryGeometry, TeamSection, inspect_summary,
                                layout_summary, lines_per_project, render_team_report, summary_geometry, summary_titles)
from weekly_report.weekly import iso_week, run_weekly
from weekly_report.worklog import WorklogError

log = logging.getLogger("worklog.reports")

ASSETS = Path(__file__).resolve().parents[2] / "report_assets"
KINDS = ("weekly", "period", "monthly")
TEMPLATES = ("weekly", "exec")
KIND_LABEL = {"weekly": "주간 보고", "period": "기간 보고", "monthly": "월간 종합"}
TEMPLATE_LABEL = {"weekly": "주간업무 양식", "exec": "경영진 1장 요약 양식", "monthly": "월간 종합 양식"}
PROMPT_LABEL = {"weekly_rollup": "주간 정리", "period_rollup": "기간 정리", "cumulative_update": "누적 요약",
                "report_exec_summary": "경영진 1장 요약", "report_monthly": "월간 종합", "fit_to_budget": "분량 줄이기",
                "project_summary": "팀장 요약"}
MS_STATUS_KO = {"planned": "예정", "in_progress": "진행", "on_hold": "보류", "completed": "완료", "cancelled": "취소"}
MAX_LOG_BLOCKS = 15
MAX_PROJECTS = 30
MAX_PERIOD_DAYS = 366
JOB_ID_RE = re.compile(r"^[0-9a-f]{32}$")
NAME_RE = re.compile(r"^[A-Za-z0-9_]+__[A-Za-z0-9-]+__[A-Za-z0-9_-]+\.json$")

_store_lock = threading.RLock()
_run_lock = threading.Lock()


class ReportError(ValueError):
    """사용자에게 그대로 보여 줄 수 있는 생성 실패 사유."""


class NeedResponse(Exception):
    """붙여넣기 방식에서 AI 응답이 필요해 작업을 멈춘다."""

    def __init__(self, prompt_id: str, name: str, prompt: str, sha: str, project_id: str):
        super().__init__(f"AI 응답 필요: {name}")
        self.prompt_id, self.name, self.prompt, self.sha, self.project_id = prompt_id, name, prompt, sha, project_id


# ── 시간·경로 ────────────────────────────────────────────────────────────────

def _now(settings: Settings) -> str:
    return datetime.now(ZoneInfo(settings.timezone)).isoformat(timespec="seconds")


def _today(settings: Settings) -> date:
    return datetime.now(ZoneInfo(settings.timezone)).date()


def _atomic_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.{os.getpid()}.{threading.get_ident()}.tmp")
    tmp.write_text(text, encoding="utf-8")
    tmp.replace(path)


def _atomic_json(path: Path, value: Any) -> None:
    _atomic_text(path, json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def workspace(settings: Settings) -> Path:
    return settings.reports_dir / "workspace"


def responses_dir(settings: Settings) -> Path:
    return workspace(settings) / "responses"


def sync_workspace(settings: Settings) -> Path:
    """프로그램 폴더의 보고 자료 원본(report_assets)을 작업공간으로 맞춘다. 사용자 생성물(정리 결과·응답)은 건드리지 않는다."""
    ws = workspace(settings)
    for sub in ("config", "schemas", "prompts"):
        for src in (ASSETS / sub).rglob("*"):
            if src.is_file():
                dst = ws / sub / src.relative_to(ASSETS / sub)
                if not dst.exists() or dst.read_bytes() != src.read_bytes():
                    dst.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copyfile(src, dst)
    for src in (ASSETS / "templates").glob("*.pptx"):  # 원본 배치와 같게 작업공간 바로 아래 (글꼴 탐색 기준)
        dst = ws / src.name
        if not dst.exists() or dst.stat().st_size != src.stat().st_size or dst.read_bytes() != src.read_bytes():
            shutil.copyfile(src, dst)
    fonts = ASSETS / "fonts"
    if fonts.is_dir():
        for src in fonts.glob("*.[tT][tT][fF]"):
            dst = ws / "fonts" / src.name
            if not dst.exists() or dst.stat().st_size != src.stat().st_size:
                dst.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(src, dst)
    (ws / "responses").mkdir(parents=True, exist_ok=True)
    return ws


def fonts_available() -> bool:
    return any((ASSETS / "fonts").glob("LGSM*.[tT][tT][fF]")) if (ASSETS / "fonts").is_dir() else False


# ── AI 응답 저장소 / 클라이언트 ───────────────────────────────────────────────

def prompt_sha(prompt: str) -> str:
    return hashlib.sha256(prompt.encode("utf-8")).hexdigest()


def response_name(prompt_id: str, project_id: str, key: str, variant: str | None = None) -> str:
    return f"{prompt_id}__{project_id}__{key}" + (f"__{variant}" if variant else "") + ".json"


def read_cached(folder: Path, name: str, sha: str, not_before: str | None) -> str | None:
    path, meta_path = folder / name, folder / f"{name}.meta.json"
    if not path.is_file() or not meta_path.is_file():
        return None
    try:
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if meta.get("promptSha") != sha:
        return None  # 업무일지·프롬프트가 바뀐 뒤의 요청 → 예전 응답은 쓰지 않는다
    if not_before and str(meta.get("savedAt") or "") < not_before:
        return None  # 'AI 응답 새로 받기'를 고른 작업
    return path.read_text(encoding="utf-8")


def save_response(folder: Path, name: str, payload: Any, sha: str, source: str, saved_at: str, saved_by: str | None = None) -> None:
    _atomic_json(folder / name, payload)
    _atomic_json(folder / f"{name}.meta.json", {"promptSha": sha, "savedAt": saved_at, "source": source, "savedBy": saved_by})


def response_format(prompt_id: str) -> str:
    shape = RESPONSE_SHAPES.get(prompt_id)
    if not shape:
        return ""
    need = [f'"{k}": {{...}}' for k in shape["objects"]] + [f'"{k}": [...]' for k in shape["lists"]]
    return "{" + ", ".join(need) + "}" + (f" (선택: {', '.join(shape['optional'])})" if shape["optional"] else "")


class ServiceClient(ExaoneClient):
    """서버 설정에 따라 live 호출 또는 붙여넣기 응답을 쓴다. 두 방식 모두 응답을 프롬프트 해시와 함께 보관한다."""

    def __init__(self, ws: Path, settings: Settings, not_before: str | None = None):
        super().__init__(ws, "live" if settings.ai_live else "mock", timeout=settings.ai_timeout_seconds, mock_dir=ws / "responses",
                         api_url=settings.ai_api_url, api_key=settings.ai_api_key, model=settings.ai_model,
                         adapter=ChatCompletionsAdapter(json_mode=settings.ai_json_mode, max_tokens=settings.ai_max_tokens),
                         transport=UrlLibTransport(settings.ai_request_id_header or None))
        self.settings = settings
        self.input_chars = settings.ai_input_chars
        self.compat = {"true": True, "false": False}.get(settings.ai_compat_mode)
        self.on_call = None  # 서버 AI 호출 직전 알림 (작업 단계 표시용): on_call(prompt_id, 입력 글자 수)
        self.not_before = not_before
        self.used: list[str] = []

    @property
    def model_label(self) -> str:
        return super().model_label if self.mode == "live" else "붙여넣기 응답 (화면에서 받은 AI 응답)"

    def complete(self, prompt_id, project_id, week, system, user, variant=None):
        self.calls.append(f"{prompt_id}:{project_id}:{week}:{variant or ''}")
        name = response_name(prompt_id, project_id, week, variant)
        prompt = f"{system}\n\n{user}"
        sha = prompt_sha(prompt)
        cached = read_cached(self.mock_dir, name, sha, self.not_before)
        if cached is not None:
            try:
                payload = read_payload(cached, prompt_id)
                self.used.append(f"{PROMPT_LABEL.get(prompt_id, prompt_id)}: 저장된 응답 사용 ({name})")
                return payload
            except (json.JSONDecodeError, ResponseFormatError):
                pass  # 손상된 저장 응답 → 새로 받는다
        if self.mode == "live":
            if self.on_call:
                self.on_call(prompt_id, len(prompt))
            try:
                payload = self._live(prompt_id, system, user)
            except AIError:
                self._save_failed_request(prompt_id, name)
                raise
            save_response(self.mock_dir, name, payload, sha, "live", _now(self.settings))
            stat = self.call_log[-1]  # 입력 크기·걸린 시간·요청 횟수 (어느 단계가 무겁고 느린지 보이게)
            self.used.append(f"{PROMPT_LABEL.get(prompt_id, prompt_id)}: 서버 AI 호출 ({name}) — 입력 {stat['chars']:,}자, "
                             f"{stat['seconds']:g}초, 요청 {stat['sends']}회{', 끝난 이유 ' + str(stat['finish']) if stat['finish'] not in (None, 'stop') else ''}")
            return payload
        if prompt_id == "fit_to_budget":  # 붙여넣기 방식에서는 묻지 않는다 → 원문 유지 + (계속) 장
            raise MockResponseMissing("붙여넣기 방식에서는 분량 줄이기를 AI에 묻지 않음")
        raise NeedResponse(prompt_id, name, prompt, sha, project_id)

    def _save_failed_request(self, prompt_id: str, name: str) -> None:
        if not self.last_request:
            return
        stat = self.call_log[-1] if self.call_log else {}
        try:
            _save_failed_request_file(self.settings, self.last_request, {
                "savedAt": _now(self.settings), "promptId": prompt_id, "responseName": name,
                "inputChars": stat.get("chars"), "sends": stat.get("sends"), "compat": bool(self.compat),
                "requestId": stat.get("request_id")})
        except OSError:
            log.warning("실패한 AI 요청을 저장하지 못함")


def failed_request_path(settings: Settings) -> Path:
    return settings.reports_dir / "ai_debug" / "last_failed_request.json"


def _save_failed_request_file(settings: Settings, body: dict, meta: dict) -> None:
    """[I40] 마지막으로 실패한 AI 요청 본문 (AI점검.bat이 다시 보내 원인 확인). 키·헤더·주소는 담지 않는다."""
    path = failed_request_path(settings)
    path.parent.mkdir(parents=True, exist_ok=True)
    _atomic_json(path, {**meta, "body": body})


# ── 작업 저장소 ──────────────────────────────────────────────────────────────

class JobStore:
    def __init__(self, settings: Settings):
        self.settings = settings
        self.dir = settings.reports_dir / "jobs"

    def _path(self, job_id: str) -> Path:
        if not JOB_ID_RE.match(job_id or ""):
            raise ApiError(404, "NOT_FOUND", "보고자료 작업을 찾을 수 없습니다.")
        return self.dir / f"{job_id}.json"

    def get(self, job_id: str) -> dict | None:
        path = self._path(job_id)
        with _store_lock:
            if not path.is_file():
                return None
            return json.loads(path.read_text(encoding="utf-8"))

    def require(self, job_id: str) -> dict:
        job = self.get(job_id)
        if job is None:
            raise ApiError(404, "NOT_FOUND", "보고자료 작업을 찾을 수 없습니다.")
        return job

    def save(self, job: dict) -> None:
        with _store_lock:
            _atomic_json(self._path(job["id"]), job)

    def update(self, job_id: str, **changes: Any) -> dict:
        with _store_lock:
            job = self.require(job_id)
            job.update(changes)
            job["updatedAt"] = _now(self.settings)
            self.save(job)
            return job

    def all(self) -> list[dict]:
        if not self.dir.is_dir():
            return []
        out = []
        with _store_lock:
            for path in self.dir.glob("*.json"):
                try:
                    out.append(json.loads(path.read_text(encoding="utf-8")))
                except (OSError, json.JSONDecodeError):
                    continue
        return sorted(out, key=lambda j: (j.get("createdAt") or "", j["id"]), reverse=True)

    def claim_next(self) -> dict | None:
        with _store_lock:
            queued = [j for j in self.all() if j.get("status") == "queued"]
            if not queued:
                return None
            job = min(queued, key=lambda j: (j.get("queuedAt") or j.get("createdAt") or "", j["id"]))
            return self.update(job["id"], status="running", startedAt=_now(self.settings), error=None)

    def recover(self) -> int:
        """서버가 작업 도중 꺼졌으면 다시 대기열로 (같은 입력이면 같은 결과가 나오는 작업이라 안전)."""
        n = 0
        for job in self.all():
            if job.get("status") == "running":
                self.update(job["id"], status="queued", stage="서버 재시작 후 다시 대기")
                n += 1
        return n

    def prune(self) -> None:
        keep = self.settings.report_job_retention
        jobs = self.all()
        for job in jobs[keep:]:
            if job.get("status") in ("running", "queued"):
                continue
            with _store_lock:
                self._path(job["id"]).unlink(missing_ok=True)
            shutil.rmtree(self.settings.reports_dir / "output" / job["id"], ignore_errors=True)


def public(job: dict, *, full: bool = False) -> dict:
    out = {k: v for k, v in job.items() if k != "need"}
    need = job.get("need")
    if need:
        out["need"] = need if full else {k: v for k, v in need.items() if k not in ("prompt", "previous")}
    return out


# ── 작업 만들기 ──────────────────────────────────────────────────────────────

def _md(d: date) -> str:
    return f"{d.month}/{d.day}"


def normalize_params(kind: str, template: str, week: str | None, date_from: date | None, date_to: date | None,
                     month: str | None) -> tuple[dict, date, date, str]:
    if kind not in KINDS:
        raise ApiError(422, "VALIDATION_ERROR", "보고 종류가 올바르지 않습니다.")
    if kind == "weekly":
        try:
            start, end = week_range(week or "")
        except ValidationError:
            raise ApiError(422, "VALIDATION_ERROR", "주차 형식은 2026-W40 처럼 입력해야 합니다.") from None
        return {"week": week}, start, end, f"{week} ({_md(start)}~{_md(end)})"
    if kind == "period":
        if not date_from or not date_to:
            raise ApiError(422, "VALIDATION_ERROR", "시작일과 마감일을 모두 정해야 합니다.")
        if date_to < date_from:
            raise ApiError(422, "VALIDATION_ERROR", "마감일은 시작일보다 빠를 수 없습니다.")
        if (date_to - date_from).days + 1 > MAX_PERIOD_DAYS:
            raise ApiError(422, "VALIDATION_ERROR", f"기간은 최대 {MAX_PERIOD_DAYS}일까지 정할 수 있습니다.")
        return ({"dateFrom": date_from.isoformat(), "dateTo": date_to.isoformat()}, date_from, date_to,
                f"{date_from.isoformat()} ~ {date_to.isoformat()}")
    m = re.fullmatch(r"(\d{4})-(\d{2})", month or "")
    if not m or not 1 <= int(m[2]) <= 12:
        raise ApiError(422, "VALIDATION_ERROR", "월 형식은 2026-10 처럼 입력해야 합니다.")
    weeks = month_weeks(int(m[1]), int(m[2]))
    start, end = week_range(weeks[0])[0], week_range(weeks[-1])[1]
    return {"month": month}, start, end, f"{int(m[1])}년 {int(m[2])}월 ({weeks[0][5:]}~{weeks[-1][5:]})"


def create_job(s, settings: Settings, actor, body: Any) -> dict:
    template = body.template if body.kind != "monthly" else "monthly"
    if body.kind != "monthly" and template not in TEMPLATES:
        raise ApiError(422, "VALIDATION_ERROR", "양식이 올바르지 않습니다.")
    params, start, end, label = normalize_params(body.kind, template, body.week, body.date_from, body.date_to, body.month)
    ids = list(dict.fromkeys(body.project_ids))
    if not ids:
        raise ApiError(422, "VALIDATION_ERROR", "프로젝트를 하나 이상 골라야 합니다.")
    if len(ids) > MAX_PROJECTS:
        raise ApiError(422, "VALIDATION_ERROR", f"한 번에 최대 {MAX_PROJECTS}개 프로젝트까지 만들 수 있습니다.")
    names = []
    for pid in ids:
        p = s.get(Project, pid)
        if p is None or p.deleted_at is not None:
            raise ApiError(409, "PROJECT_UNAVAILABLE", "삭제되었거나 없는 프로젝트가 포함되어 있습니다. 목록을 새로고침하세요.")
        names.append(p.name)
    team_summary = body.include_team_summary and body.kind in ("weekly", "period") and template == "weekly"
    now = _now(settings)
    job = {
        "id": uuid.uuid4().hex, "kind": body.kind, "kindLabel": KIND_LABEL[body.kind], "template": template,
        "templateLabel": TEMPLATE_LABEL[template], "params": params, "periodLabel": label,
        "rangeFrom": start.isoformat(), "rangeTo": end.isoformat(),
        "projectIds": ids, "projectNames": names, "orgLabel": (body.org_label or "").strip() or None,
        "options": {"includeTables": body.include_tables, "includeGantts": body.include_gantts,
                    "includeMilestoneGantt": body.include_milestone_gantt, "refreshAi": body.refresh_ai,
                    "includeTeamSummary": team_summary,
                    "summaryAuthor": ((body.summary_author or "").strip() or None) if team_summary else None},
        "aiMode": "live" if settings.ai_live else "paste",
        "status": "queued", "stage": "대기 중", "need": None, "result": None, "error": None, "responsesReceived": 0,
        "requestedBy": {"id": actor.id, "name": actor.name}, "createdAt": now, "queuedAt": now, "updatedAt": now,
        "startedAt": None, "finishedAt": None,
    }
    store = JobStore(settings)
    store.save(job)
    store.prune()
    return public(job)


def submit_response(settings: Settings, job_id: str, name: str, text: str, actor) -> dict:
    store = JobStore(settings)
    with _store_lock:
        job = store.require(job_id)
        need = job.get("need") or {}
        if job.get("status") != "need_response" or need.get("responseName") != name:
            raise ApiError(409, "RESPONSE_NOT_EXPECTED", "지금 이 응답을 기다리는 작업이 아닙니다. 화면을 새로고침하세요.")
        prompt_id = need["promptId"]
        try:
            payload = read_payload(text, prompt_id)
        except json.JSONDecodeError as exc:
            raise ApiError(422, "INVALID_AI_RESPONSE", f"JSON 형식이 아닙니다 (줄 {exc.lineno}, 칸 {exc.colno} 근처). AI 답에서 {{ 로 시작하는 JSON 부분을 "
                                                       "빠짐없이 붙여 넣었는지 확인하세요. 저장하지 않았습니다.") from None
        except ResponseFormatError as exc:
            raise ApiError(422, "INVALID_AI_RESPONSE", f"{exc}. 필요한 형식: {response_format(prompt_id)} — 저장하지 않았습니다.") from None
        if not NAME_RE.match(name):
            raise ApiError(422, "VALIDATION_ERROR", "응답 이름이 올바르지 않습니다.")
        save_response(responses_dir(settings), name, payload, need["promptSha"], "paste", _now(settings), actor.id if actor else None)
        now = _now(settings)
        return public(store.update(job_id, status="queued", queuedAt=now, stage="응답 받음 · 이어서 진행 대기",
                                   responsesReceived=int(job.get("responsesReceived") or 0) + 1))


def retry_job(settings: Settings, job_id: str) -> dict:
    store = JobStore(settings)
    with _store_lock:
        job = store.require(job_id)
        if job["status"] in ("queued", "running"):
            raise ApiError(409, "JOB_BUSY", "이미 진행 중인 작업입니다.")
        now = _now(settings)
        return public(store.update(job_id, status="queued", queuedAt=now, stage="다시 실행 대기", error=None))


def cancel_job(settings: Settings, job_id: str) -> dict:
    store = JobStore(settings)
    with _store_lock:
        job = store.require(job_id)
        if job["status"] not in ("queued", "need_response", "failed"):
            raise ApiError(409, "JOB_NOT_CANCELLABLE", "진행 중이거나 끝난 작업은 취소할 수 없습니다.")
        return public(store.update(job_id, status="cancelled", stage="취소됨", finishedAt=_now(settings)))


def job_file(settings: Settings, job_id: str, name: str) -> Path:
    job = JobStore(settings).require(job_id)
    files = {f["name"] for f in ((job.get("result") or {}).get("files") or [])}
    if name not in files:
        raise ApiError(404, "NOT_FOUND", "파일을 찾을 수 없습니다.")
    path = settings.reports_dir / "output" / job_id / name
    if not path.is_file():
        raise ApiError(410, "FILE_GONE", "보관 기간이 지나 파일이 지워졌습니다. 다시 실행하세요.")
    return path


def config_info(settings: Settings) -> dict:
    return {"aiMode": "live" if settings.ai_live else "paste", "aiModel": settings.ai_model if settings.ai_live else None,
            "aiPasteReason": settings.ai_paste_reason,
            "aiUrlConfigured": bool(settings.ai_api_url), "aiKeyConfigured": bool(settings.ai_api_key),
            "fonts": fonts_available(), "kinds": [{"id": k, "label": KIND_LABEL[k]} for k in KINDS],
            "templates": [{"id": t, "label": TEMPLATE_LABEL[t]} for t in TEMPLATES], "maxProjects": MAX_PROJECTS}


# ── 입력 준비 ────────────────────────────────────────────────────────────────

def write_exports(rt, ws: Path, project_ids: list[str], d_from: date, d_to: date) -> tuple[dict[str, str], list[str]]:
    """선택 프로젝트의 기준정보 + 기간 안 업무일지를 WorkLog export 형식으로 작업공간에 쓴다 (매 실행마다 새로)."""
    st = rt.settings
    target = ws / "data" / "worklog"
    shutil.rmtree(target, ignore_errors=True)
    target.mkdir(parents=True, exist_ok=True)
    names: dict[str, str] = {}
    notes: list[str] = []
    as_of = min(_today(st), d_to)
    stamp = f"{as_of.isoformat()}T23:59:59+09:00" if as_of < _today(st) else None
    with rt.read_factory() as s:
        for pid in project_ids:
            p = s.get(Project, pid)
            if p is None or p.deleted_at is not None:
                raise ReportError("삭제되었거나 없는 프로젝트가 포함되어 있습니다.")
            names[pid] = p.name
            for _rel, body in export_svc.build_project(s, st, pid, p.revision):
                proj = body["project"]
                if not proj.get("startDate") and not proj.get("endDate"):
                    notes.append(f"{p.name}: 프로젝트 시작·종료일이 비어 있어 일정 칸을 '미정'으로 표시")
                if proj.get("team") is None:
                    notes.append(f"{p.name}: 담당 팀이 없어 조직 칸을 '-'로 표시")
                if stamp:  # 지난 기간 보고: 마일스톤 지연 판정 기준일 = 보고 기간 마지막 날
                    body["generatedAt"] = stamp
                _atomic_json(target / pid / "project.json", body)
            days = s.execute(select(DailyLog.work_date).where(
                DailyLog.project_id == pid, DailyLog.deleted_at.is_(None), DailyLog.work_date >= d_from,
                DailyLog.work_date <= d_to).distinct()).scalars().all()
            for day in sorted(days):
                for _rel, body in export_svc.build_daily(s, st, pid, day, 0):
                    if stamp:
                        body["generatedAt"] = stamp
                    _atomic_json(target / pid / f"{day.isoformat()}.json", body)
    return names, notes


def appendix_blocks(rt, pid: str, d_from: date, d_to: date, options: dict, marker: date) -> tuple[list[Block], list[str]]:
    blocks: list[Block] = []
    notes: list[str] = []
    with rt.read_factory() as s:
        if options.get("includeMilestoneGantt"):
            rows = []
            ms = s.execute(select(Milestone).where(Milestone.project_id == pid, Milestone.deleted_at.is_(None),
                                                   Milestone.is_general.is_(False)).order_by(Milestone.sort_order, Milestone.id)).scalars().all()
            for m in ms:
                status = MS_STATUS_KO.get(m.status, m.status)
                if status not in ("완료", "취소") and m.planned_end and m.planned_end < marker:
                    status = "지연"
                start = m.planned_start or m.planned_end
                end = m.planned_end or m.planned_start
                rows.append(GanttRow(m.name, start, end, STATUS_COLOR.get(status, "4472C4"),
                                     baseline=(m.baseline_start or m.baseline_end, m.baseline_end or m.baseline_start)
                                     if (m.baseline_start or m.baseline_end) else None,
                                     actual_end=m.actual_end, note=status + (f" · 실적 {_md(m.actual_end)}" if m.actual_end else "")))
            if rows:
                blocks.append(GanttBlock("마일스톤 일정", f"프로젝트 마일스톤 (생성 시점 값) · 기준일 {marker.isoformat()}", rows,
                                         marker=marker, legend="막대: 계획 기간(색 = 상태: 진행 파랑·지연 빨강·완료 회색·예정 하늘) · "
                                                               "아래 가는 막대: Baseline · ◆: 실적 완료일 · 빨간 점선: 보고 기준일"))
            else:
                notes.append("마일스톤 일정: 일반·수시 업무 외 마일스톤이 없어 넣지 않음")
        if options.get("includeTables") or options.get("includeGantts"):
            logs = s.execute(select(DailyLog).where(DailyLog.project_id == pid, DailyLog.deleted_at.is_(None),
                                                    DailyLog.work_date >= d_from, DailyLog.work_date <= d_to)
                             .order_by(DailyLog.work_date, DailyLog.created_at, DailyLog.id)).scalars().all()
            found: list[Block] = []
            for lg in logs:
                author = (lg.author_snapshot or {}).get("name") or "작성자 미상"
                for task in lg.tasks:
                    name = (task.title or "").strip() or title_preview(task.content_doc or {}, 40) or "TASK"
                    subtitle = f"{lg.work_date.month}/{lg.work_date.day} {author} · {name}"
                    for block in blocks_from_document(task.content_doc, subtitle):
                        if isinstance(block, TableBlock) and options.get("includeTables"):
                            found.append(block)
                        elif isinstance(block, GanttBlock) and options.get("includeGantts"):
                            block.marker = None
                            found.append(block)
            if len(found) > MAX_LOG_BLOCKS:
                notes.append(f"업무일지 표·간트 {len(found)}개 중 앞의 {MAX_LOG_BLOCKS}개만 넣음 (기간을 줄이면 모두 넣을 수 있음)")
                found = found[:MAX_LOG_BLOCKS]
            if not found:
                notes.append("기간 안 업무일지에 넣을 표·간트가 없음")
            blocks += found
    return blocks, notes


# ── 실행 ─────────────────────────────────────────────────────────────────────

def render_weekly_deck(template: Path, prepared: list[dict], output: Path) -> list[str]:
    """여러 과제의 주간업무 슬라이드를 한 파일로 (팀장 요약 없이). render_team_report의 하위 호환 래퍼."""
    return render_team_report(template, [TeamSection("", prepared)], output).notes


class _Run:
    def __init__(self, rt, job: dict):
        self.rt, self.job, self.st = rt, job, rt.settings
        self.store = JobStore(self.st)
        self.ws = sync_workspace(self.st)
        not_before = job.get("createdAt") if (job.get("options") or {}).get("refreshAi") else None
        self.client = ServiceClient(self.ws, self.st, not_before)
        self.client.on_call = self._on_call
        self.out = self.st.reports_dir / "output" / job["id"]
        self.check: list[str] = []
        self.problems: list[str] = []
        self.names: dict[str, str] = {}

    def stage(self, text: str) -> None:
        self.current = text
        self.store.update(self.job["id"], stage=text)

    def _on_call(self, prompt_id: str, chars: int) -> None:
        base = getattr(self, "current", "").split(" (서버 AI", 1)[0]
        self.stage(f"{base} (서버 AI {PROMPT_LABEL.get(prompt_id, prompt_id)} 응답 대기 · 입력 {chars:,}자)")

    def section(self, title: str, lines: list[str]) -> None:
        self.check += ["", f"===== {title} ====="] + (lines or ["(없음)"])

    def run(self) -> dict:
        job = self.job
        d_from, d_to = date.fromisoformat(job["rangeFrom"]), date.fromisoformat(job["rangeTo"])
        ids = job["projectIds"]
        self.stage("입력 준비 (업무일지·기준정보 읽기)")
        self.names, notes = write_exports(self.rt, self.ws, ids, d_from, d_to)
        shutil.rmtree(self.out, ignore_errors=True)
        self.out.mkdir(parents=True, exist_ok=True)
        if job["kind"] == "weekly":
            pptx, slide_groups = self._weekly(job["params"]["week"], ids)
        elif job["kind"] == "period":
            pptx, slide_groups = self._period(date.fromisoformat(job["params"]["dateFrom"]), date.fromisoformat(job["params"]["dateTo"]), ids)
        else:
            pptx, slide_groups = self._monthly(job["params"]["month"], ids)
        appendix_notes = self._appendix(pptx, slide_groups, d_from, d_to)
        head = [f"보고 종류: {job['kindLabel']} / 양식: {job['templateLabel']}", f"기간: {job['periodLabel']}",
                f"프로젝트 {len(ids)}개: {', '.join(self.names[i] for i in ids)}", f"AI: {self.client.model_label}",
                f"생성 시각: {_now(self.st)} / 요청: {(job.get('requestedBy') or {}).get('name')}"]
        head += [f"- {n}" for n in notes]
        self.section("AI 응답 사용 내역", self.client.used)
        if self.client.notes:
            self.section("AI 호출 중 자동 조치 (재요청 등)", self.client.notes)
        self.section("참고 슬라이드", appendix_notes)
        self.section("PPT 재검사 요약", self.problems or ["통과"])
        check_name = f"{pptx.stem}_검사보고서.txt"
        _atomic_text(self.out / check_name, "\n".join(head + self.check) + "\n")
        from pptx import Presentation
        slide_count = len(Presentation(str(pptx)).slides)
        files = [{"name": pptx.name, "label": "PPT", "sizeBytes": pptx.stat().st_size},
                 {"name": check_name, "label": "검사 보고서", "sizeBytes": (self.out / check_name).stat().st_size}]
        return {"files": files, "slideCount": slide_count, "problems": self.problems[:50]}

    def _stem(self) -> str:
        job = self.job
        if job["kind"] == "weekly":
            stem = f"주간보고_{job['params']['week']}"
        elif job["kind"] == "period":
            stem = f"기간보고_{job['params']['dateFrom']}_{job['params']['dateTo']}"
        else:
            stem = f"월간종합_{job['params']['month']}"
        return stem + ("_경영진요약" if job["template"] == "exec" else "")

    def _rollup(self, pid: str, week: str, out_root: Path, label: str, **kw) -> tuple[Path, Path]:
        self.stage(f"{self.names[pid]} · {label}")
        wpath, cpath, report = run_weekly(self.ws, pid, week, out_root, client=self.client, **kw)
        mode = "서버 AI" if self.client.mode == "live" else "붙여넣기 응답"
        lines = [f"AI 모드: {mode}{line[len('AI 모드: mock'):]}" if line.startswith("AI 모드: mock") else line
                 for line in report.read_text(encoding="utf-8").splitlines()]
        self.section(f"{self.names[pid]} · {label} 의미 검증", lines)
        return wpath, cpath

    def _weekly_template(self, items: list[tuple[str, Path, Path]], *, out_root: Path | None = None,
                         period: tuple[date, date] | None = None, key: str | None = None) -> tuple[Path, list[tuple[str, int]]]:
        """주간업무 양식 한 파일. 팀장 요약을 켜면 팀별로 [팀장 요약 페이지 → 그 팀 과제 장표] 순서로 묶는다."""
        template = find_template(self.ws)
        with_summary = bool((self.job.get("options") or {}).get("includeTeamSummary"))
        projects = {pid: wr_sources.load_project(self.ws, pid) for pid, _w, _c in items}
        grouped: list[tuple[str, list[tuple[str, Path, Path]]]] = [("", items)]
        if with_summary:  # 팀 순서 = 선택 목록에서 처음 나온 순서, 팀 안 순서 = 선택 순서
            teams: dict[str, list[tuple[str, Path, Path]]] = {}
            for item in items:
                teams.setdefault((projects[item[0]].get("org") or {}).get("team") or "팀 미지정", []).append(item)
            grouped = list(teams.items())
            sgeom = summary_geometry(self.ws, open_deck_template(template)[1])
        sections: list[TeamSection] = []
        order: list[str] = []
        part_no = 0
        for team, team_items in grouped:
            prepared = []
            for i, (pid, wpath, cpath) in enumerate(team_items, 1):  # 과제 번호 (n/N): 요약을 켜면 팀 안에서
                self.stage(f"{self.names[pid]} · PPT 구성")
                project = projects[pid]
                p = prepare_ppt(self.ws, project, wpath, cpath, template, client=self.client, project_index=i, project_total=len(team_items))
                part_no += 1
                part = self.out / "parts" / f"{part_no:02d}.pptx"
                notes = generate_ppt(self.ws, project, wpath, cpath, template, part, client=self.client, prepared=p)
                self.problems += [f"{self.names[pid]}: {n[len('PPT 검사 문제: '):]}" for n in notes if n.startswith("PPT 검사 문제")]
                self.section(f"{self.names[pid]} · PPT 처리", part.with_name(f"{part.stem}_ppt_check.txt").read_text(encoding="utf-8").splitlines())
                prepared.append(p)
                order.append(pid)
            section = TeamSection(team, prepared)
            if with_summary:
                self._team_summary(section, team_items, projects, sgeom, out_root or self.ws, period, key)
            sections.append(section)
        pptx = self.out / f"{self._stem()}.pptx"
        layout = render_team_report(template, sections, pptx)
        shutil.rmtree(self.out / "parts", ignore_errors=True)
        if with_summary:
            from pptx import Presentation
            saved = Presentation(str(pptx))
            for section, start in zip(sections, layout.summary_starts):
                found = inspect_summary(saved, start, section.pages, section.geom, section.titles)
                self.problems += [f"팀장 요약 · {section.team}: {f}" for f in found]
                self.section(f"팀장 요약 · {section.team} PPT 재검사", [f"문제: {f}" for f in found] or ["통과 (제목·글꼴·글자 크기·색·파랑=대상 기간 근거·영역 경계)"])
        return pptx, list(zip(order, layout.groups))

    def _team_summary(self, section: "TeamSection", team_items: list[tuple[str, Path, Path]], projects: dict[str, dict],
                      sgeom: "SummaryGeometry", out_root: Path, period: tuple[date, date] | None, key: str | None) -> None:
        """팀 과제마다 AI 요약(project_summary)을 받아 요약 페이지를 배치한다. 붙여넣기 모드면 여기서 'AI 응답 필요'로 멈춘다."""
        budget = lines_per_project(sgeom, len(team_items))
        line_chars = sgeom.chars(FONT_SIZES[0], ITEM_MAR_IN)
        blocks, stamps = [], []
        for number, (pid, wpath, _c) in enumerate(team_items, 1):
            self.stage(f"{self.names[pid]} · 팀장 요약")
            weekly = load_json(wpath)
            stamps.append(weekly["meta"]["updated_at"])
            result = run_summary(self.ws, projects[pid], weekly["week"], out_root, client=self.client, max_lines=budget,
                                 line_chars=line_chars, period=period, response_key=key)
            mode = "서버 AI" if self.client.mode == "live" else "붙여넣기 응답"
            lines = [f"AI 모드: {mode}{line[len('AI 모드: mock'):]}" if line.startswith("AI 모드: mock") else line
                     for line in result.check.read_text(encoding="utf-8").splitlines()]
            self.section(f"{self.names[pid]} · 팀장 요약 의미 검증", lines)
            blocks.append(SummaryBlock(number, result.summary["project_name"], result.summary["items"]))
        pages, notes = layout_summary(blocks, sgeom)
        first = projects[team_items[0][0]]
        author = (self.job.get("options") or {}).get("summaryAuthor")
        section.author = f"작성자 : {author or PeopleTable.load(self.ws).with_project(first).name_with_title(first['owner'])}"
        section.updated_at = updated_at_label(max(stamps))
        section.pages, section.geom, section.titles = pages, sgeom, summary_titles(section.team, len(pages))
        self.section(f"팀장 요약 · {section.team}", [f"과제 {len(team_items)}건, 과제당 {budget}줄 배정 (한 줄 약 {int(line_chars)}자)",
                                                    section.author + (" (입력값)" if author else " (입력 없음 → 첫 과제 담당자)")] + notes)

    def _exec_template(self, fills: list[tuple[str, Any]]) -> tuple[Path, list[tuple[str, int]]]:
        pptx = self.out / f"{self._stem()}.pptx"
        notes = render_report(find_report_template(self.ws), [f for _pid, f in fills], pptx)
        problems = inspect_report(pptx, self.ws)
        self.problems += problems
        self.section("경영진 1장 요약 PPT 처리", notes + [f"문제: {p}" for p in problems])
        return pptx, [(pid, i) for i, (pid, _f) in enumerate(fills, 1)]

    def _exec_fill(self, pid: str, week: str, **kw) -> Any:
        self.stage(f"{self.names[pid]} · 경영진 1장 요약")
        built = exec_fill(self.ws, pid, week, client=self.client, today=_today(self.st), **kw)
        self.section(f"{self.names[pid]} · 경영진 1장 요약 AI 응답 처리", built["notes"] + [built["schedule"].lstrip("└ ")])
        return built["fill"]

    def _weekly(self, week: str, ids: list[str]):
        items = []
        for pid in ids:
            wpath, cpath = self._rollup(pid, week, self.ws, f"주간 정리 {week}")
            items.append((pid, wpath, cpath))
        if self.job["template"] == "weekly":
            return self._weekly_template(items)
        return self._exec_template([(pid, self._exec_fill(pid, week, dirs=[self.ws])) for pid in ids])

    def _period(self, start: date, end: date, ids: list[str]):
        key = f"{start.isoformat()}_{end.isoformat()}"
        week = iso_week(end)
        pdir = self.ws / "periods" / key
        items = []
        for pid in ids:
            wpath, cpath = self._rollup(pid, week, pdir, f"기간 정리 {start}~{end}", period=(start, end), response_key=key)
            items.append((pid, wpath, cpath))
        if self.job["template"] == "weekly":
            return self._weekly_template(items, out_root=pdir, period=(start, end), key=key)
        fills = []
        for pid, wpath, _c in items:
            block = f"■ {pid} 기간 {start.isoformat()} ~ {end.isoformat()}\n" + pv.prev_weekly_lines(load_json(wpath))
            fills.append((pid, self._exec_fill(pid, week, dirs=[pdir, self.ws], response_key=key, weekly_blocks=block)))
        return self._exec_template(fills)

    def _monthly(self, month: str, ids: list[str]):
        year, mon = (int(v) for v in month.split("-"))
        for pid in ids:
            for week in month_weeks(year, mon):
                self._rollup(pid, week, self.ws, f"주간 정리 {week}")
        self.stage("월간 종합 PPT 구성")
        pptx = self.out / f"{self._stem()}.pptx"
        scope = "S" + hashlib.sha1(",".join(sorted(ids)).encode()).hexdigest()[:10]
        result = generate_monthly(self.ws, ids, year, mon, pptx, client=self.client, dirs=[self.ws], scope=scope,
                                  today=_today(self.st), org=self.job.get("orgLabel"))
        self.problems += result["problems"]
        self.section("월간 종합 PPT 처리", result["check"].read_text(encoding="utf-8").splitlines())
        result["check"].unlink(missing_ok=True)
        from pptx import Presentation
        total = len(Presentation(str(pptx)).slides)
        return pptx, [(pid, total) for pid in ids]

    def _appendix(self, pptx: Path, groups: list[tuple[str, int]], d_from: date, d_to: date) -> list[str]:
        options = self.job.get("options") or {}
        if not any(options.get(k) for k in ("includeTables", "includeGantts", "includeMilestoneGantt")):
            return ["넣지 않음 (옵션 선택 안 함)"]
        self.stage("참고 슬라이드 (표·간트)")
        marker = min(_today(self.st), d_to)
        notes: list[str] = []
        built = []
        for pid, after in groups:
            blocks, local = appendix_blocks(self.rt, pid, d_from, d_to, options, marker)
            notes += [f"{self.names[pid]}: {n}" for n in local]
            built.append(AppendixGroup(after, self.names[pid], blocks))
        added = insert_appendix(pptx, built)
        notes.append(f"참고 슬라이드 {added}장 추가")
        return notes


def execute(rt, job: dict) -> dict:
    return _Run(rt, job).run()


def run_next(rt) -> str | None:
    """대기 작업 하나를 처리한다 (테스트에서 직접 호출). 작업공간이 하나라 한 번에 하나씩."""
    if rt.maintenance:
        return None
    with _run_lock:
        store = JobStore(rt.settings)
        job = store.claim_next()
        if job is None:
            return None
        jid = job["id"]
        try:
            result = execute(rt, job)
            store.update(jid, status="succeeded", stage="완료", result=result, need=None, error=None, finishedAt=_now(rt.settings))
        except NeedResponse as nr:
            current = store.require(jid)
            saved = responses_dir(rt.settings) / nr.name
            store.update(jid, status="need_response", stage=current.get("stage") or "", need={
                "promptId": nr.prompt_id, "promptLabel": PROMPT_LABEL.get(nr.prompt_id, nr.prompt_id), "responseName": nr.name,
                "projectId": nr.project_id, "projectName": dict(zip(job["projectIds"], job["projectNames"])).get(nr.project_id),
                "prompt": nr.prompt, "promptSha": nr.sha, "format": response_format(nr.prompt_id),
                "previous": saved.read_text(encoding="utf-8") if saved.is_file() else "",
            })
        except PermissionError:
            store.update(jid, status="failed", stage="실패", error="결과 파일이 다른 프로그램에서 열려 있어 쓸 수 없습니다. 닫고 다시 실행하세요.",
                         finishedAt=_now(rt.settings))
        except (ReportError, ValidationError, WorklogError, BudgetError, AIError, FileNotFoundError, ValueError) as exc:
            log.warning("report job %s failed: %s", jid, exc)
            where = (store.require(jid).get("stage") or "").split(" (서버 AI", 1)[0]
            error = f"[{where}] {exc}" if isinstance(exc, AIError) and where else str(exc)  # 어느 과제·단계에서 실패했는지
            store.update(jid, status="failed", stage="실패", error=error, finishedAt=_now(rt.settings))
        except Exception as exc:  # noqa: BLE001 - 작업 스레드는 계속 돈다
            log.exception("report job %s crashed", jid)
            store.update(jid, status="failed", stage="실패", error=f"내부 오류 ({type(exc).__name__}). 서버 로그를 확인하세요.",
                         finishedAt=_now(rt.settings))
        return jid


def _loop(rt) -> None:
    JobStore(rt.settings).recover()
    while not rt.stop_event.is_set():
        try:
            while run_next(rt) and not rt.stop_event.is_set():
                pass
        except Exception:  # noqa: BLE001
            log.exception("report worker cycle failed")
        rt.stop_event.wait(1.0)


def start_report_worker(rt) -> threading.Thread:
    t = threading.Thread(target=_loop, args=(rt,), name="worklog-reports", daemon=True)
    t.start()
    return t
