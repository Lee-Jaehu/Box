"""보고 자료 템플릿 초안(보고자료_Template_v1_초안.pptx)에 내용을 채운다.

- 칸은 도형 이름으로 찾고, 템플릿 칸의 글자 서식(크기·굵기)을 복제한다 (주간 PPT와 같은 Proto/write_paras).
- 표는 템플릿 데이터 행을 복제해 필요한 행 수만큼 만들고, 글자 줄 수로 행 높이를 정한다.
- 과제 현황표가 한 장에 넘치면 같은 슬라이드를 복제해 "(계속)" 장으로 나눈다.
"""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from ..core import ValidationError
from ..fonts import load_fonts
from ..ppt.model import Para, Run
from ..ppt.render import Proto, _set_cell_fill, a, duplicate_slide, fix_theme_fonts, shape_map, write_paras
from ..textmetrics import line_count
from .content import Cell

TEMPLATE_NAME = "보고자료_Template_v1_초안.pptx"
SLIDE_NAMES = ("exec_summary", "monthly_overview", "monthly_requests")
EMU_PER_IN, EMU_PER_PT = 914400, 12700
LINE_FACTOR = 1.2  # 줄 높이 / 글자 크기 (보수적)
OVERVIEW_BOTTOM_IN = 4.82  # 과제 현황표가 넘으면 안 되는 선 (주요 성과 제목 위)
REQUESTS_BOTTOM_IN = 7.0


@dataclass
class SlideFill:
    """슬라이드 1장에 채울 값. texts: 칸 이름 → 문단, tables: 표 이름 → 행(Cell 목록)."""
    source: str  # 템플릿 슬라이드 이름
    texts: dict[str, list[Para]] = field(default_factory=dict)
    tables: dict[str, list[list[Cell]]] = field(default_factory=dict)


def find_report_template(root: Path) -> Path:
    for path in (root / TEMPLATE_NAME, root / "templates" / TEMPLATE_NAME):
        if path.exists():
            return path
    raise FileNotFoundError(f"보고 자료 템플릿 없음: {TEMPLATE_NAME} (python tools/make_report_template.py 로 생성)")


def open_report_template(path: Path):
    prs = Presentation(str(path))
    names = tuple(s.name for s in prs.slides)
    if names != SLIDE_NAMES:
        raise ValidationError(f"보고 자료 템플릿 슬라이드 구성 {names} ≠ {SLIDE_NAMES}")
    return prs


# ---------------------------------------------------------------- 글자 분량

def _font_pt(tx_body) -> float:
    proto = Proto(tx_body)
    size = proto.normal.get("sz") or proto.bold.get("sz") or "1000"
    return int(size) / 100


def _width_chars(width_emu: int, margin_emu: int, size_pt: float, hangul_em: float) -> float:
    return max(1.0, (width_emu - margin_emu) / EMU_PER_PT / (size_pt * hangul_em))


def text_lines(paras: list[Para], width_chars: float) -> int:
    return sum(line_count(p.text or " ", width_chars) for p in paras)


def needed_height(lines: int, size_pt: float, padding_emu: int) -> int:
    return int(lines * size_pt * LINE_FACTOR * EMU_PER_PT + padding_emu)


# ---------------------------------------------------------------- 채우기

def cell_paras(cell: Cell) -> list[Para]:
    return [Para([Run(line, bold=cell.bold, color=cell.color)]) for line in cell.text.split("\n")]


def fill_table(shape, rows: list[list[Cell]], hangul_em: float) -> int:
    """데이터 행을 rows 수만큼 복제해 채우고, 글자 줄 수로 행 높이를 정한다. 표 높이(EMU)를 돌려준다."""
    tbl = shape.table._tbl
    trs = tbl.findall(a("tr"))
    prototype = deepcopy(trs[1])
    widths = [int(gc.get("w")) for gc in tbl.find(a("tblGrid")).findall(a("gridCol"))]
    protos = [Proto(tc.find(a("txBody"))) for tc in prototype.findall(a("tc"))]
    sizes = [_font_pt(tc.find(a("txBody"))) for tc in prototype.findall(a("tc"))]
    margin = int(Inches(0.1))
    for tr in trs[1:]:
        tbl.remove(tr)
    min_height = int(prototype.get("h"))
    total = int(trs[0].get("h"))
    for row in rows:
        tr = deepcopy(prototype)
        height = min_height
        for col, tc in enumerate(tr.findall(a("tc"))):
            paras = cell_paras(row[col])
            write_paras(tc.find(a("txBody")), paras, recolor=False, proto=protos[col])
            if row[col].fill:
                _set_cell_fill(tc, row[col].fill)
            lines = text_lines(paras, _width_chars(widths[col], margin, sizes[col], hangul_em))
            height = max(height, needed_height(lines, sizes[col], int(Inches(0.06))))
        tr.set("h", str(height))
        total += height
        tbl.append(tr)
    shape.height = total
    return total


def table_heights(shape, rows: list[list[Cell]], hangul_em: float) -> list[int]:
    """fill_table과 같은 계산으로 행 높이만 구한다 (페이지 나누기용)."""
    tbl = shape.table._tbl
    tr = tbl.findall(a("tr"))[1]
    widths = [int(gc.get("w")) for gc in tbl.find(a("tblGrid")).findall(a("gridCol"))]
    sizes = [_font_pt(tc.find(a("txBody"))) for tc in tr.findall(a("tc"))]
    margin, min_height = int(Inches(0.1)), int(tr.get("h"))
    return [max([min_height] + [needed_height(text_lines(cell_paras(c), _width_chars(widths[i], margin, sizes[i], hangul_em)),
                                              sizes[i], int(Inches(0.06))) for i, c in enumerate(row)]) for row in rows]


def paginate_rows(prs, rows: list[list[Cell]], hangul_em: float, bottom_in: float, maximum: int = 8) -> list[list[list[Cell]]]:
    shape = shape_map(_slide(prs, "monthly_overview"))["project_table"]
    room = int(bottom_in * EMU_PER_IN) - shape.top - int(shape.table._tbl.findall(a("tr"))[0].get("h"))
    pages, current, used = [], [], 0
    for row, height in zip(rows, table_heights(shape, rows, hangul_em)):
        if current and (used + height > room or len(current) >= maximum):
            pages.append(current)
            current, used = [], 0
        current.append(row)
        used += height
    return pages + [current] if current or not pages else pages


def _slide(prs, name: str):
    return next(s for s in prs.slides if s.name == name)


def _drop_slide(prs, slide) -> None:
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if prs.part.related_part(sld_id.rId) is slide.part:
            prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)
            return


def _copy_slide(prs, source):
    slide = duplicate_slide(prs, source)
    slide.name = source.name
    if source._element.get("showMasterSp") is not None:
        slide._element.set("showMasterSp", source._element.get("showMasterSp"))
    return slide


def render_report(template: Path, fills: list[SlideFill], output: Path) -> list[str]:
    """fills 순서대로 슬라이드를 만든다. 쓰지 않는 템플릿 슬라이드는 지운다."""
    prs = open_report_template(template)
    notes = fix_theme_fonts(prs)
    hangul_em = load_fonts(template.parent.resolve()).hangul_em
    originals = {name: _slide(prs, name) for name in SLIDE_NAMES}
    used: dict[str, int] = {}
    slides = []
    for fill in fills:
        source = originals[fill.source]
        slide = source if not used.get(fill.source) else _copy_slide(prs, source)
        used[fill.source] = used.get(fill.source, 0) + 1
        slides.append((slide, fill))
    # 복제는 템플릿 원본 상태에서 해야 하므로, 모두 만든 뒤에 채운다
    for slide, fill in slides:
        shapes = shape_map(slide)
        for name, paras in fill.texts.items():
            write_paras(shapes[name].text_frame._txBody, paras, recolor=False)
        for name, rows in fill.tables.items():
            fill_table(shapes[name], rows, hangul_em)
    for name, slide in originals.items():
        if not used.get(name):
            _drop_slide(prs, slide)
    # 슬라이드 순서 = fills 순서
    order = [s.part for s, _ in slides]
    sld_ids = prs.slides._sldIdLst
    ids = {prs.part.related_part(sid.rId): sid for sid in sld_ids}
    for sid in list(sld_ids):
        sld_ids.remove(sid)
    for part in order:
        sld_ids.append(ids[part])
    total = len(order)
    for index, (slide, _fill) in enumerate(slides, 1):
        write_paras(shape_map(slide)["page_no"].text_frame._txBody, [Para([Run(f"{index} / {total}", bold=True)])], recolor=False)
    output.parent.mkdir(parents=True, exist_ok=True)
    tmp = output.with_name(f".{output.name}.tmp")
    prs.save(str(tmp))
    tmp.replace(output)
    return notes


# ---------------------------------------------------------------- 재검사

def inspect_report(path: Path, root: Path) -> list[str]:
    """글꼴(ea), 도형이 슬라이드 안에 있는지, 글 상자에 글이 들어가는지(예상 줄 수)."""
    from ..fonts import EA_REGULAR

    prs = Presentation(str(path))
    hangul_em = load_fonts(root.resolve()).hangul_em
    problems = []
    width, height = prs.slide_width, prs.slide_height
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width + 1 or shape.top + shape.height > height + 1:
                problems.append(f"{number}장 {shape.name}: 슬라이드 밖으로 나감")
            bodies = [shape.text_frame._txBody] if shape.has_text_frame else []
            if shape.has_table:
                bodies = [tc.find(a("txBody")) for tc in shape.table._tbl.iter(a("tc"))]
            for body in bodies:
                for rpr in body.iter(a("rPr")):
                    ea = rpr.find(a("ea"))
                    if ea is None or ea.get("typeface") != EA_REGULAR:
                        problems.append(f"{number}장 {shape.name}: ea 글꼴이 {EA_REGULAR}가 아님")
                        break
            if shape.has_text_frame and shape.text_frame.text.strip():
                tf = shape.text_frame
                size = _font_pt(tf._txBody)
                margin = (tf.margin_left or 0) + (tf.margin_right or 0)
                paras = [Para([Run(p.text)]) for p in tf.paragraphs]
                lines = text_lines(paras, _width_chars(shape.width, margin, size, hangul_em))
                need = needed_height(lines, size, (tf.margin_top or 0) + (tf.margin_bottom or 0))
                if need > shape.height * 1.02:
                    problems.append(f"{number}장 {shape.name}: 글이 칸을 넘칠 수 있음 (예상 {lines}줄, {need / EMU_PER_IN:.2f}in > {shape.height / EMU_PER_IN:.2f}in)")
    return problems
