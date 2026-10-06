"""보고자료 팀장 요약 페이지 (서비스 수준): DB → 정리(AI 붙여넣기) → 팀별 [요약 → 과제 장표] PPT.

AI 응답은 붙여넣기 흐름(need_response → /response → 다시 실행)으로 넣는다. 응답 문장은 테스트가 프롬프트의 기록 ID만 써서 만든다.
"""
import json
import re
import uuid

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.config import Settings
from app.main import create_app
from app.services.reports import JobStore, job_file, run_next

WEEK = "2026-W40"
WORK_DAYS = ("2026-09-29", "2026-09-30")


def doc(text: str) -> dict:
    return {"documentVersion": 1, "format": "tiptap-json",
            "doc": {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": text}]}]}}


class Api:
    def __init__(self, client: TestClient):
        self.c, self.actor = client, None

    def post(self, path: str, body: dict, *, actor: bool = True) -> dict:
        headers = {"Idempotency-Key": uuid.uuid4().hex}
        if actor and self.actor:
            headers["X-Actor-Id"] = self.actor
        r = self.c.post(f"/api/v1{path}", json=body, headers=headers)
        assert r.status_code in (200, 201, 202), r.text
        return r.json()["data"]

    def put(self, path: str, body: dict) -> dict:
        r = self.c.put(f"/api/v1{path}", json=body, headers={"Idempotency-Key": uuid.uuid4().hex, "X-Actor-Id": self.actor})
        assert r.status_code == 200, r.text
        return r.json()["data"]


@pytest.fixture
def app_env(tmp_path):
    settings = Settings(data_dir=tmp_path / "data", export_worker_enabled=False, report_worker_enabled=False)
    with TestClient(create_app(settings)) as client:
        api = Api(client)
        division = api.post("/organizations", {"name": "제조DX담당", "kind": "division"}, actor=False)
        projects = []
        for team_name, names in (("자동보정팀", ["재료교체 불량 개선", "코팅 온도 관리"]), ("추적솔루션팀", ["추적 데이터 통합"])):
            team = api.post("/organizations", {"name": team_name, "kind": "team", "parentId": division["id"]}, actor=False)
            user = api.post("/users", {"name": f"{team_name[:2]}담당", "teamId": team["id"], "employeeNumber": f"E{len(projects) + 1:04d}"}, actor=False)
            api.actor = user["id"]
            for name in names:
                project = api.post("/projects", {"name": name, "teamId": team["id"], "ownerUserId": user["id"], "status": "in_progress",
                                                 "startDate": "2026-09-01", "endDate": "2026-12-31",
                                                 "backgroundDoc": doc(f"{name} 배경입니다."), "purposeDoc": doc(f"{name} 목적입니다.")})
                milestone = api.post(f"/projects/{project['id']}/milestones",
                                     {"name": "로직 개발", "plannedStart": "2026-09-01", "plannedEnd": "2026-10-31"})
                for day in WORK_DAYS:
                    api.put(f"/projects/{project['id']}/logs/{day}/{user['id']}", {
                        "expectedRevision": 0,
                        "tasks": [{"milestoneId": milestone["id"], "content": doc(f"{name} 로직 시험 적용 결과 불량률 0.18%에서 0.15%로 감소 ({day})")}]})
                projects.append((project["id"], team_name))
        yield client, client.app.state.rt, api, projects


def _ids(prompt: str) -> list[str]:
    return list(dict.fromkeys(re.findall(r"기록 ID: (\S+) \(", prompt)))


def _answer(need: dict) -> dict:
    """프롬프트의 기록 ID·과제 ID만 근거로 형식에 맞는 응답을 만든다."""
    prompt, kind = need["prompt"], need["promptId"]
    ids = _ids(prompt) or []
    if kind in ("weekly_rollup", "period_rollup"):
        return {"headline": {"text": "로직 시험 적용으로 불량률이 0.18%에서 0.15%로 감소했습니다", "source_ids": ids[:1]},
                "progress": [{"text": "로직을 시험 적용해 불량률이 0.18%에서 0.15%로 감소한 것을 확인했습니다. (9/29)", "source_ids": ids[:1]}],
                "next_plan": [], "issues": [], "milestone_updates": []}
    if kind == "cumulative_update":
        return {"items": [{"text": "로직을 시험 적용해 불량률이 0.18%에서 0.15%로 감소했습니다. (9/29)", "source_ids": ids[:1]}],
                "pinned_facts": [], "new_pinned_facts": []}
    if kind == "project_summary":
        pid = re.search(r"과제 ID / 과제명: (\S+) /", prompt).group(1)
        return {"items": [
            {"category": "background", "text": "배경과 목적에 따라 공정 불량을 줄이기 위한 로직을 개발하는 과제입니다.", "source_ids": [pid]},
            {"category": "progress", "text": "로직을 시험 적용해 불량률이 0.18%에서 0.15%로 감소한 것을 확인했습니다.(9/29)", "source_ids": ids[:1]},
            {"category": "good", "text": "시험 적용 이틀 동안 불량률 감소 효과가 같은 수준으로 유지되었습니다.(9/30)", "source_ids": ids[-1:],
             "details": [{"text": "9/29과 9/30 모두 0.15%로 확인했습니다.", "source_ids": ids[-1:]}]},
            {"category": "plan", "text": "시험 적용 결과를 바탕으로 로직 개발 단계를 마무리할 예정입니다.", "source_ids": ids[:1]}]}
    raise AssertionError(f"예상하지 못한 프롬프트: {kind}")


def _run(client: TestClient, rt, api: Api, body: dict) -> tuple[dict, list[str]]:
    job = api.post("/reports/jobs", body)
    asked = []
    for _ in range(40):
        run_next(rt)
        current = client.get(f"/api/v1/reports/jobs/{job['id']}").json()["data"]
        if current["status"] != "need_response":
            break
        need = current["need"]
        asked.append(need["promptId"])
        api.post(f"/reports/jobs/{job['id']}/response", {"responseName": need["responseName"], "text": json.dumps(_answer(need), ensure_ascii=False)})
    return current, asked


def _titles(path) -> list[tuple[str, bool]]:
    """(제목, 요약 장 여부)"""
    out = []
    for slide in Presentation(str(path)).slides:
        shapes = {s.name: s for s in slide.shapes}
        title = shapes["slide_title"].text_frame.text if "slide_title" in shapes else ""
        out.append((title, "main_table" not in shapes and "pjt_header" in shapes))
    return out


def test_weekly_report_puts_team_summary_before_each_team(app_env):
    client, rt, api, projects = app_env
    job, asked = _run(client, rt, api, {"kind": "weekly", "template": "weekly", "week": WEEK, "projectIds": [p for p, _ in projects],
                                        "includeTeamSummary": True, "summaryAuthor": "홍길동 팀장"})
    assert job["status"] == "succeeded", job.get("error")
    assert asked.count("project_summary") == 3  # 과제마다 한 번 (붙여넣기 모드)
    pptx = job_file(rt.settings, job["id"], job["result"]["files"][0]["name"])
    slides = _titles(pptx)
    summaries = [title for title, is_summary in slides if is_summary]
    assert summaries == ["1. 자동보정팀 (1/1)", "1. 추적솔루션팀 (1/1)"]
    assert [is_summary for _, is_summary in slides] == [True, False, False, True, False]  # 팀별 [요약 → 과제 장표]
    first = {s.name: s for s in Presentation(str(pptx)).slides[0].shapes}
    assert first["author"].text_frame.text == "작성자 : 홍길동 팀장"
    body = first["pjt_header"].text_frame.text
    assert "1. 재료교체 불량 개선" in body and "2. 코팅 온도 관리" in body and "추적 데이터 통합" not in body
    blue = [r.text for p in first["pjt_header"].text_frame.paragraphs for r in p.runs
            if r.font.color and r.font.color.type is not None and str(r.font.color.rgb) == "1414FE"]
    assert blue and not any("배경과 목적" in t for t in blue)  # 배경은 검정, 이번 주 근거 문장만 파랑
    check = job_file(rt.settings, job["id"], job["result"]["files"][1]["name"]).read_text(encoding="utf-8")
    assert "팀장 요약 · 자동보정팀" in check and "작성자 : 홍길동 팀장 (입력값)" in check
    assert not [p for p in job["result"]["problems"] if "팀장 요약" in p], job["result"]["problems"]


def test_summary_off_keeps_previous_output(app_env):
    client, rt, api, projects = app_env
    job, asked = _run(client, rt, api, {"kind": "weekly", "template": "weekly", "week": WEEK, "projectIds": [p for p, _ in projects]})
    assert job["status"] == "succeeded" and "project_summary" not in asked
    assert job["options"]["includeTeamSummary"] is False
    slides = _titles(job_file(rt.settings, job["id"], job["result"]["files"][0]["name"]))
    assert len(slides) == 3 and not any(is_summary for _, is_summary in slides)  # 4장 템플릿의 참고 장도 들어가지 않음


def test_period_report_summary_uses_period_key_and_default_author(app_env):
    client, rt, api, projects = app_env
    ids = [p for p, team in projects if team == "추적솔루션팀"]
    job, asked = _run(client, rt, api, {"kind": "period", "template": "weekly", "dateFrom": "2026-09-21", "dateTo": "2026-10-04",
                                        "projectIds": ids, "includeTeamSummary": True})
    assert job["status"] == "succeeded", job.get("error")
    assert "period_rollup" in asked and "project_summary" in asked
    names = [p.name for p in JobStore(rt.settings).settings.reports_dir.joinpath("workspace/responses").glob("project_summary__*.json")]
    assert f"project_summary__{ids[0]}__2026-09-21_2026-10-04.json" in names
    pptx = job_file(rt.settings, job["id"], job["result"]["files"][0]["name"])
    first = {s.name: s for s in Presentation(str(pptx)).slides[0].shapes}
    assert first["slide_title"].text_frame.text == "1. 추적솔루션팀 (1/1)"
    assert first["author"].text_frame.text.startswith("작성자 : 추적")  # 입력이 없으면 첫 과제 담당자


def test_exec_template_ignores_team_summary(app_env):
    client, rt, api, projects = app_env
    job = api.post("/reports/jobs", {"kind": "weekly", "template": "exec", "week": WEEK, "projectIds": [projects[0][0]],
                                     "includeTeamSummary": True, "summaryAuthor": "홍길동 팀장"})
    assert job["options"]["includeTeamSummary"] is False and job["options"]["summaryAuthor"] is None
