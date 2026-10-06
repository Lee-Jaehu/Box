"""가져온 weekly_report 패키지의 팀장 요약 (DB 없이): 요약 JSON(period 포함) + 팀별 [요약 → 과제 장표] 그리기.

입력은 같은 저장소의 원본 weekly-report 예시 데이터(P-ASM-001 W40, 합성 5과제 W39)와 데모 응답을 쓴다.
원본 폴더가 없는 배포본에서는 건너뛴다.
"""
import json
import shutil
from datetime import date
from pathlib import Path

import pytest
from pptx import Presentation

from weekly_report.ai import ExaoneClient
from weekly_report.pptgen import prepare_ppt
from weekly_report.ppt.render import open_deck_template, shape_map
from weekly_report.summary import run_summary
from weekly_report.team import (FONT_SIZES, ITEM_MAR_IN, Block, TeamSection, inspect_summary, layout_summary, lines_per_project,
                                render_team_report, summary_geometry, summary_titles)

BACKEND = Path(__file__).resolve().parents[1]
ASSETS = BACKEND / "report_assets"
ORIGINAL = BACKEND.parents[1] / "weekly-report"
FIXTURES = ORIGINAL / "fixtures/weekly-report-fixture-kit"
TEMPLATE_NAME = "주간업무PPT_Template_v2.pptx"
TEAMS = {"조립자동보정팀": ("2026-W40", ["P-ASM-001"]),
         "검증팀": ("2026-W39", ["P-APC-101", "P-ROL-102", "P-INV-103"])}

pytestmark = pytest.mark.skipif(not (ORIGINAL / "demo/team/mock_responses").is_dir(), reason="원본 weekly-report 예시 데이터 없음")


@pytest.fixture(scope="module")
def ws(tmp_path_factory):
    """서비스 작업공간과 같은 배치: report_assets(config·schemas·prompts·템플릿·글꼴) + 예시 데이터."""
    root = tmp_path_factory.mktemp("ws")
    for sub in ("config", "schemas", "prompts"):
        shutil.copytree(ASSETS / sub, root / sub)
    shutil.copy(ASSETS / "templates" / TEMPLATE_NAME, root / TEMPLATE_NAME)
    shutil.copytree(ASSETS / "fonts", root / "fonts")
    shutil.copytree(ORIGINAL / "data/master", root / "data/master")
    shutil.copytree(ORIGINAL / "data/raw", root / "data/raw")
    shutil.copytree(ORIGINAL / "data/derived", root / "data/derived")
    shutil.copytree(ORIGINAL / "demo/w40/raw", root / "data/raw", dirs_exist_ok=True)
    shutil.copytree(ORIGINAL / "demo/w40/out/data/derived", root / "data/derived", dirs_exist_ok=True)
    shutil.copytree(FIXTURES / "fixture-root/data/raw", root / "data/raw", dirs_exist_ok=True)
    for pid in TEAMS["검증팀"][1]:
        shutil.copy(FIXTURES / f"fixture-root/data/master/projects/{pid}.json", root / f"data/master/projects/{pid}.json")
        shutil.copytree(FIXTURES / f"results/{pid}/data/derived", root / "data/derived", dirs_exist_ok=True)
    mocks = root / "mocks"
    mocks.mkdir()
    for folder in (ORIGINAL / "demo/team/mock_responses", ORIGINAL / "demo/w40/mock_responses", FIXTURES / "fixture-root/prompts/mock_responses"):
        for path in folder.glob("*.json"):
            shutil.copy(path, mocks / path.name)
    return root


def load(ws: Path, pid: str) -> dict:
    return json.loads((ws / f"data/master/projects/{pid}.json").read_text(encoding="utf-8"))


def build_sections(ws: Path, client: ExaoneClient, *, with_summary: bool) -> list[TeamSection]:
    template = ws / TEMPLATE_NAME
    sgeom = summary_geometry(ws, open_deck_template(template)[1])
    sections = []
    for team, (week, pids) in TEAMS.items():
        prepared, blocks = [], []
        for index, pid in enumerate(pids, 1):
            project = load(ws, pid)
            wpath, cpath = ws / f"data/derived/weekly/{pid}/{week}.json", ws / f"data/derived/cumulative/{pid}/{week}.json"
            prepared.append(prepare_ppt(ws, project, wpath, cpath, template, client=client, project_index=index, project_total=len(pids)))
            if with_summary:
                result = run_summary(ws, project, week, ws, client=client, max_lines=lines_per_project(sgeom, len(pids)),
                                     line_chars=sgeom.chars(FONT_SIZES[0], ITEM_MAR_IN))
                blocks.append(Block(index, result.summary["project_name"], result.summary["items"]))
        section = TeamSection(team, prepared)
        if with_summary:
            pages, _ = layout_summary(blocks, sgeom)
            section.pages, section.geom, section.titles = pages, sgeom, summary_titles(team, len(pages))
            section.author, section.updated_at = "작성자 : 홍길동 팀장", "업데이트 시간 : 10/4 15시"
        sections.append(section)
    return sections


def test_team_sections_render_summary_then_projects(ws, tmp_path):
    client = ExaoneClient(ws, mock_dir=ws / "mocks")
    sections = build_sections(ws, client, with_summary=True)
    layout = render_team_report(ws / TEMPLATE_NAME, sections, tmp_path / "deck.pptx")
    prs = Presentation(str(tmp_path / "deck.pptx"))
    kinds = ["summary" if "main_table" not in shape_map(s) else "weekly" for s in prs.slides]
    weekly_pages = [len(p["pages"]) for s in sections for p in s.prepared]
    assert kinds == ["summary"] + ["weekly"] * weekly_pages[0] + ["summary"] + ["weekly"] * sum(weekly_pages[1:])
    assert layout.summary_starts == [0, 1 + weekly_pages[0]]
    # 참고 슬라이드 삽입 위치 = 과제마다 마지막 장 번호(1부터)
    expected, position = [], 0
    for section in sections:
        position += len(section.pages)
        for p in section.prepared:
            position += len(p["pages"])
            expected.append(position)
    assert layout.groups == expected
    for section, start in zip(sections, layout.summary_starts):
        assert inspect_summary(prs, start, section.pages, section.geom, section.titles) == []
    titles = [shape_map(prs.slides[i])["slide_title"].text_frame.text for i in layout.summary_starts]
    assert titles == ["1. 조립자동보정팀 (1/1)", "1. 검증팀 (1/1)"]
    text = " ".join(s.text_frame.text for sl in prs.slides for s in sl.shapes if s.has_text_frame)
    assert "자동보정선행개발팀" not in text  # 템플릿 1번(실제 작성본 참고)은 출력하지 않음
    # 검증팀 과제 번호는 팀 안에서 (n/3)
    assert shape_map(prs.slides[layout.groups[1] - 1])["pjt_header"].text_frame.text.endswith("(1/3)")


def test_without_summary_only_weekly_pages(ws, tmp_path):
    client = ExaoneClient(ws, mock_dir=ws / "mocks")
    sections = build_sections(ws, client, with_summary=False)
    layout = render_team_report(ws / TEMPLATE_NAME, sections, tmp_path / "plain.pptx")
    prs = Presentation(str(tmp_path / "plain.pptx"))
    assert len(prs.slides) == sum(len(p["pages"]) for s in sections for p in s.prepared) == layout.groups[-1]
    assert all("main_table" in shape_map(s) for s in prs.slides) and layout.summary_starts == [None, None]


def test_period_summary_uses_response_key_and_period_wording(ws, tmp_path):
    key = "2026-09-28_2026-10-04"
    mocks = tmp_path / "mocks"
    shutil.copytree(ws / "mocks", mocks)
    shutil.copy(mocks / "project_summary__P-ASM-001__2026-W40.json", mocks / f"project_summary__P-ASM-001__{key}.json")
    sent = {}

    class Recording(ExaoneClient):
        def complete(self, prompt_id, project_id, week, system, user, variant=None):
            sent.update(name=f"{prompt_id}__{project_id}__{week}", system=system, user=user)
            return super().complete(prompt_id, project_id, week, system, user, variant)

    out = tmp_path / "out"
    shutil.copytree(ws / "data/derived", out / "data/derived")  # 기간 작업공간(out_root)에 W40 정리 결과가 있다고 가정
    result = run_summary(ws, load(ws, "P-ASM-001"), "2026-W40", out, client=Recording(ws, mock_dir=mocks), max_lines=12,
                         line_chars=70, period=(date(2026, 9, 28), date(2026, 10, 4)), response_key=key)
    assert sent["name"] == f"project_summary__P-ASM-001__{key}"
    assert "[보고 기간 정리본] (9/28~10/4" in sent["user"] and "[보고 기간 이전까지 누적 요약]" in sent["user"]
    assert "이번 주" not in sent["system"]
    assert result.path == out / f"data/derived/summary/P-ASM-001/{key}.json"
    assert result.summary["range"] == {"from": "2026-09-28", "to": "2026-10-04"}
    assert not [i for i in result.issues if i.level == "오류"], [i.format() for i in result.issues]
