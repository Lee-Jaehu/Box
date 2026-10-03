"""v2 템플릿에 내용을 채운다.

원칙
- 칸은 도형 이름으로 찾는다. 양식 크기·글꼴 크기는 바꾸지 않는다.
- 텍스트를 바꿀 때 템플릿의 문단(pPr)·글자(rPr) 서식을 복제해 ea(LG스마트체 Regular)·크기·굵기를 유지한다.
  영문은 Arial Narrow를 명시하고, 바꾸는 것은 색(파랑/검정)뿐이다.
- 마일스톤 표는 템플릿 데이터 행을 복제해 필요한 행 수만큼 만든다.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

from ..core import ValidationError
from ..fonts import EA_REGULAR, LATIN_FONT
from .budget import GAP_EMU, INSET_EMU, Geometry, ms_row_height, para_lines
from .milestones import MsRow
from .model import PageModel, Para, Run, SlideContent, cell_paras

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
REQUIRED_SHAPES = ("slide_title", "pjt_header", "author", "updated_at", "main_table", "body_top", "ms_table", "body_main")
SLIDE_SIZE_IN = (10.83, 7.5)
BLUE, BLACK = "0000FF", "000000"
EA_FONT = EA_REGULAR  # TTF nameID 1과 같은 이름이어야 PowerPoint가 LG스마트체를 쓴다


def a(tag: str) -> str:
    return f"{{{A_NS}}}{tag}"


# ---------------------------------------------------------------- 템플릿 검사

def open_template(template: Path):
    if not template.exists():
        raise FileNotFoundError(f"공식 템플릿 누락: {template}; 필수 도형: {', '.join(sorted(REQUIRED_SHAPES))}")
    prs = Presentation(str(template))
    size = (round(prs.slide_width / Inches(1), 2), round(prs.slide_height / Inches(1), 2))
    if size != SLIDE_SIZE_IN:
        raise ValidationError(f"템플릿 슬라이드 크기 {size[0]} × {size[1]}인치 ≠ 10.83 × 7.5인치")
    if len(prs.slides) != 1:
        raise ValidationError(f"템플릿 슬라이드 수 {len(prs.slides)} ≠ 1")
    shapes = shape_map(prs.slides[0])
    missing = [name for name in REQUIRED_SHAPES if name not in shapes]
    if missing:
        raise ValidationError(f"템플릿 필수 도형 누락: {', '.join(missing)}")
    if not shapes["main_table"].has_table or not shapes["ms_table"].has_table:
        raise ValidationError("main_table/ms_table이 표가 아님")
    main, ms = shapes["main_table"].table, shapes["ms_table"].table
    if (len(main.rows), len(main.columns)) != (3, 5):
        raise ValidationError(f"main_table 크기 {len(main.rows)}×{len(main.columns)} ≠ 3×5")
    if len(ms.columns) != 7 or len(ms.rows) < 2:
        raise ValidationError(f"ms_table 크기 {len(ms.rows)}×{len(ms.columns)} (7열, 데이터 행 1개 이상 필요)")
    return prs


def shape_map(slide) -> dict:
    return {shape.name: shape for shape in slide.shapes}


def read_geometry(slide) -> Geometry:
    shapes = shape_map(slide)
    main = shapes["main_table"]
    rows = main.table.rows
    area_top = main.top + rows[0].height + rows[1].height
    ms = shapes["ms_table"].table
    tc_pr = ms.cell(1, 0)._tc.tcPr
    margin = int(tc_pr.get("marL", 91440)) + int(tc_pr.get("marR", 91440)) if tc_pr is not None else 2 * 91440
    return Geometry(area_top=area_top, area_bottom=area_top + rows[2].height,
                    ms_col_widths=[c.width for c in ms.columns], ms_row_height=ms.rows[1].height, ms_cell_margin=margin,
                    body_width=shapes["body_main"].width)


def duplicate_slide(prs, source):
    """같은 레이아웃으로 새 슬라이드를 만들고 원본 도형을 복제한다 (템플릿에 이미지 관계가 없음을 전제)."""
    slide = prs.slides.add_slide(source.slide_layout)
    tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        tree.remove(shape._element)
    for element in source.shapes._spTree:
        if etree.QName(element).localname in {"sp", "graphicFrame", "grpSp", "cxnSp", "pic"}:
            tree.append(deepcopy(element))
    return slide


# ---------------------------------------------------------------- 텍스트 쓰기

class Proto:
    """txBody 안에서 찾은 서식 원형."""

    def __init__(self, tx_body):
        self.pPr = self.normal = self.bold = self.end = None
        for p in tx_body.iter(a("p")):
            if self.pPr is None and p.find(a("pPr")) is not None:
                self.pPr = deepcopy(p.find(a("pPr")))
            for r in p.findall(a("r")):
                rpr = r.find(a("rPr"))
                if rpr is None:
                    continue
                if rpr.get("b") == "1":
                    self.bold = self.bold if self.bold is not None else deepcopy(rpr)
                else:
                    self.normal = self.normal if self.normal is not None else deepcopy(rpr)
            if self.end is None and p.find(a("endParaRPr")) is not None:
                self.end = deepcopy(p.find(a("endParaRPr")))
        if self.normal is None and self.bold is None and self.end is not None:
            # 빈 칸(글자 없음)은 문단 끝 서식(endParaRPr)에서 크기·글꼴을 가져온다
            self.normal = deepcopy(self.end)
            self.normal.tag = a("rPr")
        if self.normal is None:
            self.normal = deepcopy(self.bold) if self.bold is not None else etree.Element(a("rPr"))
            self.normal.set("b", "0")
        if self.bold is None:
            self.bold = deepcopy(self.normal)
            self.bold.set("b", "1")

    def run_props(self, bold: bool) -> etree._Element:
        rpr = deepcopy(self.bold if bold else self.normal)
        for attr in ("dirty", "err"):
            rpr.attrib.pop(attr, None)
        return rpr


def _set_color(rpr, color: str) -> None:
    for old in rpr.findall(a("solidFill")) + rpr.findall(a("noFill")):
        rpr.remove(old)
    fill = etree.Element(a("solidFill"))
    etree.SubElement(fill, a("srgbClr")).set("val", color)
    ln = rpr.find(a("ln"))
    rpr.insert(list(rpr).index(ln) + 1 if ln is not None else 0, fill)


def _set_fonts(rpr) -> None:
    latin = rpr.find(a("latin"))
    if latin is None:
        latin = etree.Element(a("latin"))
        anchor = rpr.find(a("ea")) if rpr.find(a("ea")) is not None else rpr.find(a("cs"))
        if anchor is not None:
            anchor.addprevious(latin)
        else:
            rpr.append(latin)
    for key in list(latin.attrib):
        del latin.attrib[key]
    latin.set("typeface", LATIN_FONT)
    ea = rpr.find(a("ea"))
    if ea is None:
        ea = etree.Element(a("ea"))
        latin.addnext(ea)
    if ea.get("typeface") != EA_FONT:
        for key in list(ea.attrib):
            del ea.attrib[key]
        ea.set("typeface", EA_FONT)


def write_paras(tx_body, paras: list[Para], *, recolor: bool = True, proto: Proto | None = None) -> None:
    """txBody의 문단을 paras로 교체한다. recolor=False면 템플릿 글자색을 유지한다."""
    proto = proto or Proto(tx_body)
    for p in tx_body.findall(a("p")):
        tx_body.remove(p)
    for para in paras or [Para([Run("")], "blank")]:
        p = etree.SubElement(tx_body, a("p"))
        if proto.pPr is not None:
            p.append(deepcopy(proto.pPr))
        for run in para.runs:
            if run.text == "":
                continue
            r = etree.SubElement(p, a("r"))
            rpr = proto.run_props(run.bold)
            if recolor or run.blue:
                _set_color(rpr, BLUE if run.blue else BLACK)
            _set_fonts(rpr)
            r.append(rpr)
            etree.SubElement(r, a("t")).text = run.text
        end = deepcopy(proto.end) if proto.end is not None else etree.Element(a("endParaRPr"))
        p.append(end)


def write_cell(cell, value: "Run | list[Para]", *, recolor: bool = True) -> None:
    write_paras(cell._tc.txBody, cell_paras(value), recolor=recolor)


def write_shape(shape, paras: list[Para], *, recolor: bool = True) -> None:
    write_paras(shape.text_frame._txBody, paras, recolor=recolor)


# ---------------------------------------------------------------- 마일스톤 표

def _set_cell_fill(tc, color: str) -> None:
    tc_pr = tc.find(a("tcPr"))
    if tc_pr is None:
        tc_pr = etree.SubElement(tc, a("tcPr"))
    for tag in ("solidFill", "noFill", "gradFill", "pattFill", "grpFill", "blipFill"):
        for old in tc_pr.findall(a(tag)):
            tc_pr.remove(old)
    fill = etree.Element(a("solidFill"))
    etree.SubElement(fill, a("srgbClr")).set("val", color)
    anchor = None
    for child in tc_pr:
        if etree.QName(child).localname in {"lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr", "cell3D"}:
            anchor = child
    if anchor is not None:
        anchor.addnext(fill)
    else:
        tc_pr.insert(0, fill)


def fill_ms_table(shape, rows: list[MsRow], geom: Geometry) -> int:
    """데이터 행을 rows 수만큼 만들고 채운다. 표 전체 높이(EMU)를 돌려준다."""
    tbl = shape.table._tbl
    trs = tbl.findall(a("tr"))
    prototype = deepcopy(trs[1])
    protos = [Proto(tc.find(a("txBody"))) for tc in prototype.findall(a("tc"))]
    for tr in trs[1:]:
        tbl.remove(tr)
    total = int(trs[0].get("h"))
    for row in rows:
        tr = deepcopy(prototype)
        height = ms_row_height(row, geom)
        tr.set("h", str(height))
        total += height
        for col, tc in enumerate(tr.findall(a("tc"))):
            write_paras(tc.find(a("txBody")), [Para([Run(row.cells[col], row.blue[col])])], proto=protos[col])
        tbl.append(tr)
        status_tc = tr.findall(a("tc"))[5]
        _set_cell_fill(status_tc, row.fill)
    shape.height = total
    return total


# ---------------------------------------------------------------- 슬라이드 채우기

def fill_slide(slide, page: PageModel, content: SlideContent, geom: Geometry) -> None:
    shapes = shape_map(slide)
    write_shape(shapes["slide_title"], [Para([Run(content.title)])], recolor=False)
    write_shape(shapes["pjt_header"], [Para([Run(page.pjt_header)])], recolor=False)
    write_shape(shapes["author"], [Para([Run(content.author)])], recolor=False)
    write_shape(shapes["updated_at"], [Para([Run(content.updated_at)])], recolor=False)

    main = shapes["main_table"].table
    write_cell(main.cell(0, 2), content.main["week_header"], recolor=False)
    write_cell(main.cell(1, 0), content.main["name"])
    write_cell(main.cell(1, 1), content.main["target"])
    write_cell(main.cell(1, 2), content.main["headline"])
    write_cell(main.cell(1, 3), content.main["schedule"])
    write_cell(main.cell(1, 4), content.main["owner"])
    write_cell(main.cell(2, 2), Run(""))  # r2c2는 비워 둔다 (본문은 도형으로 올린다)

    y = geom.area_top + INSET_EMU
    top = shapes["body_top"]
    write_shape(top, page.body_top)
    top_lines = sum(para_lines(p, geom.body_chars) for p in page.body_top) if page.body_top else 0
    top.top, top.height = y, max(top_lines, 1) * geom.line_emu
    if page.body_top:
        y += top.height + GAP_EMU

    ms = shapes["ms_table"]
    if page.ms_rows:
        ms.top = y
        y += fill_ms_table(ms, page.ms_rows, geom) + GAP_EMU
    else:
        slide.shapes._spTree.remove(ms._element)

    body = shapes["body_main"]
    write_shape(body, page.body)
    body.top = y
    body.height = max(geom.area_bottom - INSET_EMU - y, geom.line_emu)


def theme_parts(prs):
    return [master.part.part_related_by(RT.THEME) for master in prs.slide_masters]


def fix_theme_fonts(prs) -> list[str]:
    """테마 글꼴 체계의 ea를 실제 글꼴 이름으로 맞춘다 (출력 파일만, 원본 템플릿은 그대로).

    v2 템플릿 테마는 ea가 'LG Smart Regular'인데 실제 TTF 이름은 'LG스마트체 Regular'라
    ea를 지정하지 않은 글자(상속 텍스트)는 대체 글꼴로 표시된다.
    """
    fixed = []
    for part in theme_parts(prs):
        root = etree.fromstring(part.blob)
        changed = False
        for scheme in ("majorFont", "minorFont"):
            for ea in root.iter(a(scheme)):
                node = ea.find(a("ea"))
                if node is not None and node.get("typeface") != EA_FONT:
                    fixed.append(f"테마 {scheme} ea '{node.get('typeface')}' → '{EA_FONT}'")
                    node.set("typeface", EA_FONT)
                    changed = True
        if changed:
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return fixed


def render(template: Path, content: SlideContent, pages: list[PageModel], output: Path,
           geom: Geometry | None = None) -> list[str]:
    """geom은 페이지 나누기에 쓴 것과 같은 값을 넘긴다 (표 행 높이 계산 일치)."""
    prs = open_template(template)
    notes = fix_theme_fonts(prs)
    geom = geom or read_geometry(prs.slides[0])
    slides = [prs.slides[0]] + [duplicate_slide(prs, prs.slides[0]) for _ in pages[1:]]
    for slide, page in zip(slides, pages):
        fill_slide(slide, page, content, geom)
    output.parent.mkdir(parents=True, exist_ok=True)
    tmp = output.with_name(f".{output.name}.tmp")
    prs.save(str(tmp))
    tmp.replace(output)
    return notes
