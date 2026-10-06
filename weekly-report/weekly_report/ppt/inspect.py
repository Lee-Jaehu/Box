"""생성한 PPT를 다시 열어 규칙 준수 여부를 검사한다 (색·글꼴·크기·표 값·슬라이드 수·경계)."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation

from .budget import Geometry
from .model import PageModel, SlideContent
from .model import cell_paras, cell_text
from .render import A_NS, BLUE, EA_FONT, LATIN_FONT, REQUIRED_SHAPES, a, read_geometry, shape_map, theme_parts

MAX_SLIDES = 2
BODY_SIZE = 900  # 9pt


def _runs(tx_body):
    for p in tx_body.iter(a("p")):
        for r in p.findall(a("r")):
            yield r


def _run_info(r) -> tuple[str, str | None, str | None, str | None, str | None]:
    rpr = r.find(a("rPr"))
    text = r.findtext(a("t")) or ""
    if rpr is None:
        return text, None, None, None, None
    color = rpr.find(f"{a('solidFill')}/{a('srgbClr')}")
    latin, ea = rpr.find(a("latin")), rpr.find(a("ea"))
    return (text, rpr.get("sz"), latin.get("typeface") if latin is not None else None,
            ea.get("typeface") if ea is not None else None, color.get("val") if color is not None else None)


def expected_blue(content: SlideContent, page: PageModel) -> list[str]:
    blue = [r.text for key in ("headline", "schedule") for p in cell_paras(content.main[key]) for r in p.runs if r.blue and r.text]
    blue += [r.text for p in page.body for r in p.runs if r.blue]
    blue += [row.cells[i] for row in page.ms_rows for i, flag in enumerate(row.blue) if flag]
    return sorted(blue)


def theme_problems(prs) -> list[str]:
    problems = []
    for part in theme_parts(prs):
        root = etree.fromstring(part.blob)
        for scheme in ("majorFont", "minorFont"):
            for node in root.iter(a(scheme)):
                ea = node.find(a("ea"))
                if ea is not None and ea.get("typeface") != EA_FONT:
                    problems.append(f"테마 {scheme} ea '{ea.get('typeface')}' ≠ {EA_FONT}")
    return problems


def inspect_pptx(path: "Path | None", content: SlideContent, pages: list[PageModel], *,
                 prs=None, start: int | None = None) -> list[str]:
    """문제 목록 (빈 목록이면 통과).

    start가 없으면 과제 1건짜리 파일 전체를 검사한다(슬라이드 수·테마 포함).
    묶음 파일은 prs와 그 과제의 첫 슬라이드 번호(start, 0부터)를 넘긴다.
    """
    problems: list[str] = []
    prs = prs if prs is not None else Presentation(str(path))
    slides = list(prs.slides)
    if start is None:
        if len(slides) != len(pages):
            problems.append(f"슬라이드 수 {len(slides)} ≠ 기대 {len(pages)}")
        problems += theme_problems(prs)
        start = 0
    if len(pages) > MAX_SLIDES:
        problems.append(f"과제당 최대 {MAX_SLIDES}장 초과")
    if start + len(pages) > len(slides):
        problems.append(f"슬라이드 부족: {start + 1}~{start + len(pages)}장 기대, 전체 {len(slides)}장")
    for index, (slide, page) in enumerate(zip(slides[start:], pages), start + 1):
        tag = f"{index}장"
        shapes = shape_map(slide)
        needed = [n for n in REQUIRED_SHAPES if not (n == "ms_table" and not page.ms_rows)]
        missing = [n for n in needed if n not in shapes]
        if missing:
            problems.append(f"{tag}: 도형 누락 {missing}")
            continue
        geom: Geometry = read_geometry_from(shapes)
        # 1) 글꼴·크기 (본문 영역과 표)
        bodies = [("body_top", shapes["body_top"].text_frame._txBody), ("body_main", shapes["body_main"].text_frame._txBody)]
        main = shapes["main_table"].table
        bodies += [(f"main_table r1c{c}", main.cell(1, c)._tc.txBody) for c in range(5)]
        if page.ms_rows:
            ms = shapes["ms_table"].table
            bodies += [(f"ms_table r{r}c{c}", ms.cell(r, c)._tc.txBody) for r in range(1, len(ms.rows)) for c in range(7)]
        blue_found: list[str] = []
        for name, body in bodies:
            for r in _runs(body):
                text, size, latin, ea, color = _run_info(r)
                if not text:
                    continue
                if size != str(BODY_SIZE):
                    problems.append(f"{tag} {name}: 글자 크기 {size} ≠ 9pt ('{text[:15]}')")
                if latin != LATIN_FONT or ea != EA_FONT:
                    problems.append(f"{tag} {name}: 글꼴 latin={latin}, ea={ea} ('{text[:15]}')")
                if color not in {BLUE, "000000"}:
                    problems.append(f"{tag} {name}: 허용되지 않은 글자색 {color} ('{text[:15]}')")
                if color == BLUE:
                    blue_found.append(text)
        # 1-2) 슬라이드의 모든 글자(제목·머리글 포함)가 LG스마트체(ea)를 쓰는지
        for shape in slide.shapes:
            bodies_all = [shape.text_frame._txBody] if shape.has_text_frame else []
            if shape.has_table:
                bodies_all += [c._tc.txBody for row in shape.table.rows for c in row.cells]
            for body in bodies_all:
                for r in _runs(body):
                    text, _, _, ea, _ = _run_info(r)
                    if text.strip() and ea != EA_FONT:
                        problems.append(f"{tag} {shape.name}: 한글 글꼴(ea) {ea} ≠ {EA_FONT} ('{text[:15]}')")
        # 2) 파란색 대상 일치
        want = expected_blue(content, page)
        if sorted(blue_found) != want:
            problems.append(f"{tag}: 파란색 글자 불일치 (기대 {len(want)}개, 실제 {len(blue_found)}개)")
        # 3) 표 값 일치
        if page.ms_rows:
            ms = shapes["ms_table"].table
            if len(ms.rows) - 1 != len(page.ms_rows):
                problems.append(f"{tag}: 마일스톤 행 수 {len(ms.rows) - 1} ≠ {len(page.ms_rows)}")
            for r, row in enumerate(page.ms_rows, 1):
                if r >= len(ms.rows):
                    break
                actual = [ms.cell(r, c).text for c in range(7)]
                if actual != row.cells:
                    problems.append(f"{tag}: 마일스톤 {row.milestone_id} 값 불일치 {actual} ≠ {row.cells}")
                fill = ms.cell(r, 5)._tc.find(f"{a('tcPr')}/{a('solidFill')}/{a('srgbClr')}")
                if fill is None or fill.get("val") != row.fill:
                    problems.append(f"{tag}: 마일스톤 {row.milestone_id} 상태 칸 배경 ≠ {row.fill}")
        expected_main = {(1, 0): cell_text(content.main["name"]), (1, 2): cell_text(content.main["headline"]),
                         (1, 3): cell_text(content.main["schedule"]), (1, 4): cell_text(content.main["owner"])}
        for (r, c), value in expected_main.items():
            if main.cell(r, c).text != value:
                problems.append(f"{tag}: main_table r{r}c{c} 값 불일치")
        if main.cell(2, 2).text.strip():
            problems.append(f"{tag}: main_table r2c2는 비어 있어야 함")
        # 4) 영역 경계·분량
        for name in ("body_top", "ms_table", "body_main"):
            if name in shapes and shapes[name].top + shapes[name].height > geom.area_bottom:
                problems.append(f"{tag}: {name}이 본문 영역 아래로 넘침")
        if page.body_lines > page.body_capacity:
            problems.append(f"{tag}: 본문 {page.body_lines}줄 > 가용 {page.body_capacity}줄")
    return problems


def read_geometry_from(shapes) -> Geometry:
    class _Slide:  # read_geometry는 slide.shapes 이름 조회만 쓴다
        def __init__(self, values):
            self.shapes = values

    if "ms_table" in shapes:
        return read_geometry(_Slide(list(shapes.values())))
    main = shapes["main_table"]
    rows = main.table.rows
    top = main.top + rows[0].height + rows[1].height
    return Geometry(top, top + rows[2].height, [], 0, 0)


__all__ = ["inspect_pptx", "A_NS"]
