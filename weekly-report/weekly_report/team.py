"""팀 주간보고 묶음 PPT: 팀 요약 페이지(템플릿 0번 양식) → 과제별 주간 장표.

- 요약 페이지 문장은 과제마다 project_summary JSON(EXAONE)에서 가져오고, 배치·색·글자 크기·페이지 나누기는 Rule.
- 표시 형식은 템플릿 1번(실제 작성본)과 같다: 굵은 "n. 과제명", "- 문장"(신규 파랑 1414FE), "  . 세부".
  카테고리 이름은 표시하지 않고 순서(배경 → 진행 → 이슈·잘한점 → 계획)만 지킨다.
- 11pt로 한 장에 안 들어가면 10.5 → 10pt로 줄이고, 그래도 넘치면 과제 단위로 다음 장(제목 "(1/2)")에 넘긴다.
- 과제별 주간 장표는 단일 과제 PPT와 같은 코드(pptgen.prepare_content + render.add_weekly_pages)를 쓴다.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation

from . import sources
from .ai import ExaoneClient
from .codes import PeopleTable
from .core import ValidationError
from .fonts import load_fonts
from .ppt.budget import BODY_WIDTH_SAFETY, EMU_PER_IN, EMU_PER_PT, LINE_FACTOR_DEFAULT, LINE_FACTOR_LG
from .ppt.compose import updated_at_label
from .ppt.inspect import _run_info, _runs, inspect_pptx, theme_problems
from .ppt.model import Para, Run
from .ppt.render import (BLACK, EA_FONT, LATIN_FONT, _set_color, _set_fonts, a, drop_slide, duplicate_slide, fix_theme_fonts,
                         open_deck_template, shape_map, write_shape, add_weekly_pages)
from .pptgen import Prepared, check_lines, prepare_content, weekly_geometry
from .summary import ITEM_PREFIX, SummaryResult, run_summary
from .textmetrics import line_count, weighted_length

SUMMARY_BLUE = "1414FE"  # 템플릿 1번(실제 작성본)의 신규 표시 색
FONT_SIZES = (11.0, 10.5, 10.0)  # 실제 작성본 11pt, 넘친 과제는 10.5pt
AREA_BOTTOM_IN = 7.10  # 슬라이드 아래 쪽번호(7.27in) 위까지
ITEM_MAR_IN, DETAIL_MAR_IN, HANG_IN = 0.12, 0.30, 0.12  # 내어쓰기: "- " / ". " 뒤로 줄바꿈 정렬
LINES_PER_PROJECT = (5, 12)  # 프롬프트에 주는 과제당 줄 수 범위


@dataclass
class SummaryGeometry:
    top: int  # pjt_header 위쪽 (EMU)
    bottom: int
    width: int
    hangul_em: float = 1.0
    line_factor: float = LINE_FACTOR_DEFAULT

    def line_emu(self, size: float) -> int:
        return int(size * self.line_factor * EMU_PER_PT)

    def page_lines(self, size: float) -> int:
        return (self.bottom - self.top) // self.line_emu(size)

    def chars(self, size: float, margin_in: float = 0.0) -> float:
        """한 줄 가중 글자 수 (내어쓰기 여백 제외, 8% 여유)."""
        usable = self.width - int(margin_in * EMU_PER_IN)
        return usable / EMU_PER_PT / (size * self.hangul_em) * BODY_WIDTH_SAFETY


@dataclass
class Block:
    """요약 페이지의 과제 1건."""

    number: int
    name: str
    items: list[dict[str, Any]]

    def paras(self) -> list[tuple[str, list[tuple[str, bool, bool]]]]:
        """(종류, [(글자, 굵게, 파랑)]) 목록. 종류: title / item / detail."""
        out = [("title", [(f"{self.number}. {self.name}", True, False)])]
        for item in self.items:
            out.append(("item", [(ITEM_PREFIX + item["text"], False, item["new"])]))
            out += [("detail", [(". " + d["text"], False, d["new"])]) for d in item["details"]]
        return out

    def lines(self, geom: SummaryGeometry, size: float) -> int:
        total = 0
        for kind, runs in self.paras():
            text = "".join(r[0] for r in runs)
            margin = {"title": 0.0, "item": ITEM_MAR_IN, "detail": DETAIL_MAR_IN}[kind]
            # 첫 줄은 내어쓰기로 왼쪽 끝부터 시작하지만, 보수적으로 모든 줄을 여백 제외 폭으로 센다
            total += line_count(text, geom.chars(size, margin))
        return total


@dataclass
class SummaryPage:
    blocks: list[Block]
    size: float
    lines: int
    capacity: int


@dataclass
class TeamResult:
    pptx: Path
    check: Path
    notes: list[str] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    summaries: list[SummaryResult] = field(default_factory=list)


# ---------------------------------------------------------------- 대상 과제

def team_projects(root: Path, team_or_project: str) -> tuple[str, list[dict[str, Any]]]:
    """팀 이름 또는 과제 ID → (팀 이름, 그 팀 과제 목록). 순서: 내부 기준정보 먼저, WorkLog 과제는 이름순."""
    listed = sources.list_projects(root)
    loaded = {p["project_id"]: sources.load_project(root, p["project_id"]) for p in listed}
    team = team_or_project
    if team_or_project in loaded:
        team = loaded[team_or_project]["org"]["team"]
    projects = [loaded[p["project_id"]] for p in listed if loaded[p["project_id"]]["org"].get("team") == team]
    if not projects:
        teams = sorted({p["org"].get("team") or "-" for p in loaded.values()})
        raise ValidationError(f"팀 과제 없음: {team_or_project} (있는 팀: {', '.join(teams)})")
    return team, projects


# ---------------------------------------------------------------- 요약 페이지 배치

def summary_geometry(root: Path, slide) -> SummaryGeometry:
    shapes = shape_map(slide)
    header = shapes["pjt_header"]
    fonts = load_fonts(root.resolve())
    geom = SummaryGeometry(top=header.top, bottom=int(AREA_BOTTOM_IN * EMU_PER_IN), width=header.width,
                           hangul_em=fonts.hangul_em)
    if fonts.regular:
        geom.line_factor = LINE_FACTOR_LG
    return geom


def lines_per_project(geom: SummaryGeometry, count: int) -> int:
    """프롬프트에 주는 과제당 줄 수 = (페이지 줄 수 − 과제 수 × 2[제목·빈 줄]) ÷ 과제 수, 5~12줄."""
    low, high = LINES_PER_PROJECT
    return max(low, min(high, (geom.page_lines(FONT_SIZES[0]) - 2 * count) // max(count, 1)))


def _pack(blocks: list[Block], geom: SummaryGeometry, size: float) -> list[SummaryPage]:
    """과제 단위로 순서대로 채운다 (과제 사이 빈 줄 1)."""
    capacity = geom.page_lines(size)
    pages: list[SummaryPage] = []
    current: list[Block] = []
    used = 0
    for block in blocks:
        need = block.lines(geom, size) + (1 if current else 0)
        if current and used + need > capacity:
            pages.append(SummaryPage(current, size, used, capacity))
            current, used, need = [], 0, block.lines(geom, size)
        current.append(block)
        used += need
    if current:
        pages.append(SummaryPage(current, size, used, capacity))
    return pages


def layout_summary(blocks: list[Block], geom: SummaryGeometry) -> tuple[list[SummaryPage], list[str]]:
    """한 장에 들어가는 가장 큰 글자 크기를 고르고, 10pt로도 넘치면 11pt로 과제 단위 페이지 나누기."""
    for size in FONT_SIZES:
        pages = _pack(blocks, geom, size)
        if len(pages) == 1 and pages[0].lines <= pages[0].capacity:
            note = f"요약 페이지 {size:g}pt 1장 ({pages[0].lines}/{pages[0].capacity}줄)"
            if size != FONT_SIZES[0]:
                note += f" — {FONT_SIZES[0]:g}pt로는 넘쳐 글자 크기를 줄임"
            return pages, [note]
    pages = _pack(blocks, geom, FONT_SIZES[0])
    notes = [f"요약 페이지 {FONT_SIZES[-1]:g}pt로도 한 장 초과 → {FONT_SIZES[0]:g}pt, 과제 단위로 {len(pages)}장"]
    notes += [f"요약 {i}장: 과제 {', '.join(str(b.number) for b in p.blocks)} ({p.lines}/{p.capacity}줄)" for i, p in enumerate(pages, 1)]
    return pages, notes


# ---------------------------------------------------------------- 요약 페이지 쓰기

def _para(kind: str, runs: list[tuple[str, bool, bool]], size: float, space_before: bool) -> etree._Element:
    p = etree.Element(a("p"))
    ppr = etree.SubElement(p, a("pPr"))
    margin = {"title": 0.0, "item": ITEM_MAR_IN, "detail": DETAIL_MAR_IN}[kind]
    ppr.set("marL", str(int(margin * EMU_PER_IN)))
    ppr.set("indent", str(-int(HANG_IN * EMU_PER_IN) if kind != "title" else 0))
    ppr.set("algn", "l")
    spc = etree.SubElement(ppr, a("spcBef"))
    etree.SubElement(spc, a("spcPts")).set("val", str(int(size * 100) if space_before else 0))  # 과제 사이 빈 줄 1
    etree.SubElement(ppr, a("buNone"))
    for text, bold, blue in runs:
        r = etree.SubElement(p, a("r"))
        rpr = etree.SubElement(r, a("rPr"))
        rpr.set("lang", "ko-KR")
        rpr.set("altLang", "en-US")
        rpr.set("sz", str(int(round(size * 100))))
        rpr.set("b", "1" if bold else "0")
        _set_color(rpr, SUMMARY_BLUE if blue else BLACK)
        _set_fonts(rpr)
        etree.SubElement(r, a("t")).text = text
    end = etree.SubElement(p, a("endParaRPr"))
    end.set("lang", "ko-KR")
    end.set("sz", str(int(round(size * 100))))
    return p


def _fit_width(shape, text: str, hangul_em: float) -> None:
    """줄바꿈 없는 제목 상자: 글자보다 좁으면 왼쪽 위치는 두고 너비만 늘린다 (양쪽으로 커지며 잘리는 것 방지)."""
    sizes = [int(v) for v in shape._element.xpath(".//a:rPr/@sz")]
    size_pt = (max(sizes) if sizes else 1800) / 100
    body = shape.text_frame._txBody.find(a("bodyPr"))
    insets = int(body.get("lIns", 91440)) + int(body.get("rIns", 91440)) if body is not None else 2 * 91440
    need = int(weighted_length(text) * size_pt * max(hangul_em, 0.9) * EMU_PER_PT * 1.1) + insets
    if shape.width < need:
        shape.width = need


def fill_summary_slide(slide, page: SummaryPage, geom: SummaryGeometry, *, title: str, author: str, updated_at: str) -> None:
    shapes = shape_map(slide)
    write_shape(shapes["slide_title"], [Para([Run(title)])], recolor=False)
    _fit_width(shapes["slide_title"], title, geom.hangul_em)
    write_shape(shapes["author"], [Para([Run(author)])], recolor=False)
    write_shape(shapes["updated_at"], [Para([Run(updated_at)])], recolor=False)
    header = shapes["pjt_header"]
    body = header.text_frame._txBody
    for p in body.findall(a("p")):
        body.remove(p)
    for index, block in enumerate(page.blocks):
        for p_index, (kind, runs) in enumerate(block.paras()):
            body.append(_para(kind, runs, page.size, space_before=index > 0 and p_index == 0))
    # spcBef 1줄은 빈 줄 1개와 같은 높이로 계산했다
    header.top = geom.top
    header.height = max(page.lines, 1) * geom.line_emu(page.size)


def inspect_summary(prs, start: int, pages: list[SummaryPage], geom: SummaryGeometry, titles: list[str]) -> list[str]:
    """요약 장 재검사: 제목, 글꼴(ea·latin), 글자 크기, 색(검정·파랑만), 파랑 = 신규 항목, 영역 경계."""
    problems = []
    slides = list(prs.slides)
    for offset, (page, title) in enumerate(zip(pages, titles)):
        tag = f"{start + offset + 1}장(요약)"
        if start + offset >= len(slides):
            problems.append(f"{tag}: 슬라이드 없음")
            continue
        shapes = shape_map(slides[start + offset])
        if shapes["slide_title"].text_frame.text != title:
            problems.append(f"{tag}: 제목 '{shapes['slide_title'].text_frame.text}' ≠ '{title}'")
        want_blue = sorted(text for b in page.blocks for kind, runs in b.paras() for text, _, blue in runs if blue)
        blue = []
        for r in _runs(shapes["pjt_header"].text_frame._txBody):
            text, size, latin, ea, color = _run_info(r)
            if not text:
                continue
            if size != str(int(round(page.size * 100))):
                problems.append(f"{tag}: 글자 크기 {size} ≠ {page.size:g}pt ('{text[:15]}')")
            if latin != LATIN_FONT or ea != EA_FONT:
                problems.append(f"{tag}: 글꼴 latin={latin}, ea={ea} ('{text[:15]}')")
            if color not in {SUMMARY_BLUE, BLACK}:
                problems.append(f"{tag}: 허용되지 않은 글자색 {color} ('{text[:15]}')")
            if color == SUMMARY_BLUE:
                blue.append(text)
        if sorted(blue) != want_blue:
            problems.append(f"{tag}: 파란색 글자 불일치 (기대 {len(want_blue)}개, 실제 {len(blue)}개)")
        header = shapes["pjt_header"]
        if header.top + header.height > geom.bottom or page.lines > page.capacity:
            problems.append(f"{tag}: 본문 {page.lines}줄 > 가용 {page.capacity}줄 (영역 아래로 넘침)")
    return problems


# ---------------------------------------------------------------- 묶음 PPT

def _derived_paths(pid: str, week: str, out_root: Path, root: Path) -> tuple[Path, Path] | None:
    for base in dict.fromkeys([out_root, root]):
        weekly = base / f"data/derived/weekly/{pid}/{week}.json"
        cumulative = base / f"data/derived/cumulative/{pid}/{week}.json"
        if weekly.exists() and cumulative.exists():
            return weekly, cumulative
    return None


def generate_team_deck(root: Path, team_or_project: str, week: str, output: Path, *, out_root: Path | None = None,
                       client: ExaoneClient | None = None, mode: str = "mock", template: Path | None = None,
                       now: datetime | None = None) -> TeamResult:
    """팀 요약 페이지 + 과제별 주간 장표를 한 파일로 만든다. weekly/cumulative가 없으면 먼저 weekly를 실행한다."""
    from .pptgen import find_template
    from .weekly import run_weekly

    out_root = out_root or root
    client = client if client is not None else ExaoneClient(root, mode)
    template = find_template(root, template)
    team, projects = team_projects(root, team_or_project)
    prs, summary_slide, weekly_slide = open_deck_template(template)  # 요약·주간 양식 둘 다 필요
    sgeom = summary_geometry(root, summary_slide)
    wgeom, fonts = weekly_geometry(root, weekly_slide)
    budget = lines_per_project(sgeom, len(projects))
    notes = [f"팀 {team}: 과제 {len(projects)}건 ({', '.join(p['project_id'] for p in projects)})",
             f"요약 페이지 과제당 {budget}줄 배정 (11pt 한 장 {sgeom.page_lines(FONT_SIZES[0])}줄, 한 줄 약 {int(sgeom.chars(FONT_SIZES[0], ITEM_MAR_IN))}자)"]

    # 1) 과제별 입력 준비: weekly/cumulative(없으면 생성) → 요약 JSON → 주간 장표 내용
    summaries: list[SummaryResult] = []
    prepared: list[Prepared] = []
    for index, project in enumerate(projects, 1):
        pid = project["project_id"]
        paths = _derived_paths(pid, week, out_root, root)
        if paths is None:
            run_weekly(root, pid, week, out_root, client=client)
            notes.append(f"{pid}: {week} weekly/cumulative가 없어 주간 정리를 먼저 실행")
            paths = _derived_paths(pid, week, out_root, root)
            assert paths is not None
        prep = prepare_content(root, project, paths[0], paths[1], wgeom, client=client, project_index=index,
                               project_total=len(projects))
        prepared.append(prep)
        result = run_summary(root, project, week, out_root, client=client, max_lines=budget,
                             line_chars=sgeom.chars(FONT_SIZES[0], ITEM_MAR_IN), current_project=prep.current)
        summaries.append(result)
        counts = {lv: sum(i.level == lv for i in result.issues) for lv in ("오류", "주의")}
        notes.append(f"{pid}: 요약 {'재사용 (입력 동일)' if result.reused else '생성'}"
                     f" — 오류 {counts['오류']}건, 주의 {counts['주의']}건 ({result.check.name})")

    # 2) 요약 페이지 배치
    blocks = [Block(i, s.summary["project_name"], s.summary["items"]) for i, s in enumerate(summaries, 1)]
    pages, layout_notes = layout_summary(blocks, sgeom)
    notes += layout_notes
    people = PeopleTable.load(root).with_project(projects[0])
    author = f"작성자 : {people.name_with_title(projects[0]['owner'])}"
    # 업데이트 시간: 지정하지 않으면 과제 주간 정리본 중 가장 늦은 수정 시각 (과제별 장표의 업데이트 시간과 맞춘다)
    stamp = now.isoformat() if now else max(p.weekly["meta"]["updated_at"] for p in prepared)
    updated_at = updated_at_label(stamp)
    titles = [f"1. {team} ({i}/{len(pages)})" for i in range(1, len(pages) + 1)]

    # 3) 그리기: 깨끗한 양식을 먼저 복제한 뒤 채운다. 순서 = 요약 n장 → 과제별 주간 장표
    notes += fix_theme_fonts(prs)
    summary_slides = [summary_slide] + [duplicate_slide(prs, summary_slide) for _ in pages[1:]]
    for slide, page, title in zip(summary_slides, pages, titles):
        fill_summary_slide(slide, page, sgeom, title=title, author=author, updated_at=updated_at)
    starts = []
    position = len(pages)
    for prep in prepared:
        add_weekly_pages(prs, weekly_slide, prep.content, prep.pages, wgeom)
        starts.append(position)
        position += len(prep.pages)
    drop_slide(prs, weekly_slide)  # 양식 원본 장은 출력에서 뺀다
    output.parent.mkdir(parents=True, exist_ok=True)
    tmp = output.with_name(f".{output.name}.tmp")
    prs.save(str(tmp))
    tmp.replace(output)

    # 4) 재검사
    saved = Presentation(str(output))
    problems = theme_problems(saved)
    expected = len(pages) + sum(len(p.pages) for p in prepared)
    if len(saved.slides) != expected:
        problems.append(f"슬라이드 수 {len(saved.slides)} ≠ 기대 {expected}")
    problems += inspect_summary(saved, 0, pages, sgeom, titles)
    project_problems = [inspect_pptx(None, p.content, p.pages, prs=saved, start=s) for p, s in zip(prepared, starts)]
    for prep, found in zip(prepared, project_problems):
        problems += [f"{prep.project['project_id']}: {p}" for p in found]

    check = output.with_name(f"{output.stem}_check.txt")
    lines = [f"PPT: {output.name}", f"템플릿: {template.name}", f"팀: {team} / {week}",
             f"슬라이드: 요약 {len(pages)}장 + 과제별 {expected - len(pages)}장 = {expected}장", "", "[처리 내역]"]
    lines += [f"- {n}" for n in notes]
    lines += ["", "[요약 페이지 재검사]"]
    summary_found = [p for p in problems if "(요약)" in p or p.startswith("테마") or p.startswith("슬라이드 수")]
    lines += [f"- 문제: {p}" for p in summary_found] or ["- 통과 (제목·글꼴·글자 크기·색·파랑=이번 주 근거·영역 경계)"]
    for prep, start, found in zip(prepared, starts, project_problems):
        lines += ["", f"===== {prep.project['project_id']} {prep.project['name']} (슬라이드 {start + 1}~{start + len(prep.pages)}) ====="]
        lines += check_lines(prep, found)
    lines += ["", "[글꼴]"] + [f"- {f}" for f in fonts.summary()]
    check.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return TeamResult(output, check, notes, problems, summaries)
