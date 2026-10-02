from __future__ import annotations

from copy import deepcopy
from datetime import date
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from .core import ValidationError, load_json, validate_schema, wrapped_lines

REQUIRED_SHAPES = {"slide_title", "pjt_header", "author", "updated_at", "main_table", "body_top", "ms_table", "body_main"}
BLUE, BLACK = RGBColor(0, 0, 255), RGBColor(0, 0, 0)
STATUS_FILL = {"완료": "E7E7E7", "진행": "DDEBF7", "지연": "FBE2E2", "예정": "FFFFFF", "보류": "FFFFFF", "취소": "FFFFFF"}


def apply_updates(project: dict[str, Any], weekly: dict[str, Any]) -> tuple[dict[str, Any], list[str], set[tuple[str, str]]]:
    result, warnings, changed = deepcopy(project), [], set()
    milestones = {m["milestone_id"]: m for m in result["milestones"]}
    for index, update in enumerate(weekly.get("milestone_updates", [])):
        mid, field, value = update.get("milestone_id"), update.get("field"), update.get("to")
        if mid is None or mid not in milestones:
            warnings.append(f"milestone_updates[{index}]: 존재하지 않는 milestone_id {mid}; 적용하지 않음"); continue
        valid = field in {"plan_text", "note"} and isinstance(value, str)
        if field in {"plan", "actual"} and isinstance(value, str):
            try: date.fromisoformat(value); valid = True
            except ValueError: valid = False
        if field == "status": valid = value in STATUS_FILL
        if not valid:
            warnings.append(f"milestone_updates[{index}].to: {field} 값/타입 오류; 적용하지 않음"); continue
        milestones[mid][field] = value; changed.add((mid, field))
    return result, warnings, changed


def collapse_milestones(rows: list[dict[str, Any]], maximum: int = 9) -> list[dict[str, Any]]:
    rows = sorted(rows, key=lambda x: x["order"])
    while len(rows) > maximum:
        groups: dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            if row.get("parent_id") and row["status"] == "완료": groups.setdefault(row["parent_id"], []).append(row)
        candidate = next((group for group in groups.values() if len(group) >= 2), None)
        if not candidate: break
        first = candidate[0]; folded = deepcopy(first)
        folded["milestone_id"] = first["parent_id"]
        folded["name"] = f"{first['name']} 완료 {len(candidate)}개 사이트"
        folded["scope"] = "common"; folded["note"] = "접힌 완료 행"
        indexes = [rows.index(v) for v in candidate]; rows[indexes[0]] = folded
        rows = [v for i, v in enumerate(rows) if i not in indexes[1:]]
    return rows


def _scope(scope: Any, codes: dict[str, Any]) -> str:
    if scope == "common": return "공통"
    sites = {v["code"]: v.get("name") or v["code"] for v in codes["sites"]}
    lines = {v["code"]: v["label"] for v in codes["lines"]}; processes = {v["code"]: v["name"] for v in codes["processes"]}
    values = []
    for part in scope:
        if part.get("region"): values.append(part["region"])
        else: values.append(" ".join(filter(None, [sites.get(part.get("site"), part.get("site")), lines.get(part.get("line"), part.get("line"))])) + ("·" + processes.get(part.get("process"), part.get("process")) if part.get("process") else ""))
    return ", ".join(values)


def _set_text(shape, text: str, color=BLACK, size=9):
    frame = shape.text_frame; frame.clear(); frame.word_wrap = True
    paragraph = frame.paragraphs[0]; run = paragraph.add_run(); run.text = text
    run.font.name = "Arial Narrow"; run.font.size = Pt(size); run.font.color.rgb = color
    run._r.get_or_add_rPr().set("lang", "en-US")
    run._r.get_or_add_rPr().set("altLang", "ko-KR")
    frame.vertical_anchor = MSO_ANCHOR.TOP


def _cell(cell, text: str, color=BLACK, fill: str | None = None):
    _set_text(cell, text, color)
    if fill: cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(fill)


def _find(slide) -> dict[str, Any]: return {shape.name: shape for shape in slide.shapes}


def generate_ppt(root: Path, project_path: Path, weekly_path: Path, cumulative_path: Path, template: Path, output: Path) -> list[str]:
    if not template.exists():
        raise FileNotFoundError(f"공식 템플릿 누락: {template}; 필수 도형: {', '.join(sorted(REQUIRED_SHAPES))}")
    project, weekly, cumulative = load_json(project_path), load_json(weekly_path), load_json(cumulative_path)
    validate_schema(project, root / "schemas/project.schema.json"); validate_schema(weekly, root / "schemas/weekly.schema.json"); validate_schema(cumulative, root / "schemas/cumulative.schema.json")
    project, warnings, changed = apply_updates(project, weekly); codes = load_json(root / "config/code_table_site_process.json")
    prs = Presentation(template)
    if round(prs.slide_width / Inches(1), 2) != 10.83 or round(prs.slide_height / Inches(1), 2) != 7.5:
        raise ValidationError("템플릿 슬라이드 크기는 10.83 × 7.5인치여야 합니다")
    if not prs.slides: raise ValidationError("템플릿에 슬라이드가 없습니다")
    shapes = _find(prs.slides[0]); missing = REQUIRED_SHAPES - shapes.keys()
    if missing: raise ValidationError(f"템플릿 필수 도형 누락: {', '.join(sorted(missing))}")
    _set_text(shapes["slide_title"], "주간업무 추진 현황")
    _set_text(shapes["pjt_header"], project["name"]); _set_text(shapes["author"], project["owner"])
    _set_text(shapes["updated_at"], weekly["range"]["to"])
    main = shapes["main_table"].table
    _cell(main.cell(1, 0), project["name"]); _cell(main.cell(1, 1), _scope(project["target"], codes)); _cell(main.cell(0, 2), f"W{weekly['week'][-2:]}")
    _cell(main.cell(1, 2), weekly["headline"]["text"], BLUE); _cell(main.cell(1, 3), project["period"].get("target_text") or project["period"]["target"]); _cell(main.cell(1, 4), project["owner"]); _cell(main.cell(2, 2), "")
    _set_text(shapes["body_top"], f"[배경] {project['background']}\n[목표] {project['purpose']}")
    rows = collapse_milestones(project["milestones"]); table = shapes["ms_table"].table
    headers = ["단계", "적용 범위", "Baseline", "계획", "실적", "상태", "비고"]
    for col, label in enumerate(headers): _cell(table.cell(0, col), label)
    for row_index, milestone in enumerate(rows[:len(table.rows)-1], 1):
        mid = milestone["milestone_id"]; plan = milestone.get("plan_text") or milestone.get("plan") or "-"
        if not milestone.get("plan_text") and milestone.get("plan") and milestone.get("baseline") and milestone["plan"] != milestone["baseline"]:
            delay = (date.fromisoformat(milestone["plan"]) - date.fromisoformat(milestone["baseline"])).days; plan = f"{milestone['plan'][5:].replace('-', '/')} ({delay:+d})"
        values = [f"{mid[1:].replace('-', '-')}. {milestone['name']}", _scope(milestone["scope"], codes), (milestone.get("baseline") or "-")[5:].replace("-", "/"), plan, (milestone.get("actual") or "-")[5:].replace("-", "/"), milestone["status"], milestone.get("note") or ""]
        fields = ["name", "scope", "baseline", "plan_text" if milestone.get("plan_text") else "plan", "actual", "status", "note"]
        for col, value in enumerate(values): _cell(table.cell(row_index, col), str(value), BLUE if (mid, fields[col]) in changed else BLACK, STATUS_FILL[milestone["status"]] if col == 5 else None)
    body_parts = [("진행 현황(누적)", cumulative["items"], False), ("금주", weekly["progress"], True), ("향후 계획", weekly["next_plan"], True), ("이슈", weekly["issues"], True)]
    lines, total = [], 2
    for title, items, allow_blue in body_parts:
        lines.append((f"[{title}]", BLACK)); total += 1
        for value in items:
            lines.append(("• " + value["text"], BLUE if allow_blue and value.get("changed") else BLACK)); total += wrapped_lines(value["text"])
    if total > 36: warnings.append(f"본문 초과: 항목 {sum(len(x[1]) for x in body_parts)}개, 실제 {total}줄; 최대 2장 분할 필요")
    frame = shapes["body_main"].text_frame; frame.clear()
    for index, (text, color) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph(); run = paragraph.add_run(); run.text = text; run.font.name = "Arial Narrow"; run.font.size = Pt(9); run.font.color.rgb = color
    output.parent.mkdir(parents=True, exist_ok=True); prs.save(output)
    return warnings
