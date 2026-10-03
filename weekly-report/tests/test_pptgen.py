"""pptgen 테스트: 실제 v2 템플릿으로 PPT를 만든 뒤 다시 열어 값·색·글꼴·분량을 확인한다."""

import copy
import hashlib

import pytest
from pptx import Presentation

from conftest import ROOT, TEMPLATE, read, write
from weekly_report.ai import ExaoneClient
from weekly_report.codes import CodeTable
from weekly_report.ppt.milestones import apply_updates, collapse_milestones, layout_milestones
from weekly_report.ppt.render import EA_FONT, LATIN_FONT
from weekly_report.pptgen import find_template, generate_ppt

PID, WEEK = "P-ASM-001", "2026-W39"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def paths(root):
    return (root / f"data/master/projects/{PID}.json", root / f"data/derived/weekly/{PID}/{WEEK}.json",
            root / f"data/derived/cumulative/{PID}/{WEEK}.json")


def build(root, tmp_path, *, client=None, weekly=None, cumulative=None, project=None, template=TEMPLATE):
    p, w, c = paths(root)
    if weekly is not None:
        w = tmp_path / "in/weekly.json"
        write(w, weekly)
    if cumulative is not None:
        c = tmp_path / "in/cumulative.json"
        write(c, cumulative)
    if project is not None:
        p = tmp_path / "in/project.json"
        write(p, project)
    out = tmp_path / "output/out.pptx"
    notes = generate_ppt(root, p, w, c, template, out, client=client or ExaoneClient(root, "mock"))
    return out, notes


def shapes(slide):
    return {s.name: s for s in slide.shapes}


def runs(shape_or_cell):
    body = shape_or_cell.text_frame._txBody if hasattr(shape_or_cell, "text_frame") and not hasattr(shape_or_cell, "_tc") else shape_or_cell._tc.txBody
    for r in body.iter(f"{A}r"):
        rpr = r.find(f"{A}rPr")
        color = rpr.find(f"{A}solidFill/{A}srgbClr")
        yield r.findtext(f"{A}t"), rpr.get("sz"), rpr.find(f"{A}latin").get("typeface"), rpr.find(f"{A}ea").get("typeface"), color.get("val") if color is not None else None


def test_w39_example_full_slide(tmp_path):
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in paths(ROOT)}
    out, notes = build(ROOT, tmp_path)
    assert not [n for n in notes if n.startswith("PPT 검사 문제")], notes
    prs = Presentation(str(out))
    assert len(prs.slides) == 1 and round(prs.slide_width / 914400, 2) == 10.83
    s = shapes(prs.slides[0])
    main = s["main_table"].table
    assert main.cell(0, 2).text == "금주 진행사항 (W39)  (9/21~9/27)"
    assert main.cell(1, 0).text == "ESWA 재료교체 불량 개선" and main.cell(1, 3).text == "'26.09"
    assert "E77 18대" in main.cell(1, 1).text and main.cell(2, 2).text == ""
    assert s["pjt_header"].text_frame.text == "■ ESWA 재료교체 불량 개선 (1/1)"
    assert s["slide_title"].text_frame.text == "1. 과제 진행 현황_조립자동보정팀"
    assert s["updated_at"].text_frame.text == "업데이트 시간 : 9/25 18시"

    ms = s["ms_table"].table
    rows = [[ms.cell(r, c).text for c in range(7)] for r in range(1, len(ms.rows))]
    assert len(rows) == 9  # 9개 마일스톤 모두 표시
    assert rows[0] == ["1. 현황 분석", "공통", "07/24", "07/24", "07/24", "완료", ""]
    assert rows[7] == ["6-3. 수평전개", "Normal·조립", "09/25", "09/30 (+5)", "–", "진행", "WA·MI_HL Normal Line"]
    assert rows[8][:4] == ["6-4. 수평전개", "북미·조립", "09/25", "10월초"]
    fill = lambda r: ms.cell(r, 5)._tc.find(f"{A}tcPr/{A}solidFill/{A}srgbClr").get("val")
    assert fill(1) == "E7E7E7" and fill(8) == "DDEBF7" and fill(9) == "FFFFFF"

    blue = [t for name in ("body_main",) for t, *_, color in runs(s[name]) if color == "0000FF"]
    blue += [t for r in range(1, len(ms.rows)) for c in range(7) for t, *_, color in runs(ms.cell(r, c)) if color == "0000FF"]
    blue += [t for t, *_, color in runs(main.cell(1, 2)) if color == "0000FF"]
    assert sorted(blue) == sorted([
        " - ESWA(MEB E77 전호기) 0.171% → 0.144% (적용 전·후 4일, 단기)",
        " - ESMI1 #2-2·3 0.367% → 0.230% (단기, 장기 모니터링 필요)",
        " - Normal Line(WA·MI_HL) 수평전개 진행 중",
        " - Normal Line 수평전개 완료(~9/30), 북미 Site 수평전개(~10월초)",
        " - 효과 수치가 4일 단기 기준 → 장기 모니터링 결과로 재확인 필요",
        "09/30 (+5)", "WA·MI_HL Normal Line", "10월초", "과제 기한('26.09) 초과",
        "금주(W39)에는 로직 개선 효과 모니터링 – ESWA 0.144%, ESMI1 0.230%로 단기 개선 확인",
    ])
    body_text = s["body_main"].text_frame.text
    assert "1차 적용 E77 위치 불량률 0.189% → 0.164%" in body_text  # 누적 요약에 빠진 고정 사실 보충
    for name in ("body_top", "body_main"):
        for text, size, latin, ea, color in runs(s[name]):
            assert (size, latin, ea) == ("900", LATIN_FONT, EA_FONT) and color in {"000000", "0000FF"}
    assert before == {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in paths(ROOT)}


def test_missing_template_reports_required_shapes(tmp_path):
    with pytest.raises(FileNotFoundError, match="main_table"):
        build(ROOT, tmp_path, template=tmp_path / "missing.pptx")
    with pytest.raises(FileNotFoundError, match="body_main"):
        find_template(tmp_path)


def test_template_missing_shape_is_reported(tmp_path):
    prs = Presentation(str(TEMPLATE))
    s = shapes(prs.slides[0])
    s["body_main"]._element.getparent().remove(s["body_main"]._element)
    broken = tmp_path / "broken.pptx"
    prs.save(str(broken))
    with pytest.raises(ValueError, match="필수 도형 누락: body_main"):
        build(ROOT, tmp_path, template=broken)


def test_milestone_updates_apply_only_to_memory_copy():
    project = read(ROOT / "data/master/projects/P-ASM-001.json")
    original = copy.deepcopy(project)
    weekly = {"milestone_updates": [
        {"milestone_id": "M6-3", "field": "plan", "to": "2026-10-02", "source_ids": []},
        {"milestone_id": None, "field": "plan", "to": "2026-10-01", "source_ids": []},
        {"milestone_id": "M6-3", "field": "status", "to": 3, "source_ids": []},
        {"milestone_id": "M6-3", "field": "actual", "to": "10/02", "source_ids": []},
    ]}
    result, warnings, changed = apply_updates(project, weekly)
    assert project == original and len(warnings) == 3 and changed == {("M6-3", "plan")}
    m = next(x for x in result["milestones"] if x["milestone_id"] == "M6-3")
    assert m["plan"] == "2026-10-02" and m["baseline"] == "2026-09-25"


def test_completed_children_fold_and_unfoldable_rows_carry_over():
    codes = CodeTable.load(ROOT)
    rows = [{"milestone_id": f"M6-{i}", "parent_id": "M6", "status": "완료", "order": i, "name": "수평전개",
             "scope": [{"site": "WA", "process": "AS"}], "baseline": "2026-09-25", "plan": "2026-09-25", "plan_text": None,
             "actual": f"2026-09-{10 + i:02d}", "note": None} for i in range(1, 11)]
    folded = collapse_milestones(rows, codes=codes)
    assert len(folded) == 1 and folded[0]["name"] == "수평전개 완료 10개 사이트" and folded[0]["actual"] == "2026-09-20"
    for row in rows:
        row["status"] = "진행"
    first, rest = layout_milestones(rows + rows[:2], codes=codes)
    assert len(first) == 9 and len(rest) == 3  # 행을 버리지 않고 (계속)으로 넘김


def _long_weekly():
    weekly = read(ROOT / "data/derived/weekly/P-ASM-001/2026-W39.json")
    extra = [{"text": f"추가 진행 항목 {i}번 – ESWA 0.144% 유지 확인 및 후속 데이터 수집과 정리 작업 진행 중", "source_ids": ["D-260922-ljh-01"],
              "kind": "fact", "changed": i % 2 == 0} for i in range(1, 9)]
    weekly["progress"] += extra
    return weekly


def test_overflow_goes_to_continued_slide_with_max_two(tmp_path):
    out, notes = build(ROOT, tmp_path, weekly=_long_weekly())
    assert not [n for n in notes if n.startswith("PPT 검사 문제")], notes
    assert any("fit_to_budget 미적용" in n for n in notes)  # mock 없음 → 원문 유지
    prs = Presentation(str(out))
    assert len(prs.slides) == 2
    second = shapes(prs.slides[1])
    assert second["pjt_header"].text_frame.text.endswith("(계속)")
    assert "ms_table" not in second  # 이월할 마일스톤 행이 없으면 표를 두지 않음
    body2 = second["body_main"].text_frame.text
    assert "(계속)" in body2 and "추가 진행 항목" in body2
    # 항목 수와 줄 수는 따로 계산: 첫 장 진행 현황은 한도 7개 이하
    first_body = shapes(prs.slides[0])["body_main"].text_frame.text
    assert first_body.count("추가 진행 항목") <= 4


def test_three_slides_needed_is_an_error(tmp_path):
    weekly = _long_weekly()
    weekly["progress"] += weekly["progress"] * 4
    with pytest.raises(ValueError, match="최대 2장"):
        build(ROOT, tmp_path, weekly=weekly)


def test_fit_to_budget_keeps_changed_and_sources(repo, tmp_path):
    weekly = _long_weekly()
    write(repo / "data/derived/weekly/P-ASM-001/2026-W39.json", weekly)
    progress = weekly["progress"]
    fit = {"items": [{"text": p["text"].replace("및 후속 데이터 수집과 정리 작업 ", ""), "source_ids": p["source_ids"]} for p in progress[:7]],
           "dropped": [{"text": p["text"], "source_ids": p["source_ids"]} for p in progress[7:]]}
    write(repo / "prompts/mock_responses/fit_to_budget__P-ASM-001__2026-W39__progress.json", fit)
    out, notes = build(repo, tmp_path)
    assert any("progress: fit_to_budget 적용" in n for n in notes), notes
    assert not [n for n in notes if n.startswith("PPT 검사 문제")], notes
    prs = Presentation(str(out))
    body = shapes(prs.slides[0])["body_main"]
    colors = {t: c for t, *_, c in runs(body)}
    # 원래 changed=false였던 '추가 진행 항목 1번'은 줄인 뒤에도 검정, changed=true인 2번은 파랑
    one = next(t for t in colors if "추가 진행 항목 1번" in t)
    two = next(t for t in colors if "추가 진행 항목 2번" in t)
    assert colors[one] == "000000" and colors[two] == "0000FF"
    assert len(prs.slides) == 2  # dropped는 (계속) 장으로


def test_fit_to_budget_result_with_new_number_is_rejected(repo, tmp_path):
    weekly = _long_weekly()
    write(repo / "data/derived/weekly/P-ASM-001/2026-W39.json", weekly)
    fit = {"items": [{"text": "ESWA 0.999% 개선", "source_ids": ["D-260922-ljh-01"]}], "dropped": []}
    write(repo / "prompts/mock_responses/fit_to_budget__P-ASM-001__2026-W39__progress.json", fit)
    out, notes = build(repo, tmp_path)
    assert any("원본에 없는 수치" in n for n in notes)
    assert "0.999" not in "".join(s.text_frame.text for s in Presentation(str(out)).slides[0].shapes if s.has_text_frame)


def test_no_issue_shows_default_text(tmp_path):
    weekly = read(ROOT / "data/derived/weekly/P-ASM-001/2026-W39.json")
    weekly["issues"] = []
    out, _ = build(ROOT, tmp_path, weekly=weekly)
    assert "특이사항 없음" in shapes(Presentation(str(out)).slides[0])["body_main"].text_frame.text


def test_project_week_mismatch_is_rejected(tmp_path):
    cumulative = read(ROOT / "data/derived/cumulative/P-ASM-001/2026-W39.json")
    cumulative["as_of_week"] = "2026-W38"
    with pytest.raises(ValueError, match="주차 불일치"):
        build(ROOT, tmp_path, cumulative=cumulative)


def test_milestone_rows_beyond_nine_go_to_continued_slide(tmp_path):
    project = read(ROOT / "data/master/projects/P-ASM-001.json")
    base = project["milestones"][-1]
    for i in range(5, 9):
        extra = copy.deepcopy(base)
        extra.update({"milestone_id": f"M6-{i}", "order": 4 + i, "scope": [{"site": "HD", "process": "AS"}]})
        project["milestones"].append(extra)
    out, notes = build(ROOT, tmp_path, project=project)
    assert not [n for n in notes if n.startswith("PPT 검사 문제")], notes
    assert any("완료 하위 행 접기: M6-1, M6-2" in n for n in notes)
    prs = Presentation(str(out))
    assert len(prs.slides) == 2
    first = shapes(prs.slides[0])["ms_table"].table
    second = shapes(prs.slides[1])["ms_table"].table
    assert len(first.rows) - 1 == 9 and len(second.rows) - 1 == 3
    assert first.cell(6, 0).text == "6. 수평전개 완료 2개 사이트"
