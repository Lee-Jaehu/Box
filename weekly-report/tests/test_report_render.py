"""보고 자료 렌더러: 템플릿 초안, 경영진 1장 요약, 월간 종합 보고, 칸 단위 대체, 웹 흐름."""

import importlib.util
import json
import shutil
from datetime import date
from pathlib import Path

import pytest
from pptx import Presentation

from conftest import ROOT, read
from weekly_report import workbench as wb
from weekly_report.ai import ExaoneClient
from weekly_report.fonts import EA_REGULAR
from weekly_report.report.generate import generate_exec_summary, generate_monthly
from weekly_report.report.render import SLIDE_NAMES

EXEC_EXAMPLE = ROOT / "docs/보고자료_예시/report_exec_summary__예시응답.json"
DIRS = [ROOT, ROOT / "demo/w40/out"]
TODAY = date(2026, 10, 4)
WORKLOG_SAMPLE = "00000000-0000-4000-8000-000000000001"  # data/worklog/WorkLog-sample


def load_demo():
    spec = importlib.util.spec_from_file_location("report_demo", ROOT / "demo/report/run_report_demo.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def mock_client(root: Path, mock_dir: Path, name: str, payload: dict) -> ExaoneClient:
    mock_dir.mkdir(parents=True, exist_ok=True)
    (mock_dir / name).write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return ExaoneClient(root, mock_dir=mock_dir)


def texts(slide) -> dict[str, str]:
    return {s.name: s.text_frame.text for s in slide.shapes if s.has_text_frame}


def test_template_draft_has_named_slots(tmp_path):
    spec = importlib.util.spec_from_file_location("tpl", ROOT / "tools/make_report_template.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    out = module.build(output=tmp_path / "t.pptx")
    prs = Presentation(str(out))
    assert tuple(s.name for s in prs.slides) == SLIDE_NAMES
    names = {s.name for s in prs.slides[0].shapes}
    assert {"title", "head_message", "background_conclusion", "left_items", "right_items", "kpi_table", "page_no"} <= names
    assert {"project_table", "highlights", "risks"} <= {s.name for s in prs.slides[1].shapes}
    eas = {ea.get("typeface") for s in prs.slides for ea in s._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")}
    assert eas == {EA_REGULAR}
    # 저장소의 템플릿 파일은 스크립트 결과와 같은 구성이어야 한다
    repo = Presentation(str(ROOT / "보고자료_Template_v1_초안.pptx"))
    assert [sorted(s.name for s in sl.shapes) for sl in repo.slides] == [sorted(s.name for s in sl.shapes) for sl in prs.slides]


def test_exec_summary_one_slide_with_emphasis_and_kpi_color(tmp_path):
    client = mock_client(ROOT, tmp_path / "mock", "report_exec_summary__P-ASM-001__2026-W40.json", read(EXEC_EXAMPLE))
    result = generate_exec_summary(ROOT, "P-ASM-001", "2026-W40", tmp_path / "exec.pptx", client=client, dirs=DIRS, today=TODAY)
    assert result["problems"] == [] and result["notes"] == ["대체 없음 (AI 응답 그대로 사용)"]
    prs = Presentation(str(result["pptx"]))
    assert len(prs.slides) == 1
    t = texts(prs.slides[0])
    assert t["title"] == "ESWA 재료교체 불량 개선 진행 결과 및 향후 계획" and t["page_no"] == "1 / 1"
    assert t["background_conclusion"].count("•") == 3 and "10/13 (+18일)" in t["schedule_note"]
    xml = prs.slides[0]._element.xml
    assert xml.count("<a:highlight>") == 3  # 강조 구절 3개
    kpi = next(s for s in prs.slides[0].shapes if s.name == "kpi_table").table
    assert kpi.cell(1, 4).text == "Δ0.045%p" and 'val="0000CC"' in kpi.cell(1, 4)._tc.xml  # 감소 = 개선 → 파랑


def test_bad_ai_slots_are_replaced_and_reported(tmp_path):
    payload = read(EXEC_EXAMPLE)
    payload["head_message"]["text"] = "재료교체 불량률 개선 효과 유지"  # 경어체 아님
    payload["left_items"][1]["text"] = "- 누적 적용 27대"  # 입력에 없는 수치
    payload["emphasis"][2] = "없는 구절"  # 문장에 없는 강조 구절 → 그 구절만 제외
    client = mock_client(ROOT, tmp_path / "mock", "report_exec_summary__P-ASM-001__2026-W40.json", payload)
    result = generate_exec_summary(ROOT, "P-ASM-001", "2026-W40", tmp_path / "exec.pptx", client=client, dirs=DIRS, today=TODAY)
    notes = "\n".join(result["notes"])
    assert "head_message" in notes and "27" in notes and "없는 구절" in notes
    t = texts(Presentation(str(result["pptx"])).slides[0])
    assert t["head_message"] == "ESWA 재료교체 불량 개선 과제의 진행 현황을 보고드립니다."
    assert "27대" not in t["left_items"] and "0.270% → 0.185%" in t["left_items"]  # 그 항목만 빠지고 나머지는 유지
    assert "[대체]" in result["check"].read_text(encoding="utf-8")


def test_monthly_demo_six_projects(tmp_path):
    demo = load_demo()
    root = demo.build_root(tmp_path)
    client = ExaoneClient(root, mock_dir=ROOT / "demo/report/mock_responses")
    ids = ["P-ASM-001", *demo.FIXTURE_PROJECTS]
    result = generate_monthly(root, ids, 2026, 9, tmp_path / "m.pptx", client=client, today=TODAY)
    assert result["problems"] == [] and result["notes"] == ["대체 없음 (AI 응답 그대로 사용)"]
    prs = Presentation(str(result["pptx"]))
    assert len(prs.slides) == 1  # 요청 없음 → 의사결정 장 생략
    table = next(s for s in prs.slides[0].shapes if s.name == "project_table").table
    assert len(table.rows) == 7 and table.cell(1, 0).text == "ESWA 재료교체 불량 개선"
    assert "09/30 (+5일)" in table.cell(1, 3).text and 'val="C00000"' in table.cell(1, 3)._tc.xml


def test_monthly_splits_long_table_and_adds_request_slide(tmp_path):
    demo = load_demo()
    root = demo.build_root(tmp_path)
    ids = ["P-ASM-001", *demo.FIXTURE_PROJECTS]
    payload = read(ROOT / "demo/report/mock_responses/report_monthly__ALL__2026-09.json")
    payload["requests"] = [{"text": "WA 적용 9/30 완료를 위한 설비 일정 협조", "due": "9/30", "dept": None, "source_ids": ["D-260922-tester-02"]}]
    client = mock_client(root, tmp_path / "mock", "report_monthly__ALL__2026-09.json", payload)
    result = generate_monthly(root, ids * 2, 2026, 9, tmp_path / "m.pptx", client=client, today=TODAY)
    prs = Presentation(str(result["pptx"]))
    titles = [texts(s)["title"] for s in prs.slides]
    assert len(prs.slides) >= 3 and titles[1].endswith("(계속)") and titles[-1].startswith("의사결정 및 업무협조 요청")
    assert [texts(s)["page_no"] for s in prs.slides][-1] == f"{len(prs.slides)} / {len(prs.slides)}"
    requests = next(s for s in prs.slides[-1].shapes if s.name == "requests_table").table
    assert requests.cell(1, 0).text == "[합성 예시] 검증 로직 수평전개" and requests.cell(1, 2).text == "9/30"
    assert result["problems"] == []


@pytest.fixture
def ws(tmp_path):
    wb.init_workspace(tmp_path / "ws")
    return tmp_path / "ws"


def test_web_report_flow(ws):
    assert wb.run(ws, "P-ASM-001", "2026-W40")["status"] == "ok"
    exec_result = wb.run_report(ws, "exec", "P-ASM-001", "2026-W40")  # 데모 응답이 작업공간에 있음
    assert exec_result["status"] == "ok" and exec_result["files"]["pptx"].endswith("경영진요약_P-ASM-001_2026-W40.pptx")
    need = wb.run_report(ws, "monthly", "P-ASM-001", "2026-W40")
    assert need["status"] == "need_response" and need["response_name"] == "report_monthly__ALL__2026-10.json"
    assert '"head_message": {...}' in need["format"] and "’26.10월" in need["prompt"]
    with pytest.raises(wb.WorkbenchError, match="head_message"):
        wb.save_response(ws, need["response_name"], json.dumps({"summary": "x"}))
    answer = {"head_message": {"text": "10월 재료교체 과제는 Normal Line 수평전개를 완료했습니다", "source_ids": ["D-260930-khw-01"]},
              "project_comments": [{"project_id": "P-ASM-001", "text": "Normal Line 수평전개 完(9/30)", "source_ids": ["D-260930-khw-01"]},
                                   {"project_id": WORKLOG_SAMPLE, "text": "업무기록 시스템 개발 진행", "source_ids": [WORKLOG_SAMPLE]}],  # 작업공간의 WorkLog 예시 과제
              "highlights": [], "risks": [], "requests": []}
    wb.save_response(ws, need["response_name"], "결과입니다\n```json\n" + json.dumps(answer, ensure_ascii=False) + "\n```")
    done = wb.run_report(ws, "monthly", "P-ASM-001", "2026-W40")
    assert done["status"] == "ok" and "대체 없음" in done["check"], done.get("check")
