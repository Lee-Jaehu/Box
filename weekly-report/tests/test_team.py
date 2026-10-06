"""팀 주간보고: 과제 요약 JSON(EXAONE → Rule 색·순서·검증) + 팀 요약 페이지 + 과제별 주간 장표 묶음."""

import importlib.util
import json
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from conftest import ROOT, TEMPLATE, read, write
from weekly_report import workbench as wb
from weekly_report.ai import ExaoneClient
from weekly_report.core import ValidationError
from weekly_report.ppt.render import drop_slide, shape_map
from weekly_report.summary import build_items
from weekly_report.team import FONT_SIZES, SUMMARY_BLUE, Block, SummaryGeometry, generate_team_deck, layout_summary

WORKLOG_SAMPLE = "00000000-0000-4000-8000-000000000001"


def load_demo():
    spec = importlib.util.spec_from_file_location("team_demo", ROOT / "demo/team/run_team_demo.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def demo_decks(tmp_path_factory):
    """데모 root(조립자동보정팀 W40 + 검증팀 W39)로 두 묶음 PPT를 만든다."""
    demo = load_demo()
    tmp = tmp_path_factory.mktemp("team")
    root = demo.build_root(tmp)
    out = tmp / "out"
    client = ExaoneClient(root, "mock")
    decks = {team: generate_team_deck(root, team, week, out / f"{team}.pptx", out_root=out, client=client)
             for team, week in demo.RUNS}
    return root, out, decks


def blue_texts(shape) -> list[str]:
    return [r.text for p in shape.text_frame.paragraphs for r in p.runs
            if r.font.color and r.font.color.type is not None and str(r.font.color.rgb) == SUMMARY_BLUE]


def test_deck_order_summary_then_weekly_pages(demo_decks):
    _, _, decks = demo_decks
    result = decks["검증팀"]
    assert result.problems == [], result.problems
    slides = list(Presentation(str(result.pptx)).slides)
    assert len(slides) == 6  # 요약 1장 + 과제 5건 × 1장
    first = shape_map(slides[0])
    assert "main_table" not in first and first["slide_title"].text_frame.text == "1. 검증팀 (1/1)"
    body = first["pjt_header"].text_frame.text
    for number, name in enumerate(["자동보정 로직 개선", "수율 개선 분석", "설비 투자", "검증 로직 수평전개", "대시보드 시스템 구축"], 1):
        assert f"{number}. [합성 예시] {name}" in body  # 굵은 "n. 과제명", 카테고리 이름은 표시하지 않음
    assert "배경" not in body and "진행 현황" not in body
    for index, slide in enumerate(slides[1:], 1):
        shapes = shape_map(slide)
        assert "main_table" in shapes and shapes["pjt_header"].text_frame.text.endswith(f"({index}/5)")
    # 템플릿의 참고 장(실제 작성본·주간 예시)은 들어가지 않는다
    all_text = " ".join(s.text_frame.text for sl in slides for s in sl.shapes if s.has_text_frame)
    assert "자동보정선행개발팀" not in all_text and "과제명 (ESWA" not in all_text


def test_blue_marks_only_this_week_sources(demo_decks):
    _, out, decks = demo_decks
    summary = read(out / "data/derived/summary/P-ASM-001/2026-W40.json")
    assert summary["items"][0]["category"] == "background" and summary["items"][0]["new"] is False
    old = [i for i in summary["items"] if i["category"] == "progress" and not i["new"]]
    assert old and all(not set(i["source_ids"]) & set(summary["week_source_ids"]) for i in old)
    header = shape_map(Presentation(str(decks["조립자동보정팀"].pptx)).slides[0])["pjt_header"]
    want = [f"- {i['text']}" for i in summary["items"] if i["new"]] + [f". {d['text']}" for i in summary["items"] for d in i["details"] if d["new"]]
    assert sorted(blue_texts(header)) == sorted(want)
    sizes = {r.font.size.pt for p in header.text_frame.paragraphs for r in p.runs}
    assert sizes == {11.0}  # 1과제는 11pt로 충분
    bold = [r.text for p in header.text_frame.paragraphs for r in p.runs if r.font.bold]
    assert bold == ["1. ESWA 재료교체 불량 개선"]


def test_summary_check_report_lists_issues(demo_decks):
    _, out, decks = demo_decks
    check = (out / "output/P-ASM-001/summary_check_2026-W40.txt").read_text(encoding="utf-8")
    assert "■ [잘한점]" in check and "□ [배경/목적]" in check and "오류 0건" in check
    deck_check = decks["조립자동보정팀"].check.read_text(encoding="utf-8")
    assert "[요약 페이지 재검사]" in deck_check and "- 통과" in deck_check


def test_build_items_rule_order_color_and_background_fallback():
    project = {"project_id": "P-X", "background": "배경 문장입니다.", "purpose": "목적 문장입니다."}
    payload = {"items": [
        {"category": "plan", "text": "계획", "source_ids": ["D-261001-a-01"]},
        {"category": "good", "text": "잘한점", "source_ids": ["D-261001-a-01"], "details": [{"text": "세부", "source_ids": []}]},
        {"category": "progress", "text": "(D-260901-a-01) 기존 진행", "source_ids": ["D-260901-a-01"]},
        {"category": "issue", "text": "이슈", "source_ids": ["D-261001-a-01", "D-260901-a-01"]},
        {"category": "risk", "text": "알 수 없는 칸", "source_ids": []},
    ]}
    notes = []
    items = build_items(payload, project, {"D-261001-a-01"}, notes)
    assert [i["category"] for i in items] == ["background", "progress", "good", "issue", "plan"]
    assert items[0]["text"] == "배경 문장입니다. 목적 문장입니다." and items[0]["new"] is False
    assert items[1]["text"] == "기존 진행" and items[1]["new"] is False  # 앞에 붙은 기록 ID 제거, 이전 기록 → 검정
    assert items[2]["new"] and items[2]["details"][0]["new"]  # 근거 없는 세부는 항목 색을 따른다
    assert items[3]["new"]
    messages = " ".join(n.format() for n in notes)
    assert "알 수 없는 카테고리" in messages and "배경(background)이 없어" in messages and "함께 근거" in messages


def test_unsupported_number_and_style_are_reported(repo, tmp_path):
    mock_dir = tmp_path / "mocks"
    shutil.copytree(ROOT / "prompts/mock_responses", mock_dir)
    write(mock_dir / "project_summary__P-ASM-001__2026-W39.json", {"items": [
        {"category": "background", "text": "재료교체 Y축 불량을 보정 로직으로 개선하는 과제입니다.", "source_ids": ["P-ASM-001"]},
        {"category": "progress", "text": "ESWA 불량률이 0.171% → 0.120%로 감소했습니다.(9/22)", "source_ids": ["D-260922-ljh-01"]},
        {"category": "issue", "text": "장기 모니터링이 필요함", "source_ids": ["D-260922-ljh-01"]},
    ]})
    result = generate_team_deck(repo, "P-ASM-001", "2026-W39", tmp_path / "out/deck.pptx", out_root=tmp_path / "out",
                                client=ExaoneClient(repo, mock_dir=mock_dir))
    issues = " ".join(i.format() for s in result.summaries for i in s.issues)
    assert "입력에서 확인되지 않는 수치 0.120%" in issues
    assert '경어체 종결' in issues and '진행 날짜 "(M/D)" 없음' in issues
    assert result.problems == []  # 의미 검증 문제는 저장·PPT 생성을 막지 않는다 (보고서에 표시)
    assert (tmp_path / "out/data/derived/summary/P-ASM-001/2026-W39.json").exists()


def test_layout_shrinks_font_then_paginates_by_project():
    geom = SummaryGeometry(top=int(0.94 * 914400), bottom=int(7.10 * 914400), width=int(10.56 * 914400), hangul_em=0.891, line_factor=1.17)
    item = {"text": "가" * 60 + "했습니다.(9/22)", "new": True, "details": []}

    def blocks(count, per):
        return [Block(n, f"과제 {n}", [dict(item) for _ in range(per)]) for n in range(1, count + 1)]

    pages, notes = layout_summary(blocks(5, 4), geom)  # 5 × (제목 1 + 4) + 빈 줄 4 = 29줄 → 11pt
    assert len(pages) == 1 and pages[0].size == 11.0
    pages, notes = layout_summary(blocks(5, 5), geom)  # 34줄 → 11pt 한도(34) 안
    assert pages[0].size == 11.0
    pages, notes = layout_summary(blocks(6, 5), geom)  # 41줄 → 10pt도 넘침 → 11pt 과제 단위 2장
    assert len(pages) == 2 and all(p.size == FONT_SIZES[0] for p in pages)
    assert [b.number for b in pages[1].blocks] == [6] and "2장" in notes[0]
    pages, _ = layout_summary(blocks(6, 4), geom)  # 6 × 5 + 5 = 35줄 → 10.5pt 1장
    assert len(pages) == 1 and pages[0].size == 10.5


def test_many_projects_split_summary_pages(demo_decks, tmp_path):
    """요약이 한 장을 넘으면 과제 단위로 "(1/2)", "(2/2)"로 나눈다."""
    root, _, _ = demo_decks
    mock_dir = tmp_path / "mocks"
    shutil.copytree(root / "prompts/mock_responses", mock_dir)
    long = "보고 자동화 검증을 위해 실제 업무 데이터 없이 파이프라인을 확인하고 결과를 정리했습니다.(9/22)"
    for pid, daily in (("P-APC-101", "01"), ("P-ROL-102", "02"), ("P-INV-103", "03"), ("P-SYS-104", "04"), ("P-DAT-105", "05")):
        source = [f"D-260922-tester-{daily}"]
        items = [{"category": "progress", "text": long, "source_ids": source,
                  "details": [{"text": "세부 내용을 자세히 설명하기 위한 문장입니다.", "source_ids": source}] * 2}] * 2
        write(mock_dir / f"project_summary__{pid}__2026-W39.json", {"items": items})
    out = tmp_path / "out"
    result = generate_team_deck(root, "검증팀", "2026-W39", out / "deck.pptx", out_root=out, client=ExaoneClient(root, mock_dir=mock_dir))
    slides = list(Presentation(str(result.pptx)).slides)
    titles = [shape_map(s)["slide_title"].text_frame.text for s in slides if "main_table" not in shape_map(s)]
    assert titles == ["1. 검증팀 (1/2)", "1. 검증팀 (2/2)"]
    assert len(slides) == 2 + 5 and result.problems == [], result.problems


def test_worklog_project_team_deck(repo, tmp_path):
    out = tmp_path / "out"
    shutil.copytree(ROOT / "demo/worklog/out/data", out / "data")
    weekly = read(out / f"data/derived/weekly/{WORKLOG_SAMPLE}/2026-W40.json")
    ids = weekly["source_daily_ids"]
    mock_dir = tmp_path / "mocks"
    write(mock_dir / f"project_summary__{WORKLOG_SAMPLE}__2026-W40.json", {"items": [
        {"category": "progress", "text": "업무기록 시스템 화면과 저장 기능 개발을 진행했습니다.", "source_ids": ids[:1]}]})
    result = generate_team_deck(repo, WORKLOG_SAMPLE, "2026-W40", out / "deck.pptx", out_root=out,
                                client=ExaoneClient(repo, mock_dir=mock_dir))
    slides = list(Presentation(str(result.pptx)).slides)
    assert shape_map(slides[0])["slide_title"].text_frame.text == "1. 예시DX팀 (1/1)"
    assert len(slides) >= 2 and result.problems == [], result.problems


def test_template_without_summary_slide_is_reported(repo, tmp_path):
    prs = Presentation(str(TEMPLATE))
    drop_slide(prs, prs.slides[0])
    template = tmp_path / "no_summary.pptx"
    prs.save(str(template))
    with pytest.raises(ValidationError, match="팀 요약 장표"):
        generate_team_deck(repo, "P-ASM-001", "2026-W39", tmp_path / "out.pptx", out_root=tmp_path, template=template)


def test_web_team_report_asks_then_completes(tmp_path):
    ws = tmp_path / "ws"
    wb.init_workspace(ws)
    assert wb.run(ws, "P-ASM-001", "2026-W40")["status"] == "ok"
    done = wb.run_report(ws, "team", "P-ASM-001", "2026-W40")  # 데모 응답이 작업공간에 있음
    assert done["status"] == "ok" and done["files"]["pptx"].endswith("팀주간보고_조립자동보정팀_2026-W40.pptx"), done
    answer = ws / "prompts/mock_responses/project_summary__P-ASM-001__2026-W40.json"
    saved = answer.read_text(encoding="utf-8")
    answer.unlink()
    need = wb.run_report(ws, "team", "P-ASM-001", "2026-W40")
    assert need["status"] == "need_response" and need["prompt_id"] == "project_summary"
    assert need["response_name"] == "project_summary__P-ASM-001__2026-W40.json" and "이 과제에 쓸 수 있는 줄 수" in need["prompt"]
    wb.save_response(ws, need["response_name"], "응답입니다\n```json\n" + saved + "\n```")
    again = wb.run_report(ws, "team", "P-ASM-001", "2026-W40")
    assert again["status"] == "ok" and json.loads(saved)["items"]
