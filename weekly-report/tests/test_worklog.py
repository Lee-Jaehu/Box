"""WorkLog export 연동: 변환기, 입력 읽기, 주간 파이프라인, 마일스톤 변경 표시, 보고 자료."""

import json
import shutil
from datetime import date
from pathlib import Path

import pytest
from pptx import Presentation

from conftest import ROOT, TEMPLATE, read, write
from weekly_report import sources
from weekly_report.ai import ExaoneClient
from weekly_report.core import validate_schema
from weekly_report.pptgen import generate_ppt
from weekly_report.validate import id_dates
from weekly_report.weekly import run_weekly
from weekly_report.worklog import WorklogError, load_mapping, tiptap_inline, tiptap_lines, to_dailies, to_project

SAMPLE = ROOT / "data/worklog/WorkLog-sample"
PID = "00000000-0000-4000-8000-000000000001"
DEMO_MOCK = ROOT / "demo/worklog/mock_responses"


def doc(*nodes):
    return {"documentVersion": 1, "format": "tiptap-json", "doc": {"type": "doc", "content": list(nodes)}}


def para(*texts):
    return {"type": "paragraph", "content": [{"type": "text", "text": t} if t != "\n" else {"type": "hardBreak"} for t in texts]}


def items(kind, *texts, **attrs):
    return {"type": kind, **({"attrs": attrs} if attrs else {}), "content": [{"type": "listItem", "content": [para(t)]} for t in texts]}


def test_tiptap_lists_paragraphs_and_breaks():
    value = doc(items("bulletList", "A 발생", "B 부족"), {"type": "paragraph"}, items("orderedList", "C 축적", "D 해소", start=1),
                para("첫 줄", "\n", "둘째 줄"))
    assert tiptap_lines(value) == ["- A 발생", "- B 부족", "1. C 축적", "2. D 해소", "첫 줄", "둘째 줄"]
    assert tiptap_inline(doc(items("orderedList", "C 축적", "D 해소", start=1))) == "① C 축적 ② D 해소"
    assert tiptap_inline(doc(items("bulletList", "A 발생", "B 부족"))) == "A 발생 / B 부족"
    assert tiptap_lines(None) == [] and tiptap_lines("문자열\n그대로") == ["문자열", "그대로"]


def test_project_conversion_follows_decisions():
    notes = []
    project = to_project(read(SAMPLE / "project.json"), load_mapping(ROOT), notes)
    validate_schema(project, ROOT / "schemas/project.schema.json")
    assert project["project_id"] == PID and project["status"] == "예정" and project["health"] == "정상"
    assert project["org"] == {"group": "예시DX담당", "dept": "예시DX담당", "team": "예시DX팀"}
    assert project["target_label"] == "예시DX팀"  # 결정: 대상 = 팀명
    names = [m["name"] for m in project["milestones"]]
    assert names == ["요건 분석", "설계", "Work log 개발", "PPT Generator 개발"]  # 일반 업무·삭제 단계 제외
    design = project["milestones"][1]
    assert design["scope_label"] == "10/01~10/04" and design["status"] == "진행"  # 결정: 적용 범위 = 계획 기간
    assert project["milestones"][2]["baseline"] is None and project["milestones"][0]["actual"] == "2026-10-01"
    assert project["people"]["00000000-0000-4000-8000-000000000010"]["name"] == "김예시"
    assert project["purpose"].startswith("① ") and any("일반 업무" in n for n in notes)


def test_overdue_milestone_and_health_are_computed():
    export = read(SAMPLE / "project.json")
    export["generatedAt"] = "2026-10-07T09:00:00+09:00"  # 설계 계획 종료(10/04)가 지난 시점
    export["milestones"][3]["plannedEnd"] = "2026-10-08"  # Baseline 없음 → 주의 판단에는 쓰지 않음
    project = to_project(export, load_mapping(ROOT))
    assert project["milestones"][1]["status"] == "지연" and project["health"] == "지연"
    export["project"]["status"] = "unknown_state"
    notes = []
    assert to_project(export, load_mapping(ROOT), notes)["status"] == "진행" and any("unknown_state" in n for n in notes)
    export["project"]["deletedAt"] = "2026-10-07T09:00:00+09:00"
    with pytest.raises(WorklogError, match="삭제된 과제"):
        to_project(export, load_mapping(ROOT))


def test_daily_conversion_keeps_structure_and_stable_ids():
    mapping = load_mapping(ROOT)
    project = to_project(read(SAMPLE / "project.json"), mapping)
    first = to_dailies(read(SAMPLE / "2026-10-04.json"), project, mapping)
    again = to_dailies(read(SAMPLE / "2026-10-04.json"), project, mapping)
    assert [d["daily_id"] for d in first] == [d["daily_id"] for d in again] == ["D-261004-00000010-01"]
    daily = first[0]
    validate_schema(daily, ROOT / "schemas/daily.schema.json")
    text = daily["raw_text"]
    assert "[수행] (설계) AI 기반 프롬프트 구체화 (10/04)" in text and "  1. FAST API, node js 기반 Web Service 구축" in text
    assert "[이슈] 업무용 컴퓨터 Codex 사용 불가 (상태 미해결)" in text and "[배운 점] Codex는 업무용 노트북으로 할 수 없다." in text
    assert daily["author_name"] == "김예시" and len(daily["source_refs"]["records"]) == 4
    assert id_dates({daily["daily_id"]}) == ["10/4"]  # 문장 끝 진행 날짜의 근거


@pytest.fixture
def wl_root(tmp_path) -> Path:
    root = tmp_path / "root"
    for name in ("config", "schemas", "prompts"):
        shutil.copytree(ROOT / name, root / name)
    shutil.copytree(SAMPLE, root / "data/worklog/WorkLog-sample")
    shutil.copy(TEMPLATE, root / TEMPLATE.name)
    for path in ROOT.glob("LGSM*.[tT][tT][fF]"):
        shutil.copy(path, root / path.name)
    return root


def test_sources_list_and_select_worklog_dailies(wl_root):
    projects = sources.list_projects(wl_root)
    assert projects == [{"project_id": PID, "name": "WorkLog", "source": "worklog"}]
    from weekly_report.weekly import select_dailies

    picked = select_dailies(wl_root, PID, "2026-W40")
    assert [d["daily_id"] for d in picked] == ["D-261001-00000011-01", "D-261004-00000010-01"]
    assert select_dailies(wl_root, PID, "2026-W41") == []


def run_week(root: Path, mock: Path, week: str, out: Path):
    client = ExaoneClient(root, mock_dir=mock)
    weekly_path, cum_path, report = run_weekly(root, PID, week, out, client=client)
    pptx = out / f"output/{week}.pptx"
    notes = generate_ppt(root, sources.load_project(root, PID), weekly_path, cum_path, root / TEMPLATE.name, pptx, client=client)
    return read(weekly_path), report.read_text(encoding="utf-8"), pptx, notes


def test_weekly_pipeline_uses_worklog_milestones_and_ignores_ai_updates(wl_root, tmp_path):
    mock = tmp_path / "mock"
    shutil.copytree(DEMO_MOCK, mock)
    name = f"weekly_rollup__{PID}__2026-W40.json"
    answer = read(mock / name)
    answer["milestone_updates"] = [{"milestone_id": "M2", "field": "actual", "to": "2026-10-04", "reason": "AI 추정", "source_ids": ["D-261004-00000010-01"]}]
    write(mock / name, answer)
    weekly, report, pptx, notes = run_week(wl_root, mock, "2026-W40", tmp_path / "out")
    assert weekly["milestone_updates"] == [] and "AI가 낸 일정 변화 1건 무시" in report
    assert weekly["milestone_snapshot"]["00000000-0000-4000-8000-000000000102"]["status"] == "진행"
    assert not [n for n in notes if n.startswith("PPT 검사 문제")]
    slide = Presentation(str(pptx)).slides[0]
    shapes = {s.name: s for s in slide.shapes}
    assert shapes["main_table"].table.cell(1, 1).text == "예시DX팀"
    ms = shapes["ms_table"].table
    assert [ms.cell(r, 1).text for r in range(1, 3)] == ["09/29~10/01", "10/01~10/04"] and ms.cell(2, 4).text == "–"
    assert shapes["author"].text_frame.text == "작성자 : 김예시"


def test_changed_milestones_are_blue_next_week(wl_root, tmp_path):
    mock = tmp_path / "mock"
    shutil.copytree(DEMO_MOCK, mock)
    run_week(wl_root, mock, "2026-W40", tmp_path / "out")
    # 다음 주: WorkLog에서 설계 완료 처리 (rev 19) + 10/05 업무일지
    export_path = wl_root / "data/worklog/WorkLog-sample/project.json"
    export = read(export_path)
    export["generatedAt"], export["sourceRevision"] = "2026-10-06T10:00:00+09:00", 19
    design = next(m for m in export["milestones"] if m["name"] == "설계")
    design.update(status="completed", actualEnd="2026-10-05")
    write(export_path, export)
    daily = read(wl_root / "data/worklog/WorkLog-sample/2026-10-04.json")
    daily["date"], daily["generatedAt"] = "2026-10-05", "2026-10-05T18:00:00+09:00"
    daily["logs"][0]["tasks"] = daily["logs"][0]["tasks"][:1]
    daily["logs"][0]["tasks"][0]["title"] = "설계 검토 회의 및 설계 완료"
    daily["logs"][0]["tasks"][0]["performedStart"] = daily["logs"][0]["tasks"][0]["performedEnd"] = "2026-10-05"
    daily["logs"][0]["issueRecords"] = daily["logs"][0]["todoRecords"] = []
    daily["logs"][0]["lessonLearned"] = None
    write(wl_root / "data/worklog/WorkLog-sample/2026-10-05.json", daily)
    src = "D-261005-00000010-01"
    write(mock / f"weekly_rollup__{PID}__2026-W41.json", {
        "headline": {"text": "금주(W41)에는 설계 검토 회의를 거쳐 설계 단계를 완료했습니다", "source_ids": [src]},
        "progress": [{"text": "AI 기반 프롬프트 설계 검토 회의를 진행하고 설계 단계를 완료했습니다. (10/05)", "source_ids": [src]}],
        "next_plan": [], "issues": [], "milestone_updates": []})
    write(mock / f"cumulative_update__{PID}__2026-W41.json", {"items": [
        {"text": "팀원 5명 인터뷰로 반복 보고 업무 3종을 도출해 요건 분석을 완료했습니다. (10/01)", "source_ids": ["D-261001-00000011-01"]},
        {"text": "AI 기반 프롬프트 설계 검토 회의를 거쳐 설계 단계를 완료했습니다. (10/05)", "source_ids": [src]}],
        "pinned_facts": [], "new_pinned_facts": []})
    weekly, report, pptx, notes = run_week(wl_root, mock, "2026-W41", tmp_path / "out")
    assert {(u["milestone_id"], u["field"], u["to"]) for u in weekly["milestone_updates"]} == {("M2", "status", "완료"), ("M2", "actual", "2026-10-05")}
    assert "WorkLog 일정 변경 2건" in report
    ms = {s.name: s for s in Presentation(str(pptx)).slides[0].shapes}["ms_table"].table
    row = next(r for r in range(1, len(ms.rows)) if ms.cell(r, 0).text == "설계")
    blue = lambda c: 'val="0000FF"' in ms.cell(row, c)._tc.xml
    assert ms.cell(row, 4).text == "10/05" and blue(4) and blue(5) and not blue(3)  # 실적·상태만 파랑
    other = next(r for r in range(1, len(ms.rows)) if ms.cell(r, 0).text == "요건 분석")
    assert not any('val="0000FF"' in ms.cell(other, c)._tc.xml for c in range(7))


def test_exec_summary_for_worklog_project(wl_root, tmp_path):
    from weekly_report.report.generate import generate_exec_summary

    mock = tmp_path / "mock"
    shutil.copytree(DEMO_MOCK, mock)
    run_week(wl_root, mock, "2026-W40", tmp_path / "out")
    shutil.copy(ROOT / "보고자료_Template_v1_초안.pptx", wl_root / "보고자료_Template_v1_초안.pptx")
    write(mock / f"report_exec_summary__{PID}__2026-W40.json", {
        "title": {"text": "WorkLog 보고 자동화 과제 진행 현황", "source_ids": [PID]},
        "head_message": {"text": "요건 분석을 완료하고 설계와 업무기록 웹 개발을 진행하고 있습니다", "source_ids": ["D-261001-00000011-01", "D-261004-00000010-01"]},
        "background_conclusion": [
            {"text": "보고서 작성 관련 VoE가 지속되어 AI Readable한 업무기록 축적을 추진하고 있습니다.", "source_ids": [PID]},
            {"text": "팀원 5명 인터뷰로 반복 보고 업무 3종을 도출해 요건 분석을 완료했습니다.", "source_ids": ["D-261001-00000011-01"]},
            {"text": "업무용 컴퓨터 Codex 사용 불가 문제를 해소하고 개발을 이어가겠습니다.", "source_ids": ["D-261004-00000010-01"]}],
        "emphasis": ["반복 보고 업무 3종"], "left_title": {"text": "추진 경과", "source_ids": [PID]},
        "left_items": [{"text": "[요건 분석] 팀원 5명 인터뷰, 반복 보고 업무 3종 도출 (10/1)", "source_ids": ["D-261001-00000011-01"]}],
        "right_title": {"text": "향후 계획", "source_ids": [PID]},
        "right_items": [{"text": "[개발] Claude 개발 버전 테스트", "source_ids": ["D-261004-00000010-01"]}]})
    result = generate_exec_summary(wl_root, PID, "2026-W40", tmp_path / "exec.pptx", client=ExaoneClient(wl_root, mock_dir=mock),
                                   dirs=[tmp_path / "out", wl_root], today=date(2026, 10, 4))
    assert result["problems"] == [] and result["notes"] == ["대체 없음 (AI 응답 그대로 사용)"], result["notes"]
    kpi = {s.name: s for s in Presentation(str(result["pptx"])).slides[0].shapes}["kpi_table"].table
    assert kpi.cell(1, 0).text == "KPI 없음"


def test_internal_examples_still_valid():
    for path in (ROOT / "data/master/projects").glob("*.json"):
        validate_schema(read(path), ROOT / "schemas/project.schema.json")
    assert {p["project_id"] for p in sources.list_projects(ROOT)} >= {"P-ASM-001", PID}
