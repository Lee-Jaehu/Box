"""보고 자료 템플릿 초안 만들기 (재현 가능한 스크립트).

python tools/make_report_template.py  →  보고자료_Template_v1_초안.pptx

- 주간 템플릿(주간업무PPT_Template_v2.pptx)을 열어 크기·테마·마스터를 물려받고, 슬라이드를 비운 뒤 3장을 그린다.
- 배치는 장표모음집 분석(docs/보고자료_양식_분석.md 1·3장)을 따른다.
  ① exec_summary (유형 A, 모음집 8번) ② monthly_overview (유형 D, 7번) ③ monthly_requests (유형 C, 5번)
- 모든 칸에 도형 이름을 붙인다. 렌더러는 이 이름으로 칸을 찾고, 여기서 정한 글자 크기·굵기를 복제해 쓴다.
- 공식 양식이 아니므로 하단에 "초안 – 공식 양식 아님"을 표시한다.
"""

from __future__ import annotations

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from weekly_report.fonts import EA_REGULAR, LATIN_FONT  # noqa: E402
from weekly_report.ppt.render import a  # noqa: E402

BASE = ROOT / "주간업무PPT_Template_v2.pptx"
OUTPUT = ROOT / "보고자료_Template_v1_초안.pptx"

BLACK, GRAY_TEXT, GREEN = "000000", "7F7F7F", "006600"
LABEL_FILL, HEADER_FILL, LINE = "F2F2F2", "D9D9D9", "7F7F7F"

# 표 열 너비(in)와 머리글: 렌더러가 같은 이름의 표를 채운다
PROJECT_COLUMNS = [("과제", 2.15), ("담당", 0.85), ("상태", 0.8), ("목표 일정 (지연 단계)", 1.75), ("KPI 기준 → 최신", 1.9), ("이번 달 주요 내용", 2.88)]
KPI_COLUMNS = [("지표", 1.6), ("기준", 0.65), ("최신 (기간)", 1.1), ("목표", 0.55), ("변화", 0.83)]
REQUEST_COLUMNS = [("과제", 2.3), ("요청 내용", 5.0), ("기한", 1.0), ("유관 부서", 2.03)]


def _fonts(rpr) -> None:
    etree.SubElement(rpr, a("latin")).set("typeface", LATIN_FONT)
    etree.SubElement(rpr, a("ea")).set("typeface", EA_REGULAR)


def style_runs(tf, size: float, bold: bool, color: str) -> None:
    for para in tf.paragraphs:
        runs = para.runs or [para.add_run()]
        for run in runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)
            rpr = run._r.get_or_add_rPr()
            for old in rpr.findall(a("latin")) + rpr.findall(a("ea")):
                rpr.remove(old)
            _fonts(rpr)
        end = para._p.get_or_add_endParaRPr()
        end.set("sz", str(int(size * 100)))
        end.set("b", "1" if bold else "0")
        for old in end.findall(a("latin")) + end.findall(a("ea")):
            end.remove(old)
        _fonts(end)


def textbox(slide, name: str, x: float, y: float, w: float, h: float, text: str, *, size: float, bold: bool = False,
            color: str = BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill: str | None = None,
            line: str | None = None, inset: float = 0.05, wrap: bool = True, space_before: float = 0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    style = shape._element.find(qn("p:style"))  # 기본 도형 스타일(그림자·흰 글자) 제거 → 서식은 아래에서 직접 지정
    if style is not None:
        shape._element.remove(style)
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(inset))
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = text.split("\n")
    tf.text = lines[0]
    for extra in lines[1:]:
        tf.add_paragraph().text = extra
    for para in tf.paragraphs:
        para.alignment = align
        if space_before:
            para.space_before = Pt(space_before)  # 렌더러가 첫 문단 pPr을 복제하므로 모든 문단에 적용됨
    style_runs(tf, size, bold, color)
    return shape


def hline(slide, x: float, y: float, w: float, color: str = LINE, width: float = 1.0) -> None:
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    style = conn._element.find(qn("p:style"))
    if style is not None:
        conn._element.remove(style)
    conn.line.color.rgb = RGBColor.from_string(color)
    conn.line.width = Pt(width)


def _cell_border(cell, color: str = LINE, width_pt: float = 0.75) -> None:
    """칸 테두리 (tcPr 자식 순서: lnL, lnR, lnT, lnB, …, 채우기)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for side in ("lnL", "lnR", "lnT", "lnB"):
        for old in tc_pr.findall(a(side)):
            tc_pr.remove(old)
    for index, side in enumerate(("lnL", "lnR", "lnT", "lnB")):
        ln = etree.Element(a(side), w=str(int(width_pt * 12700)), cmpd="sng")
        fill = etree.SubElement(ln, a("solidFill"))
        etree.SubElement(fill, a("srgbClr")).set("val", color)
        tc_pr.insert(index, ln)


def table(slide, name: str, x: float, y: float, columns: list[tuple[str, float]], *, header_h: float, row_h: float,
          rows: int = 1, size: float = 10, header_size: float = 10) -> None:
    width = sum(w for _, w in columns)
    shape = slide.shapes.add_table(rows + 1, len(columns), Inches(x), Inches(y), Inches(width), Inches(header_h + row_h * rows))
    shape.name = name
    tbl = shape.table
    tbl_pr = tbl._tbl.tblPr
    for flag in ("firstRow", "bandRow"):
        tbl_pr.set(flag, "0")
    style_id = tbl_pr.find(a("tableStyleId"))
    if style_id is not None:
        tbl_pr.remove(style_id)
    for col, (_, w) in enumerate(columns):
        tbl.columns[col].width = Inches(w)
    tbl.rows[0].height = Inches(header_h)
    for r in range(1, rows + 1):
        tbl.rows[r].height = Inches(row_h)
    for r in range(rows + 1):
        for col, (title, _) in enumerate(columns):
            cell = tbl.cell(r, col)
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = title if r == 0 else " "
            for para in cell.text_frame.paragraphs:
                para.alignment = PP_ALIGN.CENTER if r == 0 else PP_ALIGN.LEFT
            style_runs(cell.text_frame, header_size if r == 0 else size, r == 0, BLACK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(HEADER_FILL if r == 0 else "FFFFFF")
            _cell_border(cell)


def header_block(slide, title: str, head: str) -> None:
    textbox(slide, "title", 0.25, 0.15, 7.7, 0.45, title, size=20, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    textbox(slide, "org_date", 7.95, 0.12, 2.63, 0.5, "부서\n2026. 10. 4", size=10, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM)
    hline(slide, 0.25, 0.66, 10.33)
    textbox(slide, "head_message", 0.25, 0.74, 10.33, 0.66, head, size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def footer(slide) -> None:
    textbox(slide, "confidential", 0.25, 7.13, 1.35, 0.24, "대외비 (Confidential)", size=8, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, line=LINE)
    textbox(slide, "draft_mark", 1.7, 7.13, 2.6, 0.24, "초안 – 공식 양식 아님 (weekly-report 생성)", size=8, color=GRAY_TEXT,
            anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, "page_no", 5.0, 7.13, 0.83, 0.24, "1 / 1", size=10, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, "no_forward", 6.9, 7.13, 3.68, 0.24, "* 외부전달 및 임의변경 삼가를 부탁드립니다.", size=8, bold=True,
            align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def band_label(slide, name: str, y: float, h: float, text: str) -> None:
    textbox(slide, name, 0.25, y, 0.62, h, text, size=13, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
            fill=LABEL_FILL, line=LINE, inset=0.03)


def section_title(slide, name: str, x: float, y: float, w: float, text: str) -> None:
    textbox(slide, name, x, y, w, 0.3, text, size=12, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    hline(slide, x, y + 0.32, w, color=LINE, width=0.75)


def exec_summary(slide) -> None:
    header_block(slide, "과제명 진행 결과 및 향후 계획", "헤드메시지: 결론 · 핵심 수치 · 다음 조치를 경어체 1~2문장으로 씁니다.")
    band_label(slide, "band1_label", 1.5, 1.85, "보고\n배경\n및\n결론")
    textbox(slide, "background_conclusion", 0.95, 1.5, 9.63, 1.85, "• 배경·목적\n• 지금까지의 결론(핵심 수치)\n• 남은 과제와 조치·요청",
            size=13, bold=True, anchor=MSO_ANCHOR.MIDDLE, line=LINE, inset=0.12, space_before=4)
    band_label(slide, "band2_label", 3.45, 3.6, "상세\n검토\n결과")
    section_title(slide, "left_title", 0.95, 3.45, 4.7, "[ 추진 경과 및 성과 ]")
    textbox(slide, "left_items", 0.95, 3.85, 4.7, 3.2, "[ 소제목 ] 내용\n- 세부 내용", size=11, space_before=3)
    table(slide, "kpi_table", 5.85, 3.45, KPI_COLUMNS, header_h=0.28, row_h=0.3, size=10, header_size=10)
    textbox(slide, "schedule_note", 5.85, 4.08, 4.73, 0.4, "└ 일정: 지연 단계 (코드 계산)", size=9, color=GREEN)
    section_title(slide, "right_title", 5.85, 4.5, 4.73, "[ 향후 계획 및 요청 사항 ]")
    textbox(slide, "right_items", 5.85, 4.9, 4.73, 2.15, "[ 소제목 ] 내용\n▶ 요청: 내용", size=11, space_before=3)
    footer(slide)


def monthly_overview(slide) -> None:
    header_block(slide, "과제 종합 현황 / ’26.10월", "헤드메시지: 이번 달 전체 진척 결론 · 핵심 성과 수치 · 다음 달 조치를 씁니다.")
    textbox(slide, "unit_note", 7.58, 1.42, 3.0, 0.2, "(KPI 기준 → 최신, 지연 단계는 코드 계산값)", size=8, color=GRAY_TEXT, align=PP_ALIGN.RIGHT)
    table(slide, "project_table", 0.25, 1.62, PROJECT_COLUMNS, header_h=0.3, row_h=0.42, size=10, header_size=10)
    section_title(slide, "highlights_title", 0.25, 4.9, 5.05, "[ 주요 성과 ]")
    textbox(slide, "highlights", 0.25, 5.3, 5.05, 1.75, "- 성과", size=11, space_before=3)
    section_title(slide, "risks_title", 5.53, 4.9, 5.05, "[ 리스크 및 대응 ]")
    textbox(slide, "risks", 5.53, 5.3, 5.05, 1.75, "- 리스크 → 대응", size=11, space_before=3)
    footer(slide)


def monthly_requests(slide) -> None:
    header_block(slide, "의사결정 및 업무협조 요청 / ’26.10월", "헤드메시지: 요청 건수와 가장 급한 결정 사항을 씁니다.")
    table(slide, "requests_table", 0.25, 1.62, REQUEST_COLUMNS, header_h=0.3, row_h=0.5, size=11, header_size=10)
    footer(slide)


def build(base: Path = BASE, output: Path = OUTPUT) -> Path:
    prs = Presentation(str(base))
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):  # 주간 슬라이드 제거 (마스터·테마는 유지)
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)
    blank = prs.slide_layouts[0]
    for name, draw in (("exec_summary", exec_summary), ("monthly_overview", monthly_overview), ("monthly_requests", monthly_requests)):
        slide = prs.slides.add_slide(blank)
        slide.name = name
        slide._element.set("showMasterSp", "0")  # 마스터의 "‹#› / 140" 쪽 번호 숨김 → page_no 칸을 쓴다
        draw(slide)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return output


if __name__ == "__main__":
    print(build())
